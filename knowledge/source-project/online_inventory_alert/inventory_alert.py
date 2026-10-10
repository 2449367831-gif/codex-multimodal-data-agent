from __future__ import annotations

import argparse
import base64
import hashlib
import json
import logging
import math
import os
import re
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import datetime, timezone, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Sequence


CN_TZ = timezone(timedelta(hours=8))
VERSION = "2026.07.21-available-stock-query"

ALL_REGIONS = [
    "北京市",
    "天津市",
    "河北省",
    "山西省",
    "内蒙古自治区",
    "辽宁省",
    "吉林省",
    "黑龙江省",
    "上海市",
    "江苏省",
    "浙江省",
    "安徽省",
    "福建省",
    "江西省",
    "山东省",
    "河南省",
    "湖北省",
    "湖南省",
    "广东省",
    "广西壮族自治区",
    "海南省",
    "重庆市",
    "四川省",
    "贵州省",
    "云南省",
    "西藏自治区",
    "陕西省",
    "甘肃省",
    "青海省",
    "宁夏回族自治区",
    "新疆维吾尔自治区",
]

REGION_ALIASES = {
    "北京": "北京市",
    "北京市": "北京市",
    "天津": "天津市",
    "天津市": "天津市",
    "河北": "河北省",
    "河北省": "河北省",
    "山西": "山西省",
    "山西省": "山西省",
    "内蒙古": "内蒙古自治区",
    "内蒙古自治区": "内蒙古自治区",
    "辽宁": "辽宁省",
    "辽宁省": "辽宁省",
    "吉林": "吉林省",
    "吉林省": "吉林省",
    "黑龙江": "黑龙江省",
    "黑龙江省": "黑龙江省",
    "上海": "上海市",
    "上海市": "上海市",
    "江苏": "江苏省",
    "江苏省": "江苏省",
    "浙江": "浙江省",
    "浙江省": "浙江省",
    "安徽": "安徽省",
    "安徽省": "安徽省",
    "福建": "福建省",
    "福建省": "福建省",
    "江西": "江西省",
    "江西省": "江西省",
    "山东": "山东省",
    "山东省": "山东省",
    "河南": "河南省",
    "河南省": "河南省",
    "湖北": "湖北省",
    "湖北省": "湖北省",
    "湖南": "湖南省",
    "湖南省": "湖南省",
    "广东": "广东省",
    "广东省": "广东省",
    "广西": "广西壮族自治区",
    "广西壮族自治区": "广西壮族自治区",
    "海南": "海南省",
    "海南省": "海南省",
    "重庆": "重庆市",
    "重庆市": "重庆市",
    "四川": "四川省",
    "四川省": "四川省",
    "贵州": "贵州省",
    "贵州省": "贵州省",
    "云南": "云南省",
    "云南省": "云南省",
    "西藏": "西藏自治区",
    "西藏自治区": "西藏自治区",
    "陕西": "陕西省",
    "陕西省": "陕西省",
    "甘肃": "甘肃省",
    "甘肃省": "甘肃省",
    "甘": "甘肃省",
    "青海": "青海省",
    "青海省": "青海省",
    "青": "青海省",
    "宁夏": "宁夏回族自治区",
    "宁夏回族自治区": "宁夏回族自治区",
    "宁": "宁夏回族自治区",
    "新疆": "新疆维吾尔自治区",
    "新疆维吾尔自治区": "新疆维吾尔自治区",
}


@dataclass(frozen=True)
class ProductRule:
    sku_code: str
    sku_name: str
    strategy: str
    product_type: str
    category: str
    replacement_code: str = ""
    alert_group_code: str = ""
    alert_group_name: str = ""
    alert_group_sku_codes: tuple[str, ...] = ()


@dataclass(frozen=True)
class WarehouseRule:
    wms_desc: str
    short_name: str
    region_text: str
    raw_region_text: str


def app_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


ROOT = app_root()
DEFAULT_CONFIG = ROOT / "config" / "config.env"
DEFAULT_STATE = ROOT / "data" / "last_status_region.json"
DEFAULT_SUMMARY = ROOT / "data" / "last_run_summary.json"
DEFAULT_LOG_DIR = ROOT / "logs"


def now_cn() -> datetime:
    return datetime.now(CN_TZ)


def text(value) -> str:
    return "" if value is None else str(value).strip()


def to_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return number if math.isfinite(number) else 0.0


def parse_bool(value: str | None, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on", "是", "启用"}


def parse_list(value: str | None) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[|,，;；\r\n]+", value)
    return [part.strip() for part in parts if part and part.strip()]


def load_env_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"配置文件不存在：{path}")
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key:
            os.environ[key] = value


def env(name: str, default: str = "") -> str:
    return os.environ.get(name, default).strip()


def resolve_path(value: str, default_path: Path) -> Path:
    raw = (value or "").strip()
    if not raw:
        return default_path
    path = Path(raw)
    if not path.is_absolute():
        path = ROOT / path
    return path


def setup_logging(log_dir: Path) -> None:
    log_dir.mkdir(parents=True, exist_ok=True)
    cleanup_old_logs(log_dir, int(env("LOG_RETENTION_DAYS", "3") or "3"))
    log_file = log_dir / f"inventory_alert_{now_cn().strftime('%Y%m%d')}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(log_file, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )


def cleanup_old_logs(log_dir: Path, retention_days: int) -> None:
    if retention_days <= 0:
        return
    cutoff = now_cn() - timedelta(days=retention_days - 1)
    cutoff_date = cutoff.date()
    for pattern in ("inventory_alert_*.log", "inventory_alert_detail_*.jsonl"):
        for path in log_dir.glob(pattern):
            try:
                if datetime.fromtimestamp(path.stat().st_mtime, CN_TZ).date() < cutoff_date:
                    path.unlink()
            except OSError:
                pass


def require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("缺少依赖 openpyxl，请执行：py -m pip install -r requirements.txt") from exc
    return load_workbook


def header_index(header: Iterable[object]) -> dict[str, int]:
    return {text(value): idx for idx, value in enumerate(header) if text(value)}


def first_existing_column(idx: dict[str, int], names: Iterable[str]) -> int | None:
    for name in names:
        if name in idx:
            return idx[name]
    return None


def attach_replacement_groups(products: dict[str, ProductRule]) -> tuple[dict[str, ProductRule], dict]:
    parent = {sku_code: sku_code for sku_code in products}
    ignored_replacements: list[dict] = []

    def find(sku_code: str) -> str:
        root = parent.setdefault(sku_code, sku_code)
        if root != sku_code:
            parent[sku_code] = find(root)
        return parent[sku_code]

    def union(left: str, right: str) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    for sku_code, product in products.items():
        replacement_code = product.replacement_code
        if not replacement_code:
            continue
        if replacement_code not in products:
            ignored_replacements.append({
                "sku_code": sku_code,
                "replacement_code": replacement_code,
                "reason": "替代品不在当前监控范围内",
            })
            continue
        union(sku_code, replacement_code)

    groups: dict[str, list[str]] = defaultdict(list)
    for sku_code in products:
        groups[find(sku_code)].append(sku_code)

    grouped_products: dict[str, ProductRule] = {}
    replacement_group_examples: list[dict] = []
    for sku_codes in groups.values():
        ordered_codes = tuple(sorted(sku_codes))
        if len(ordered_codes) == 1:
            sku_code = ordered_codes[0]
            product = products[sku_code]
            grouped_products[sku_code] = replace(
                product,
                alert_group_code=sku_code,
                alert_group_name=product.sku_name,
                alert_group_sku_codes=ordered_codes,
            )
            continue

        group_code = "replacement:" + ",".join(ordered_codes)
        names: list[str] = []
        for sku_code in ordered_codes:
            sku_name = products[sku_code].sku_name
            if sku_name and sku_name not in names:
                names.append(sku_name)
        group_name = " / ".join(names)
        replacement_group_examples.append({
            "group_code": group_code,
            "sku_codes": list(ordered_codes),
            "sku_names": names,
        })
        for sku_code in ordered_codes:
            product = products[sku_code]
            grouped_products[sku_code] = replace(
                product,
                alert_group_code=group_code,
                alert_group_name=group_name,
                alert_group_sku_codes=ordered_codes,
            )

    audit = {
        "replacement_group_count": len(replacement_group_examples),
        "replacement_group_examples": replacement_group_examples[:20],
        "ignored_replacement_count": len(ignored_replacements),
        "ignored_replacement_examples": ignored_replacements[:20],
    }
    return grouped_products, audit


def load_base_excel(path: Path) -> tuple[dict[str, ProductRule], list[WarehouseRule], dict]:
    load_workbook = require_openpyxl()
    if not path.exists():
        raise FileNotFoundError(f"基础数据文件不存在：{path}")

    product_sheet = env("PRODUCT_SHEET", "产品信息表")
    warehouse_sheet = env("WAREHOUSE_SHEET", "仓库映射表")
    allowed_categories = set(parse_list(env("ALLOWED_PRODUCT_CATEGORIES", "产成品")))
    full_strategy_names = set(parse_list(env("FULL_STRATEGY_NAMES", "全铺")))
    ambient_strategy_names = set(parse_list(env("AMBIENT_STRATEGY_NAMES", "常温（成都万纬、南京普洛斯、雨花常温、郑州常温）")))

    wb = load_workbook(path, read_only=True, data_only=True)
    if product_sheet not in wb.sheetnames:
        raise ValueError(f"基础数据缺少工作表：{product_sheet}")
    if warehouse_sheet not in wb.sheetnames:
        raise ValueError(f"基础数据缺少工作表：{warehouse_sheet}")

    products: dict[str, ProductRule] = {}
    duplicate_products: list[str] = []
    skipped_strategy = Counter()
    skipped_category = Counter()

    ws = wb[product_sheet]
    rows = ws.iter_rows(values_only=True)
    try:
        idx = header_index(next(rows))
    except StopIteration:
        raise ValueError(f"{product_sheet} 是空表")

    required_cols = ["产品编码", "商品名称", "布仓策略"]
    missing = [col for col in required_cols if col not in idx]
    if missing:
        raise ValueError(f"{product_sheet} 缺少列：{', '.join(missing)}")

    category_col = idx.get("产品大类")
    replacement_col = first_existing_column(idx, ["产品替换", "商品替换", "替代品编码", "替代SKU", "替代品"])
    for row in rows:
        sku_code = text(row[idx["产品编码"]] if idx["产品编码"] < len(row) else "")
        if not sku_code:
            continue
        sku_name = text(row[idx["商品名称"]] if idx["商品名称"] < len(row) else "") or sku_code
        strategy = text(row[idx["布仓策略"]] if idx["布仓策略"] < len(row) else "")
        category = text(row[category_col] if category_col is not None and category_col < len(row) else "")

        if allowed_categories and category not in allowed_categories:
            skipped_category[category or "<空>"] += 1
            continue

        if strategy in full_strategy_names:
            product_type = "frozen"
        elif strategy in ambient_strategy_names or "常温" in strategy:
            product_type = "ambient"
        else:
            skipped_strategy[strategy or "<空>"] += 1
            continue

        if sku_code in products:
            duplicate_products.append(sku_code)
            continue
        products[sku_code] = ProductRule(
            sku_code=sku_code,
            sku_name=sku_name,
            strategy=strategy,
            product_type=product_type,
            category=category,
            replacement_code=text(row[replacement_col] if replacement_col is not None and replacement_col < len(row) else ""),
        )

    products, replacement_audit = attach_replacement_groups(products)

    warehouses: list[WarehouseRule] = []
    ws = wb[warehouse_sheet]
    rows = ws.iter_rows(values_only=True)
    try:
        idx = header_index(next(rows))
    except StopIteration:
        raise ValueError(f"{warehouse_sheet} 是空表")

    required_cols = ["wms_desc", "仓库简称", "可发省份"]
    missing = [col for col in required_cols if col not in idx]
    if missing:
        raise ValueError(f"{warehouse_sheet} 缺少列：{', '.join(missing)}")

    raw_region_col = idx.get("可发区域_原文")
    seen_wms = set()
    duplicate_wms: list[str] = []
    for row in rows:
        wms_desc = text(row[idx["wms_desc"]] if idx["wms_desc"] < len(row) else "")
        if not wms_desc:
            continue
        if wms_desc in seen_wms:
            duplicate_wms.append(wms_desc)
            continue
        seen_wms.add(wms_desc)
        warehouses.append(
            WarehouseRule(
                wms_desc=wms_desc,
                short_name=text(row[idx["仓库简称"]] if idx["仓库简称"] < len(row) else "") or wms_desc,
                region_text=text(row[idx["可发省份"]] if idx["可发省份"] < len(row) else ""),
                raw_region_text=text(row[raw_region_col] if raw_region_col is not None and raw_region_col < len(row) else ""),
            )
        )

    audit = {
        "base_excel": str(path),
        "product_count": len(products),
        "warehouse_count": len(warehouses),
        "duplicate_products": sorted(set(duplicate_products)),
        "duplicate_wms": duplicate_wms,
        "skipped_strategy": dict(skipped_strategy),
        "skipped_category": dict(skipped_category),
        "replacement_col_found": replacement_col is not None,
        **replacement_audit,
    }
    return products, warehouses, audit


def detect_regions(fragment: str, allow_single_char_alias: bool = False) -> set[str]:
    found: set[str] = set()
    normalized = fragment.strip()
    if not normalized:
        return found

    for region in ALL_REGIONS:
        if region in normalized:
            found.add(region)

    tokens = [part.strip() for part in re.split(r"[、,，/&\s]+", normalized) if part.strip()]
    for token in tokens:
        if token in REGION_ALIASES:
            found.add(REGION_ALIASES[token])

    for alias, region in REGION_ALIASES.items():
        if len(alias) == 1 and not allow_single_char_alias:
            continue
        if alias in normalized:
            found.add(region)

    return found


def parse_region_text(region_text: str, product_type: str) -> list[str]:
    text_value = text(region_text)
    if not text_value:
        return []

    selected_parts: list[str] = []
    labeled = False
    for segment in re.split(r"[;；\r\n]+", text_value):
        segment = segment.strip()
        if not segment:
            continue
        if "：" in segment or ":" in segment:
            labeled = True
            label, content = re.split(r"[：:]", segment, maxsplit=1)
            label = label.strip()
            content = content.strip()
            if product_type == "ambient" and "常温" in label:
                selected_parts.append(content)
            elif product_type != "ambient" and ("冻" in label or "冷冻" in label):
                selected_parts.append(content)
        else:
            selected_parts.append(segment)

    if labeled and not selected_parts:
        selected_parts = [text_value]

    regions: set[str] = set()
    for part in selected_parts:
        if "全国" in part and ("不含" in part or "除" in part):
            excluded_part = part
            if "不含" in part:
                excluded_part = part.split("不含", 1)[1]
            elif "除" in part:
                excluded_part = part.split("除", 1)[1]
            excluded = detect_regions(excluded_part, allow_single_char_alias=True)
            regions.update(region for region in ALL_REGIONS if region not in excluded)
            continue
        if "全国" in part:
            regions.update(ALL_REGIONS)
            continue
        regions.update(detect_regions(part))

    return [region for region in ALL_REGIONS if region in regions]


def warehouse_allowed(product: ProductRule, warehouse: WarehouseRule) -> bool:
    ambient_warehouses = set(parse_list(env(
        "AMBIENT_WAREHOUSES",
        "皇家小虎南京-普洛斯仓xs|皇家小虎长沙雨花常温-万纬仓xs|皇家小虎郑州-常温品代发仓xs|皇家小虎成都-万纬仓",
    )))
    ambient_only_warehouses = set(parse_list(env(
        "AMBIENT_ONLY_WAREHOUSES",
        "皇家小虎长沙雨花常温-万纬仓xs|皇家小虎郑州-常温品代发仓xs|皇家小虎成都-万纬仓",
    )))
    if product.product_type == "ambient":
        return warehouse.wms_desc in ambient_warehouses
    if warehouse.wms_desc in ambient_only_warehouses:
        return False
    return True


def connect_doris():
    try:
        import pymysql
    except ImportError as exc:
        raise SystemExit("缺少依赖 PyMySQL，请执行：py -m pip install -r requirements.txt") from exc

    host = env("DB_HOST") or env("PICKUP_DB_HOST")
    user = env("DB_USER") or env("PICKUP_DB_USER")
    password = os.environ.get("DB_PASSWORD", os.environ.get("PICKUP_DB_PASSWORD", ""))
    database = env("DB_NAME") or env("PICKUP_DB_NAME", "ywzt_ods") or "ywzt_ods"
    port = int(env("DB_PORT") or env("PICKUP_DB_PORT", "9030") or "9030")

    missing = [name for name, value in {
        "DB_HOST": host,
        "DB_USER": user,
        "DB_PASSWORD": password,
    }.items() if not value]
    if missing:
        raise SystemExit(f"缺少 Doris 配置：{', '.join(missing)}")

    import pymysql

    return pymysql.connect(
        host=host,
        port=port,
        user=user,
        password=password,
        database=database,
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=int(env("DB_READ_TIMEOUT", "300")),
        write_timeout=60,
    )


STOCK_SQL = """
WITH tall AS (
    SELECT p.take_delivery_wms AS wms_code, p.cargo_owner_id, d.sku_code
    FROM ywzt_ods.scm_po_header p
    LEFT JOIN ywzt_ods.scm_po_detail d ON p.po_code = d.po_code
    WHERE p.order_status IN ('10', '11', '12', '13', '14', '15', '16')
      AND p.cargo_owner_id = 'HJXH'

    UNION
    SELECT wms_code, cargo_owner_id, sku_code
    FROM ywzt_ods.scm_warehous_manage
    WHERE cargo_owner_id = 'HJXH'

    UNION
    SELECT wms_code, cargo_owner_id, sku_code
    FROM ywzt_ods.scm_stock_history
    WHERE cargo_owner_id = 'HJXH'

    UNION
    SELECT p.receive_wms AS wms_code, p.receiver_cargo_owner AS cargo_owner_id, d.sku_code
    FROM ywzt_ods.scm_transfer_order_header p
    LEFT JOIN ywzt_ods.scm_transfer_order_detail d
        ON p.transfer_order_code = d.transfer_order_code
    WHERE p.order_status IN ('13', '14', '15')
      AND p.receiver_cargo_owner = 'HJXH'
),
inventory AS (
    SELECT
        wms_code,
        cargo_owner_id,
        sku_code,
        MAX(sku_name) AS sku_name,
        SUM(IFNULL(in_stock_quantity, 0)) AS in_stock_quantity,
        SUM(IFNULL(CAST(NULLIF(stock_lock, '') AS DECIMAL(20, 3)), 0)) AS stock_lock,
        SUM(IFNULL(freeze_quantity, 0)) AS freeze_qty,
        COUNT(*) AS inventory_row_count
    FROM ywzt_ods.scm_warehous_manage
    WHERE (inventory_status = 'ZP' OR inventory_status IS NULL OR inventory_status = '')
      AND cargo_owner_id = 'HJXH'
    GROUP BY cargo_owner_id, wms_code, sku_code
),
allocate AS (
    SELECT
        wms_code,
        cargo_owner_id,
        sku_code,
        SUM(IFNULL(freeze_quantity, 0)) AS allocate_qty
    FROM ywzt_ods.scm_warehous_occupy
    WHERE created_date > DATE_SUB(NOW(), INTERVAL 60 DAY)
      AND status != 1
      AND cargo_owner_id = 'HJXH'
    GROUP BY cargo_owner_id, sku_code, wms_code
)
SELECT
    tall.wms_code AS warehouse_code,
    mwms.wms_desc AS warehouse_name,
    tall.sku_code,
    COALESCE(ms.sku_name, inventory.sku_name) AS sku_name,
    GREATEST(
        IFNULL(inventory.in_stock_quantity, 0)
        - IFNULL(allocate.allocate_qty, 0)
        - CASE
            WHEN IFNULL(inventory.stock_lock, 0) = 0 THEN IFNULL(inventory.freeze_qty, 0)
            ELSE IFNULL(inventory.stock_lock, 0)
          END,
        0
    ) AS stock_qty,
    IFNULL(inventory.in_stock_quantity, 0) AS all_stock_qty,
    IFNULL(allocate.allocate_qty, 0) AS allocate_qty,
    IFNULL(inventory.stock_lock, 0) AS stock_lock_qty,
    IFNULL(inventory.freeze_qty, 0) AS freeze_qty,
    IFNULL(inventory.inventory_row_count, 0) AS inventory_row_count
FROM tall
LEFT JOIN inventory
    ON tall.wms_code = inventory.wms_code
   AND tall.cargo_owner_id = inventory.cargo_owner_id
   AND tall.sku_code = inventory.sku_code
LEFT JOIN allocate
    ON tall.wms_code = allocate.wms_code
   AND tall.cargo_owner_id = allocate.cargo_owner_id
   AND tall.sku_code = allocate.sku_code
LEFT JOIN ywzt_ods.master_wms mwms ON tall.wms_code = mwms.wms_code
LEFT JOIN ywzt_ods.master_sku ms ON ms.sku_code = tall.sku_code
WHERE mwms.wms_type IN (10, 11)
  AND mwms.wms_desc NOT LIKE '%残次%'
"""


SALES_SQL = """
WITH shipment_by_day AS (
    SELECT
        DATE_FORMAT(h.ship_at, '%Y-%m-%d') AS ship_date,
        h.facility_code AS warehouse_code,
        w.wms_desc AS warehouse_name,
        i.sku_code AS sku_code,
        m.sku_name AS sku_name,
        SUM(i.ship_qty) AS total_ship_qty,
        SUM(CASE WHEN s.type = 'Z004' THEN i.ship_qty ELSE 0 END) AS online_ship_qty,
        SUM(CASE WHEN s.type = 'Z004' THEN 0 ELSE i.ship_qty END) AS offline_ship_qty
    FROM ywzt_ods.scm_sales_shipment_header h
    INNER JOIN ywzt_ods.scm_sales_shipment_detail i
        ON i.outbound_code = h.outbound_code
    LEFT JOIN ywzt_ods.master_store s
        ON s.code = h.store_code
    LEFT JOIN ywzt_ods.master_sku m
        ON m.sku_code = i.sku_code
    LEFT JOIN ywzt_ods.master_wms w
        ON w.wms_code = h.facility_code
    WHERE h.ship_at >= CURDATE() - INTERVAL 7 DAY
      AND h.ship_at < CURDATE()
      AND h.outbound_code LIKE 'HJXH%'
      AND h.status IN ('700', '900', '701', '901')
      AND (
          h.ref_order_type IN ('JYCK', 'HHCK', 'BFCK', 'QTCK')
          OR h.ref_order_type LIKE 'SO%'
      )
    GROUP BY
        DATE_FORMAT(h.ship_at, '%Y-%m-%d'),
        h.facility_code,
        w.wms_desc,
        i.sku_code,
        m.sku_name
)
SELECT
    warehouse_code,
    warehouse_name,
    sku_code,
    MAX(sku_name) AS sku_name,
    SUM(total_ship_qty) AS total_sales_7d_qty,
    SUM(online_ship_qty) AS online_sales_7d_qty,
    SUM(offline_ship_qty) AS offline_sales_7d_qty
FROM shipment_by_day
GROUP BY
    warehouse_code,
    warehouse_name,
    sku_code
"""


def fetch_rows(connection, sql: str) -> list[dict]:
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return list(cursor.fetchall())


def fetch_doris_data() -> tuple[list[dict], list[dict]]:
    connection = connect_doris()
    try:
        stock_rows = fetch_rows(connection, STOCK_SQL)
    finally:
        connection.close()
    return stock_rows, []


def row_warehouse_name(row: dict) -> str:
    return text(row.get("warehouse_name"))


def build_quantity_maps(stock_rows: list[dict], sales_rows: list[dict]) -> tuple[dict[tuple[str, str], float], dict[tuple[str, str], float], dict]:
    stock_map: dict[tuple[str, str], float] = defaultdict(float)
    sales_map: dict[tuple[str, str], float] = defaultdict(float)

    for row in stock_rows:
        warehouse = row_warehouse_name(row)
        sku_code = text(row.get("sku_code"))
        if warehouse and sku_code and to_number(row.get("inventory_row_count", 1)) > 0:
            stock_map[(warehouse, sku_code)] += to_number(row.get("stock_qty"))

    audit = {
        "stock_row_count": len(stock_rows),
        "sales_row_count": len(sales_rows),
        "stock_key_count": len(stock_map),
        "sales_key_count": len(sales_map),
        "sales_scope": "not_queried",
    }
    return stock_map, sales_map, audit


def format_qty(value: float) -> int | float:
    if abs(value - round(value)) < 0.000001:
        return int(round(value))
    return round(value, 2)


def build_current_snapshot(
    products: dict[str, ProductRule],
    warehouses: list[WarehouseRule],
    stock_map: dict[tuple[str, str], float],
    sales_map: dict[tuple[str, str], float],
) -> tuple[dict, dict]:
    low_stock_qty = float(env("LOW_STOCK_QTY", "10"))
    regions_by_warehouse_type: dict[tuple[str, str], list[str]] = {}
    missing_region: list[dict] = []
    warehouse_short = {warehouse.wms_desc: warehouse.short_name for warehouse in warehouses}

    items: dict[str, dict] = {}
    aggregate_audit = {
        "empty_region_mapping_count": 0,
        "empty_region_mapping_examples": [],
        "candidate_region_sku_count": 0,
        "replacement_group_region_count": 0,
    }
    group_details: dict[tuple[str, str], dict] = {}

    for product in products.values():
        for warehouse in warehouses:
            if not warehouse_allowed(product, warehouse):
                continue
            cache_key = (warehouse.wms_desc, product.product_type)
            if cache_key not in regions_by_warehouse_type:
                regions_by_warehouse_type[cache_key] = parse_region_text(warehouse.region_text, product.product_type)
            regions = regions_by_warehouse_type[cache_key]
            if not regions:
                missing_region.append({
                    "warehouse": warehouse.wms_desc,
                    "product_type": product.product_type,
                    "region_text": warehouse.region_text,
                })
                continue

            stock_key = (warehouse.wms_desc, product.sku_code)
            stock_qty = stock_map.get(stock_key, 0.0)
            stock_key_observed = stock_key in stock_map
            sales_7d_qty = sales_map.get((warehouse.wms_desc, product.sku_code), 0.0)
            alert_group_code = product.alert_group_code or product.sku_code
            alert_group_sku_codes = tuple(product.alert_group_sku_codes or (product.sku_code,))
            for region in regions:
                detail = group_details.setdefault(
                    (region, alert_group_code),
                    {
                        "region": region,
                        "sku_code": alert_group_code,
                        "sku_codes": set(alert_group_sku_codes),
                        "sku_name": product.alert_group_name or product.sku_name,
                        "strategies": set(),
                        "product_types": set(),
                        "stock_qty": 0.0,
                        "sales_7d_qty": 0.0,
                        "warehouses": [],
                        "observed_stock_keys": set(),
                    },
                )
                detail["sku_codes"].add(product.sku_code)
                detail["strategies"].add(product.strategy)
                detail["product_types"].add(product.product_type)
                detail["stock_qty"] += stock_qty
                detail["sales_7d_qty"] += sales_7d_qty
                detail["warehouses"].append(warehouse_short.get(warehouse.wms_desc, warehouse.wms_desc))
                if stock_key_observed:
                    detail["observed_stock_keys"].add(f"{warehouse.wms_desc}\t{product.sku_code}")

    for (region, group_code), detail in group_details.items():
        avg_daily_sales = detail["sales_7d_qty"] / 7 if detail["sales_7d_qty"] > 0 else 0.0
        inventory_days = detail["stock_qty"] / avg_daily_sales if avg_daily_sales > 0 else None
        low_stock = detail["stock_qty"] < low_stock_qty
        sku_codes = sorted(detail["sku_codes"])
        key = f"{region}|{group_code}"
        items[key] = {
            "region": region,
            "sku_code": group_code,
            "sku_codes": sku_codes,
            "sku_name": detail["sku_name"],
            "strategy": "、".join(sorted(item for item in detail["strategies"] if item)),
            "product_type": "、".join(sorted(item for item in detail["product_types"] if item)),
            "replacement_group": len(sku_codes) > 1,
            "stock_qty": format_qty(detail["stock_qty"]),
            "sales_7d_qty": format_qty(detail["sales_7d_qty"]),
            "avg_daily_sales": round(avg_daily_sales, 4),
            "inventory_days": None if inventory_days is None else round(inventory_days, 2),
            "low_stock": low_stock,
            "warehouses": sorted(set(detail["warehouses"])),
            "observed_stock_keys": sorted(detail["observed_stock_keys"]),
            "legacy_keys": [f"{region}|{sku_code}" for sku_code in sku_codes],
        }

    if missing_region:
        aggregate_audit["empty_region_mapping_count"] = len(missing_region)
        aggregate_audit["empty_region_mapping_examples"] = missing_region[:20]
    aggregate_audit["candidate_region_sku_count"] = len(items)
    aggregate_audit["replacement_group_region_count"] = sum(1 for item in items.values() if item["replacement_group"])

    relevant_stock_keys = {
        stock_key
        for item in items.values()
        for stock_key in item.get("observed_stock_keys", [])
    }
    snapshot = {
        "version": 1,
        "program_version": VERSION,
        "updated_at": now_cn().strftime("%Y-%m-%d %H:%M:%S"),
        "low_stock_qty": low_stock_qty,
        "restock_rule": "stock_qty_gte_low_stock_qty",
        "source_audit": {
            "relevant_stock_key_count": len(relevant_stock_keys),
        },
        "items": items,
    }
    return snapshot, aggregate_audit


def load_state(path: Path) -> dict | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"状态文件不是有效 JSON：{path}") from exc


def previous_low_stock_for_item(key: str, curr: dict, previous_items: dict) -> bool | None:
    prev = previous_items.get(key)
    if prev is not None:
        return bool(prev.get("low_stock"))

    legacy_keys = curr.get("legacy_keys") or []
    legacy_states = [previous_items[item_key] for item_key in legacy_keys if item_key in previous_items]
    if not legacy_states:
        return None
    if all("stock_qty" in item for item in legacy_states):
        legacy_stock_qty = sum(to_number(item.get("stock_qty")) for item in legacy_states)
        low_stock_qty = float(env("LOW_STOCK_QTY", "10"))
        return legacy_stock_qty < low_stock_qty
    return any(bool(item.get("low_stock")) for item in legacy_states)


def atomic_write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    for attempt in range(10):
        try:
            tmp.replace(path)
            return
        except PermissionError:
            time.sleep(0.3 * (attempt + 1))
    tmp.replace(path)


def compare_snapshots(previous: dict | None, current: dict) -> tuple[list[dict], list[dict], list[dict]]:
    push_on_first_run = parse_bool(env("PUSH_ON_FIRST_RUN", "false"))
    current_items = current.get("items", {})
    previous_items = {} if previous is None else previous.get("items", {})

    stockout_alerts: list[dict] = []
    restock_alerts: list[dict] = []
    skipped_restock: list[dict] = []

    for key, curr in current_items.items():
        prev_low = previous_low_stock_for_item(key, curr, previous_items)
        if prev_low is None:
            if push_on_first_run and curr.get("low_stock"):
                stockout_alerts.append(curr)
            continue

        curr_low = bool(curr.get("low_stock"))

        if not prev_low and curr_low:
            stockout_alerts.append(curr)
        elif prev_low and not curr_low:
            restock_alerts.append(curr)

    stockout_alerts.sort(key=lambda item: (item["region"], item["sku_name"], item["sku_code"]))
    restock_alerts.sort(key=lambda item: (item["region"], item["sku_name"], item["sku_code"]))
    skipped_restock.sort(key=lambda item: (item["region"], item["sku_name"], item["sku_code"]))
    return stockout_alerts, restock_alerts, skipped_restock


def item_source_keys(item: dict | None) -> set[str]:
    if not item:
        return set()
    return {text(value) for value in item.get("observed_stock_keys", []) if text(value)}


def inspect_snapshot_quality(previous: dict | None, current: dict) -> dict:
    previous_items = {} if previous is None else previous.get("items", {})
    current_items = current.get("items", {})
    min_source_key_ratio = float(env("MIN_SOURCE_KEY_RATIO", "0.7") or "0.7")
    previous_key_count = to_number((previous or {}).get("source_audit", {}).get("relevant_stock_key_count"))
    current_key_count = to_number(current.get("source_audit", {}).get("relevant_stock_key_count"))
    global_reasons: list[str] = []

    if current_key_count <= 0:
        global_reasons.append("no_relevant_stock_rows")
    elif previous_key_count > 0 and current_key_count < previous_key_count * min_source_key_ratio:
        global_reasons.append("relevant_stock_rows_dropped")

    item_issues: list[dict] = []
    suspicious_low_keys: list[str] = []
    for key, curr in current_items.items():
        prev_item = previous_item_for_audit(key, curr, previous_items)
        if prev_item is None:
            continue
        prev_source_keys = item_source_keys(prev_item)
        curr_source_keys = item_source_keys(curr)
        missing_source_keys = sorted(prev_source_keys - curr_source_keys)
        reason = ""
        if not curr_source_keys and (
            prev_source_keys
            or (not bool(prev_item.get("low_stock")) and bool(curr.get("low_stock")))
        ):
            reason = "no_current_source_record"
        elif missing_source_keys:
            reason = "previous_source_record_missing"
        if reason:
            item_issues.append({
                "key": key,
                "reason": reason,
                "prev_stock_qty": prev_item.get("stock_qty"),
                "curr_stock_qty": curr.get("stock_qty"),
                "missing_source_keys": missing_source_keys,
            })
        if not bool(prev_item.get("low_stock")) and bool(curr.get("low_stock")):
            suspicious_low_keys.append(key)

    return {
        "global_reasons": global_reasons,
        "item_issues": item_issues,
        "suspicious_low_keys": suspicious_low_keys,
        "previous_relevant_stock_key_count": format_qty(previous_key_count),
        "current_relevant_stock_key_count": format_qty(current_key_count),
        "min_source_key_ratio": min_source_key_ratio,
    }


def apply_data_quality_guard(previous: dict | None, current: dict, quality: dict) -> tuple[dict, dict]:
    required_confirmations = max(1, int(env("LOW_STOCK_CONFIRM_RUNS", "2") or "2"))
    previous_items = {} if previous is None else previous.get("items", {})
    guarded = dict(current)
    guarded_items = {key: dict(item) for key, item in current.get("items", {}).items()}
    guarded["items"] = guarded_items
    blocked_records: list[dict] = []
    pending_records: list[dict] = []

    issues_by_key = {issue["key"]: issue for issue in quality.get("item_issues", [])}
    for key, issue in issues_by_key.items():
        curr = current.get("items", {}).get(key)
        if curr is None:
            continue
        prev_item = previous_item_for_audit(key, curr, previous_items)
        if prev_item is None:
            guarded_items.pop(key, None)
        else:
            preserved = dict(prev_item)
            preserved["data_quality_blocked_at"] = current.get("updated_at")
            preserved["data_quality_blocked_reason"] = issue.get("reason")
            guarded_items[key] = preserved
        blocked_records.append({**issue, "current_item": curr, "previous_item": prev_item})

    for key, curr in list(guarded_items.items()):
        if key in issues_by_key:
            continue
        prev_item = previous_item_for_audit(key, curr, previous_items)
        if prev_item is None or bool(prev_item.get("low_stock")) or not bool(curr.get("low_stock")):
            continue
        pending_count = int(to_number(prev_item.get("pending_low_count"))) + 1
        if pending_count >= required_confirmations:
            curr.pop("pending_low_count", None)
            curr.pop("pending_stock_qty", None)
            curr.pop("pending_observed_at", None)
            continue
        preserved = dict(prev_item)
        preserved["pending_low_count"] = pending_count
        preserved["pending_stock_qty"] = curr.get("stock_qty")
        preserved["pending_observed_at"] = current.get("updated_at")
        guarded_items[key] = preserved
        pending_records.append({
            "key": key,
            "confirmation_count": pending_count,
            "required_confirmations": required_confirmations,
            "curr_stock_qty": curr.get("stock_qty"),
            "current_item": curr,
            "previous_item": prev_item,
        })

    guard_audit = {
        "run_blocked": bool(quality.get("global_reasons")),
        "global_reasons": quality.get("global_reasons", []),
        "blocked_item_count": len(blocked_records),
        "pending_confirmation_count": len(pending_records),
        "required_low_stock_confirmations": required_confirmations,
        "blocked_records": blocked_records,
        "pending_records": pending_records,
    }
    return guarded, guard_audit


def previous_item_for_audit(key: str, curr: dict, previous_items: dict) -> dict | None:
    prev = previous_items.get(key)
    if prev is not None:
        return prev
    legacy_keys = curr.get("legacy_keys") or []
    legacy_states = [previous_items[item_key] for item_key in legacy_keys if item_key in previous_items]
    if not legacy_states:
        return None
    low_stock_qty = float(env("LOW_STOCK_QTY", "10"))
    stock_qty = sum(to_number(item.get("stock_qty")) for item in legacy_states)
    return {
        "stock_qty": format_qty(stock_qty),
        "low_stock": stock_qty < low_stock_qty,
        "legacy_merged": True,
    }


def audit_reason(prev_low: bool | None, curr_low: bool, first_run: bool) -> str:
    if prev_low is None:
        if first_run:
            return "first_run_baseline"
        return "new_item_baseline"
    if not prev_low and curr_low:
        return "stockout"
    if prev_low and curr_low:
        return "still_low_no_push"
    if prev_low and not curr_low:
        return "restock"
    return "normal"


def should_write_audit(reason: str, prev_item: dict | None, curr: dict) -> bool:
    mode = env("DETAIL_LOG_MODE", "changed").lower()
    if mode == "all":
        return True
    if reason in {"stockout", "restock", "still_low_no_push"}:
        return True
    if curr.get("low_stock"):
        return True
    if prev_item is None:
        return False
    return to_number(prev_item.get("stock_qty")) != to_number(curr.get("stock_qty"))


def write_detail_audit(
    log_dir: Path,
    previous: dict | None,
    current: dict,
    guard_audit: dict | None = None,
    raw_current: dict | None = None,
) -> dict:
    log_dir.mkdir(parents=True, exist_ok=True)
    cleanup_old_logs(log_dir, int(env("LOG_RETENTION_DAYS", "3") or "3"))
    current_items = current.get("items", {})
    previous_items = {} if previous is None else previous.get("items", {})
    first_run = previous is None
    records: list[dict] = []
    run_id = now_cn().strftime("%Y%m%d%H%M%S")
    guard_audit = guard_audit or {}
    raw_current = raw_current or current
    special_keys = {
        record.get("key")
        for record in guard_audit.get("blocked_records", []) + guard_audit.get("pending_records", [])
        if record.get("key")
    }

    for key, curr in sorted(current_items.items()):
        if guard_audit.get("run_blocked") or key in special_keys:
            continue
        prev_item = previous_item_for_audit(key, curr, previous_items)
        prev_low = None if prev_item is None else bool(prev_item.get("low_stock"))
        curr_low = bool(curr.get("low_stock"))
        reason = audit_reason(prev_low, curr_low, first_run)
        if not should_write_audit(reason, prev_item, curr):
            continue
        records.append({
            "run_id": run_id,
            "run_time": current.get("updated_at"),
            "key": key,
            "reason": reason,
            "will_push": reason in {"stockout", "restock"},
            "prev_low_stock": prev_low,
            "curr_low_stock": curr_low,
            "prev_stock_qty": None if prev_item is None else prev_item.get("stock_qty"),
            "curr_stock_qty": curr.get("stock_qty"),
            "threshold": current.get("low_stock_qty"),
            "region": curr.get("region"),
            "sku_code": curr.get("sku_code"),
            "sku_codes": curr.get("sku_codes"),
            "sku_name": curr.get("sku_name"),
            "strategy": curr.get("strategy"),
            "product_type": curr.get("product_type"),
            "replacement_group": curr.get("replacement_group"),
            "warehouses": curr.get("warehouses"),
        })

    for blocked in guard_audit.get("blocked_records", []):
        curr = blocked.get("current_item") or {}
        prev_item = blocked.get("previous_item")
        records.append({
            "run_id": run_id,
            "run_time": raw_current.get("updated_at"),
            "key": blocked.get("key"),
            "reason": "data_missing_blocked",
            "data_quality_reason": blocked.get("reason"),
            "will_push": False,
            "prev_low_stock": None if prev_item is None else bool(prev_item.get("low_stock")),
            "curr_low_stock": bool(curr.get("low_stock")),
            "prev_stock_qty": blocked.get("prev_stock_qty"),
            "curr_stock_qty": blocked.get("curr_stock_qty"),
            "threshold": raw_current.get("low_stock_qty"),
            "region": curr.get("region"),
            "sku_code": curr.get("sku_code"),
            "sku_codes": curr.get("sku_codes"),
            "sku_name": curr.get("sku_name"),
            "warehouses": curr.get("warehouses"),
            "missing_source_keys": blocked.get("missing_source_keys", []),
        })

    for pending in guard_audit.get("pending_records", []):
        curr = pending.get("current_item") or {}
        prev_item = pending.get("previous_item")
        records.append({
            "run_id": run_id,
            "run_time": raw_current.get("updated_at"),
            "key": pending.get("key"),
            "reason": "low_stock_pending_confirmation",
            "will_push": False,
            "prev_low_stock": None if prev_item is None else bool(prev_item.get("low_stock")),
            "curr_low_stock": bool(curr.get("low_stock")),
            "prev_stock_qty": None if prev_item is None else prev_item.get("stock_qty"),
            "curr_stock_qty": pending.get("curr_stock_qty"),
            "threshold": raw_current.get("low_stock_qty"),
            "confirmation_count": pending.get("confirmation_count"),
            "required_confirmations": pending.get("required_confirmations"),
            "region": curr.get("region"),
            "sku_code": curr.get("sku_code"),
            "sku_codes": curr.get("sku_codes"),
            "sku_name": curr.get("sku_name"),
            "warehouses": curr.get("warehouses"),
        })

    if guard_audit.get("run_blocked"):
        records.append({
            "run_id": run_id,
            "run_time": raw_current.get("updated_at"),
            "reason": "data_quality_run_blocked",
            "will_push": False,
            "global_reasons": guard_audit.get("global_reasons", []),
            "message": "Doris returned incomplete inventory data; push and state update were skipped.",
        })

    detail_file = log_dir / f"inventory_alert_detail_{now_cn().strftime('%Y%m%d')}.jsonl"
    if records:
        with detail_file.open("a", encoding="utf-8") as handle:
            for record in records:
                handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
    return {
        "detail_log_file": str(detail_file),
        "detail_log_mode": env("DETAIL_LOG_MODE", "changed").lower(),
        "detail_record_count": len(records),
    }


def compact_warehouses(warehouses: list[str], limit: int = 8) -> str:
    values = [item for item in warehouses if item]
    if len(values) <= limit:
        return "、".join(values)
    return "、".join(values[:limit]) + f"等{len(values)}个仓"


def compact_sku_codes(row: dict, limit: int = 4) -> str:
    values = [text(item) for item in row.get("sku_codes", []) if text(item)]
    if not values:
        value = text(row.get("sku_code"))
        return value
    if len(values) <= limit:
        return "、".join(values)
    return "、".join(values[:limit]) + f"等{len(values)}个"


def require_pillow():
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError as exc:
        raise SystemExit("缺少依赖 Pillow，请执行：py -m pip install -r requirements.txt") from exc
    return Image, ImageDraw, ImageFont


def font_path() -> str:
    configured = env("IMAGE_FONT_PATH")
    if configured and Path(configured).exists():
        return configured
    for candidate in [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\Deng.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]:
        if Path(candidate).exists():
            return candidate
    return ""


def load_font(size: int, bold: bool = False):
    _image, _draw, ImageFont = require_pillow()
    path = font_path()
    if bold:
        for bold_candidate in [
            r"C:\Windows\Fonts\msyhbd.ttc",
            r"C:\Windows\Fonts\Dengb.ttf",
            r"C:\Windows\Fonts\simhei.ttf",
        ]:
            if Path(bold_candidate).exists():
                path = bold_candidate
                break
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def text_bbox(draw, value: str, font) -> tuple[int, int]:
    bbox = draw.textbbox((0, 0), value, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw, value: str, font, max_width: int, max_lines: int | None = None) -> list[str]:
    value = text(value)
    if not value:
        return [""]

    lines: list[str] = []
    for raw_part in re.split(r"[\r\n]+", value):
        part = raw_part.strip()
        if not part:
            continue
        current = ""
        for char in part:
            candidate = current + char
            if current and text_bbox(draw, candidate, font)[0] > max_width:
                lines.append(current)
                current = char
                if max_lines and len(lines) >= max_lines:
                    lines[-1] = lines[-1].rstrip("，、 ") + "..."
                    return lines
            else:
                current = candidate
        if current:
            lines.append(current)
            if max_lines and len(lines) >= max_lines:
                if len(current) < len(part):
                    lines[-1] = lines[-1].rstrip("，、 ") + "..."
                return lines
    return lines or [""]


def draw_round_rect(draw, xy: tuple[int, int, int, int], radius: int, fill: str, outline: str | None = None, width: int = 1) -> None:
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def chunk_rows_for_images(rows: list[dict]) -> list[list[dict]]:
    max_rows = int(env("IMAGE_MAX_ROWS_PER_CARD", "12") or "12")
    max_rows = max(4, min(max_rows, 24))
    return [rows[idx:idx + max_rows] for idx in range(0, len(rows), max_rows)]


def hex_to_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[idx:idx + 2], 16) for idx in (0, 2, 4))


def blend_color(left: str, right: str, ratio: float) -> tuple[int, int, int]:
    l = hex_to_rgb(left)
    r = hex_to_rgb(right)
    return tuple(int(l[idx] + (r[idx] - l[idx]) * ratio) for idx in range(3))


def fill_vertical_gradient(draw, xy: tuple[int, int, int, int], top_color: str, bottom_color: str) -> None:
    x1, y1, x2, y2 = xy
    height = max(1, y2 - y1)
    for offset in range(height):
        ratio = offset / height
        draw.line((x1, y1 + offset, x2, y1 + offset), fill=blend_color(top_color, bottom_color, ratio))


def render_alert_image(title: str, action: str, rows: list[dict], page_index: int, page_count: int) -> Path:
    Image, ImageDraw, _font = require_pillow()
    width = int(env("IMAGE_WIDTH", "1120") or "1120")
    width = max(900, min(width, 1600))
    margin = 32
    header_h = 170
    footer_h = 76
    row_gap = 14
    row_h = 138
    height = margin * 2 + header_h + footer_h + len(rows) * row_h + max(0, len(rows) - 1) * row_gap

    danger = action == "下架"
    if danger:
        palette = {
            "bg_top": "#FFF5F3",
            "bg_bottom": "#F8FAFC",
            "hero_left": "#B42318",
            "hero_right": "#F04438",
            "accent": "#D92D20",
            "accent_dark": "#912018",
            "accent_soft": "#FEF3F2",
            "accent_line": "#FDA29B",
            "chip_bg": "#FFFBFA",
            "chip_text": "#B42318",
        }
        title_text = "线上缺货提醒"
        subtitle = "库存跌破阈值，需要安排下架"
        tag_text = "安排下架"
        icon_text = "!"
    else:
        palette = {
            "bg_top": "#F0FDF4",
            "bg_bottom": "#F8FAFC",
            "hero_left": "#067647",
            "hero_right": "#12B76A",
            "accent": "#079455",
            "accent_dark": "#05603A",
            "accent_soft": "#ECFDF3",
            "accent_line": "#75E0A7",
            "chip_bg": "#F6FEF9",
            "chip_text": "#067647",
        }
        title_text = "线上到货提醒"
        subtitle = "库存恢复到阈值，可以安排上架"
        tag_text = "安排上架"
        icon_text = ""

    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    fill_vertical_gradient(draw, (0, 0, width, height), palette["bg_top"], palette["bg_bottom"])

    title_font = load_font(38, bold=True)
    subtitle_font = load_font(21)
    meta_font = load_font(18)
    tag_font = load_font(23, bold=True)
    row_title_font = load_font(24, bold=True)
    row_meta_font = load_font(18)
    small_font = load_font(16)
    index_font = load_font(22, bold=True)

    card = (margin, margin, width - margin, height - margin)
    draw_round_rect(draw, card, 24, "#FFFFFF", "#E4E7EC")

    hero = (margin, margin, width - margin, margin + header_h)
    draw.rounded_rectangle(hero, radius=24, fill=palette["hero_left"])
    fill_vertical_gradient(draw, hero, palette["hero_right"], palette["hero_left"])
    draw.rectangle((margin, margin + header_h - 28, width - margin, margin + header_h), fill=palette["hero_left"])

    header_x = margin + 32
    header_y = margin + 30
    icon_box = (header_x, header_y + 8, header_x + 64, header_y + 72)
    draw.ellipse(icon_box, fill="#FFFFFF")
    if danger:
        icon_w, icon_h = text_bbox(draw, icon_text, load_font(38, bold=True))
        draw.text((header_x + (64 - icon_w) / 2, header_y + 8 + (64 - icon_h) / 2 - 5), icon_text, fill=palette["accent"], font=load_font(38, bold=True))
    else:
        draw.line(
            (header_x + 18, header_y + 43, header_x + 29, header_y + 55, header_x + 48, header_y + 28),
            fill=palette["accent"],
            width=7,
            joint="curve",
        )

    draw.text((header_x + 84, header_y), title_text, fill="#FFFFFF", font=title_font)
    draw.text((header_x + 86, header_y + 52), subtitle, fill="#FFEDEB" if danger else "#DCFAE6", font=subtitle_font)
    timestamp = now_cn().strftime("%Y-%m-%d %H:%M")
    page_suffix = f" · 第{page_index}/{page_count}张" if page_count > 1 else ""
    draw.text((header_x + 86, header_y + 86), f"{timestamp}{page_suffix} · 共 {len(rows)} 条", fill="#FFFFFF", font=meta_font)

    tag_w = 146
    tag_h = 50
    tag_x = width - margin - 34 - tag_w
    tag_y = header_y + 24
    draw_round_rect(draw, (tag_x, tag_y, tag_x + tag_w, tag_y + tag_h), 25, "#FFFFFF")
    tag_text_w, tag_text_h = text_bbox(draw, tag_text, tag_font)
    draw.text((tag_x + (tag_w - tag_text_w) / 2, tag_y + (tag_h - tag_text_h) / 2 - 2), tag_text, fill=palette["accent_dark"], font=tag_font)

    x = margin + 28
    y = margin + header_h + 24
    row_width = width - margin * 2 - 56
    for idx, row in enumerate(rows, start=1):
        row_box = (x, y, x + row_width, y + row_h)
        draw_round_rect(draw, row_box, 16, "#FFFFFF", "#E6EAF0")
        draw_round_rect(draw, (x + 14, y + 20, x + 64, y + 70), 25, palette["accent_soft"], palette["accent_line"])
        index_text = f"{idx:02d}"
        idx_w, idx_h = text_bbox(draw, index_text, index_font)
        draw.text((x + 14 + (50 - idx_w) / 2, y + 20 + (50 - idx_h) / 2 - 2), index_text, fill=palette["accent_dark"], font=index_font)

        region = text(row.get("region"))
        sku_name = text(row.get("sku_name"))
        title_lines = wrap_text(draw, f"{region}｜{sku_name}", row_title_font, row_width - 112, max_lines=2)
        title_y = y + 18
        for line in title_lines:
            draw.text((x + 82, title_y), line, fill="#111827", font=row_title_font)
            title_y += 31

        meta_y = y + 82
        warehouse_text = compact_warehouses(row.get("warehouses", []), limit=10)
        warehouse_lines = wrap_text(draw, f"覆盖仓：{warehouse_text}", row_meta_font, row_width - 112, max_lines=1)
        draw.text((x + 82, meta_y), warehouse_lines[0], fill="#475467", font=row_meta_font)

        if row.get("replacement_group"):
            sku_text = f"替代组SKU：{compact_sku_codes(row, limit=6)}"
            sku_lines = wrap_text(draw, sku_text, small_font, row_width - 112, max_lines=1)
            draw.text((x + 82, y + row_h - 30), sku_lines[0], fill="#667085", font=small_font)

        draw_round_rect(draw, (x + row_width - 104, y + row_h - 42, x + row_width - 22, y + row_h - 14), 14, palette["chip_bg"], palette["accent_line"])
        chip_w, chip_h = text_bbox(draw, tag_text, small_font)
        draw.text((x + row_width - 104 + (82 - chip_w) / 2, y + row_h - 42 + (28 - chip_h) / 2 - 1), tag_text, fill=palette["chip_text"], font=small_font)
        y += row_h + row_gap

    footer_text = "规则：区域覆盖仓总库存 < 10 推缺货，恢复到 >= 10 推上架；已低库存不重复提醒。"
    footer_y = height - margin - 48
    draw_round_rect(draw, (margin + 28, footer_y - 10, width - margin - 28, footer_y + 34), 18, "#F8FAFC", "#EAECF0")
    draw.text((margin + 48, footer_y), footer_text, fill="#667085", font=small_font)

    image_dir = resolve_path(env("IMAGE_OUTPUT_DIR", "data/alert_images"), ROOT / "data" / "alert_images")
    image_dir.mkdir(parents=True, exist_ok=True)
    filename = f"alert_{now_cn().strftime('%Y%m%d_%H%M%S')}_{'stockout' if danger else 'restock'}_{page_index}.png"
    path = image_dir / filename
    image.save(path, format="PNG", optimize=True)
    return path


def build_image_messages(title: str, action: str, rows: list[dict]) -> list[Path]:
    if not rows:
        return []
    chunks = chunk_rows_for_images(rows)
    return [render_alert_image(title, action, chunk, idx, len(chunks)) for idx, chunk in enumerate(chunks, start=1)]


def build_message(title: str, action: str, rows: list[dict]) -> list[str]:
    if not rows:
        return []

    timestamp = now_cn().strftime("%m/%d %H:%M")
    chunks: list[list[dict]] = []
    current: list[dict] = []
    current_len = 0
    for row in rows:
        line_len = len(row.get("region", "")) + len(row.get("sku_name", "")) + len("".join(row.get("warehouses", []))) + 40
        if current and current_len + line_len > 2600:
            chunks.append(current)
            current = []
            current_len = 0
        current.append(row)
        current_len += line_len
    if current:
        chunks.append(current)

    messages = []
    for idx, chunk in enumerate(chunks, start=1):
        suffix = f"（第{idx}/{len(chunks)}批）" if len(chunks) > 1 else ""
        lines = [f"{title} {timestamp}{suffix}", "----------------------------"]
        for row in chunk:
            lines.append(f"· {row['region']} - {row['sku_name']}")
            if row.get("replacement_group"):
                lines.append(f"  替代组SKU：{compact_sku_codes(row)}")
            lines.append(f"  仓库：{compact_warehouses(row.get('warehouses', []))}")
        lines.extend([
            "----------------------------",
            f"共 {len(rows)} 条，麻烦安排{action}",
        ])
        messages.append("\n".join(lines))
    return messages


def send_wecom_text(webhook_url: str, content: str) -> None:
    payload = json.dumps({"msgtype": "text", "text": {"content": content}}, ensure_ascii=False).encode("utf-8")
    send_wecom_payload(webhook_url, payload)


def send_wecom_image(webhook_url: str, image_path: Path) -> None:
    data = image_path.read_bytes()
    payload = json.dumps({
        "msgtype": "image",
        "image": {
            "base64": base64.b64encode(data).decode("ascii"),
            "md5": hashlib.md5(data).hexdigest(),
        },
    }, ensure_ascii=False).encode("utf-8")
    send_wecom_payload(webhook_url, payload)


def send_wecom_payload(webhook_url: str, payload: bytes) -> None:
    request = urllib.request.Request(
        webhook_url,
        data=payload,
        headers={"Content-Type": "application/json; charset=utf-8"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=20) as response:
        body = response.read().decode("utf-8", errors="replace")
    try:
        result = json.loads(body)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"企微 Webhook 返回非 JSON：{body}") from exc
    if result.get("errcode") != 0:
        raise RuntimeError(f"企微 Webhook 推送失败：{body}")


def push_alerts(stockout_alerts: list[dict], restock_alerts: list[dict], no_push: bool = False) -> dict:
    enabled = parse_bool(env("WECOM_ENABLED", "true"), default=True)
    webhook_url = env("WECOM_WEBHOOK_URL")
    dry_run = no_push or parse_bool(env("DRY_RUN", "false"))
    push_format = env("PUSH_FORMAT", "image").lower()

    if push_format == "text":
        messages: Sequence[str | Path] = [
            *build_message("[!] 【缺货提醒】", "下架", stockout_alerts),
            *build_message("[+] 【到货提醒】", "上架", restock_alerts),
        ]
    else:
        messages = [
            *build_image_messages("线上缺货提醒", "下架", stockout_alerts),
            *build_image_messages("线上到货提醒", "上架", restock_alerts),
        ]

    result = {
        "enabled": enabled,
        "dry_run": dry_run,
        "format": "text" if push_format == "text" else "image",
        "message_count": len(messages),
        "sent_count": 0,
        "image_files": [str(message) for message in messages if isinstance(message, Path)],
    }

    if not messages:
        return result
    if dry_run:
        for message in messages:
            if isinstance(message, Path):
                logging.info("DRY_RUN image: %s", message)
            else:
                logging.info("DRY_RUN message:\n%s", message)
        return result
    if not enabled:
        logging.info("WECOM_ENABLED=false，跳过企微推送")
        return result
    if not webhook_url:
        raise ValueError("已启用企微推送，但 WECOM_WEBHOOK_URL 为空")

    for message in messages:
        if isinstance(message, Path):
            send_wecom_image(webhook_url, message)
        else:
            send_wecom_text(webhook_url, message)
        result["sent_count"] += 1
    return result


def run_once(args: argparse.Namespace) -> dict:
    base_excel = resolve_path(env("BASE_EXCEL_PATH", "config/基础数据线上.xlsx"), ROOT / "config" / "基础数据线上.xlsx")
    state_file = resolve_path(env("STATE_FILE", "data/last_status_region.json"), DEFAULT_STATE)
    summary_file = resolve_path(env("SUMMARY_FILE", "data/last_run_summary.json"), DEFAULT_SUMMARY)
    log_dir = resolve_path(env("LOG_DIR", "logs"), DEFAULT_LOG_DIR)

    if args.base_excel:
        base_excel = Path(args.base_excel)
        if not base_excel.is_absolute():
            base_excel = (Path.cwd() / base_excel).resolve()

    products, warehouses, base_audit = load_base_excel(base_excel)
    logging.info("基础数据读取完成：products=%s warehouses=%s", len(products), len(warehouses))
    previous = load_state(state_file)
    fetch_attempts = max(1, int(env("DORIS_DATA_RETRY_COUNT", "3") or "3"))
    retry_delay_seconds = max(0.0, float(env("DORIS_DATA_RETRY_DELAY_SECONDS", "10") or "10"))
    attempt_audits: list[dict] = []
    for attempt in range(1, fetch_attempts + 1):
        try:
            stock_rows, sales_rows = fetch_doris_data()
        except Exception as exc:
            attempt_audits.append({
                "attempt": attempt,
                "fetch_error": f"{type(exc).__name__}: {exc}",
                "retry_required": True,
            })
            if attempt >= fetch_attempts:
                logging.exception("Doris 连接或查询连续 %s 次失败，本次终止", fetch_attempts)
                raise
            logging.warning(
                "Doris 连接或查询失败：attempt=%s/%s error=%s，%s 秒后重试",
                attempt,
                fetch_attempts,
                exc,
                retry_delay_seconds,
            )
            time.sleep(retry_delay_seconds)
            continue
        stock_map, sales_map, doris_audit = build_quantity_maps(stock_rows, sales_rows)
        current, aggregate_audit = build_current_snapshot(products, warehouses, stock_map, sales_map)
        quality_audit = inspect_snapshot_quality(previous, current)
        retry_required = bool(
            quality_audit.get("global_reasons")
            or quality_audit.get("item_issues")
            or quality_audit.get("suspicious_low_keys")
        )
        attempt_audits.append({
            "attempt": attempt,
            "stock_row_count": len(stock_rows),
            "stock_key_count": doris_audit.get("stock_key_count"),
            "retry_required": retry_required,
            "global_reasons": quality_audit.get("global_reasons", []),
            "item_issue_count": len(quality_audit.get("item_issues", [])),
            "suspicious_low_count": len(quality_audit.get("suspicious_low_keys", [])),
        })
        logging.info(
            "Doris 取数完成：attempt=%s/%s stock_rows=%s stock_keys=%s suspicious_low=%s item_issues=%s global_issues=%s",
            attempt,
            fetch_attempts,
            len(stock_rows),
            doris_audit.get("stock_key_count"),
            len(quality_audit.get("suspicious_low_keys", [])),
            len(quality_audit.get("item_issues", [])),
            quality_audit.get("global_reasons", []),
        )
        if not retry_required or attempt >= fetch_attempts:
            break
        logging.warning("Doris 数据可能不完整或库存突降，%s 秒后重试", retry_delay_seconds)
        time.sleep(retry_delay_seconds)

    guarded_current, guard_audit = apply_data_quality_guard(previous, current, quality_audit)
    guard_audit["fetch_attempts"] = attempt_audits
    if guard_audit.get("run_blocked"):
        stockout_alerts, restock_alerts, skipped_restock = [], [], []
        logging.error("Doris 整批数据异常，本次禁止推送且不更新状态：%s", guard_audit.get("global_reasons"))
    else:
        stockout_alerts, restock_alerts, skipped_restock = compare_snapshots(previous, guarded_current)
    detail_audit = write_detail_audit(
        log_dir,
        previous,
        guarded_current,
        guard_audit=guard_audit,
        raw_current=current,
    )

    push_result = push_alerts(
        stockout_alerts,
        restock_alerts,
        no_push=args.no_push or args.dry_run or guard_audit.get("run_blocked"),
    )
    if args.dry_run:
        dry_run_snapshot = state_file.with_name("dry_run_last_status_region.json")
        atomic_write_json(dry_run_snapshot, current)
        state_write = {
            "updated": False,
            "reason": "dry-run 不更新正式状态快照",
            "dry_run_snapshot": str(dry_run_snapshot),
        }
    elif guard_audit.get("run_blocked"):
        state_write = {
            "updated": False,
            "reason": "Doris 数据完整性校验失败，本次不更新状态快照",
            "dry_run_snapshot": "",
        }
    else:
        atomic_write_json(state_file, guarded_current)
        state_write = {
            "updated": True,
            "reason": "",
            "dry_run_snapshot": "",
        }

    summary = {
        "ok": True,
        "program_version": VERSION,
        "updated_at": current["updated_at"],
        "state_file": str(state_file),
        "base_audit": base_audit,
        "doris_audit": doris_audit,
        "aggregate_audit": aggregate_audit,
        "data_quality_guard": guard_audit,
        "detail_audit": detail_audit,
        "alert_counts": {
            "stockout": len(stockout_alerts),
            "restock": len(restock_alerts),
            "skipped_restock": len(skipped_restock),
            "first_run": previous is None,
        },
        "state_write": state_write,
        "push_result": push_result,
        "skipped_restock_examples": skipped_restock[:20],
    }
    atomic_write_json(summary_file, summary)
    logging.info(
        "运行完成：stockout=%s restock=%s skipped_restock=%s detail_records=%s sent=%s",
        len(stockout_alerts),
        len(restock_alerts),
        len(skipped_restock),
        detail_audit.get("detail_record_count"),
        push_result.get("sent_count"),
    )
    return summary


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="线上缺货/到货预警：Doris 取数并推送企业微信")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="配置文件路径，默认 config/config.env")
    parser.add_argument("--base-excel", type=Path, default=None, help="临时指定基础数据 Excel")
    parser.add_argument("--dry-run", action="store_true", help="取数和生成结果，但不推送")
    parser.add_argument("--no-push", action="store_true", help="不推送企微，但正常更新状态快照")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    load_env_file(args.config)
    log_dir = resolve_path(env("LOG_DIR", "logs"), DEFAULT_LOG_DIR)
    setup_logging(log_dir)
    logging.info("inventory_alert version=%s root=%s", VERSION, ROOT)
    try:
        summary = run_once(args)
    except Exception:
        logging.exception("运行失败")
        return 1

    print(json.dumps({
        "ok": True,
        "updated_at": summary["updated_at"],
        "alert_counts": summary["alert_counts"],
        "push_result": summary["push_result"],
        "state_file": summary["state_file"],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
