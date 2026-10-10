from __future__ import annotations

import json
import os
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pymysql
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PROJECT_ROOT = Path(__file__).resolve().parents[1]
ENV_FILE = ROOT / "config" / "pickup.env"
PRODUCT_FILE = ROOT / "online_inventory_alert" / "config" / "基础数据线上.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "scope_inspection.json"


def load_env_file(path: Path) -> None:
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip())


def connect() -> pymysql.connections.Connection:
    return pymysql.connect(
        host=os.environ["PICKUP_DB_HOST"],
        port=int(os.environ.get("PICKUP_DB_PORT", "9030")),
        user=os.environ["PICKUP_DB_USER"],
        password=os.environ["PICKUP_DB_PASSWORD"],
        database=os.environ.get("PICKUP_DB_NAME", "ywzt_ods"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=15,
        read_timeout=120,
        write_timeout=30,
    )


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_candidate_products() -> tuple[list[dict[str, str]], dict[str, Any]]:
    wb = load_workbook(PRODUCT_FILE, read_only=True, data_only=True)
    ws = wb["产品信息表"]
    rows = ws.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    index = {name: idx for idx, name in enumerate(headers) if name}

    products: dict[str, dict[str, str]] = {}
    duplicate_codes: list[str] = []
    skipped_strategy = Counter()
    skipped_category = Counter()

    for row in rows:
        sku_code = text(row[index["产品编码"]] if index["产品编码"] < len(row) else "")
        if not sku_code:
            continue
        sku_name = text(row[index["商品名称"]] if index["商品名称"] < len(row) else "") or sku_code
        strategy = text(row[index["布仓策略"]] if index["布仓策略"] < len(row) else "")
        category = text(row[index["产品大类"]] if index.get("产品大类") is not None else "")

        if category != "产成品":
            skipped_category[category or "<空>"] += 1
            continue
        if strategy == "全铺":
            product_type = "frozen"
        elif "常温" in strategy:
            product_type = "ambient"
        else:
            skipped_strategy[strategy or "<空>"] += 1
            continue

        if sku_code in products:
            duplicate_codes.append(sku_code)
            continue
        products[sku_code] = {
            "sku_code": sku_code,
            "sku_name": sku_name,
            "strategy": strategy,
            "product_type": product_type,
        }

    return list(products.values()), {
        "source_file": str(PRODUCT_FILE),
        "candidate_product_count": len(products),
        "duplicate_codes": duplicate_codes,
        "skipped_strategy": dict(skipped_strategy),
        "skipped_category": dict(skipped_category),
    }


def load_candidate_warehouses() -> list[dict[str, str]]:
    wb = load_workbook(PRODUCT_FILE, read_only=True, data_only=True)
    ws = wb["仓库映射表"]
    rows = ws.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    index = {name: idx for idx, name in enumerate(headers) if name}
    warehouses: list[dict[str, str]] = []
    for row in rows:
        wms_desc = text(row[index["wms_desc"]] if index["wms_desc"] < len(row) else "")
        if not wms_desc:
            continue
        warehouses.append({
            "wms_desc": wms_desc,
            "short_name": text(row[index["仓库简称"]] if index["仓库简称"] < len(row) else ""),
            "service_provinces": text(row[index["可发省份"]] if index["可发省份"] < len(row) else ""),
        })
    return warehouses


def json_default(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def fetch_all(cursor: pymysql.cursors.DictCursor, sql: str, params: Any = None) -> list[dict[str, Any]]:
    cursor.execute(sql, params)
    return list(cursor.fetchall())


def fetch_one(cursor: pymysql.cursors.DictCursor, sql: str, params: Any = None) -> dict[str, Any]:
    rows = fetch_all(cursor, sql, params)
    return rows[0] if rows else {}


def main() -> None:
    load_env_file(ENV_FILE)
    products, workbook_audit = load_candidate_products()
    warehouses = load_candidate_warehouses()
    sku_codes = [product["sku_code"] for product in products]
    placeholders = ",".join(["%s"] * len(sku_codes))
    warehouse_names = [warehouse["wms_desc"] for warehouse in warehouses]
    warehouse_placeholders = ",".join(["%s"] * len(warehouse_names))

    result: dict[str, Any] = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "workbook_scope": workbook_audit,
        "tables": {},
        "scope_validation": {},
        "date_coverage": {},
        "samples": {},
    }

    with connect() as conn, conn.cursor() as cursor:
        for table in [
            "master_sku",
            "master_wms",
            "scm_sales_shipment_header",
            "scm_sales_shipment_detail",
            "scm_po_header",
            "scm_po_detail",
            "scm_warehous_manage",
            "scm_warehous_manage_backups",
        ]:
            columns = fetch_all(cursor, f"SHOW COLUMNS FROM {table}")
            result["tables"][table] = [row.get("Field") for row in columns]

        master_rows = fetch_all(
            cursor,
            f"""
            SELECT sku_code,
                   sku_name,
                   sku_group,
                   finished_product_category,
                   sku_class,
                   group_number,
                   net_weight,
                   use_status,
                   status,
                   controlsale,
                   sku_cycle_status,
                   sku_sale_status,
                   sku_launch_date,
                   sku_delisting_date
            FROM master_sku
            WHERE sku_code IN ({placeholders})
            """,
            sku_codes,
        )
        matched_codes = {text(row.get("sku_code")) for row in master_rows}
        missing_codes = [code for code in sku_codes if code not in matched_codes]
        category_counts = Counter(
            text(row.get("finished_product_category"))
            or text(row.get("sku_class"))
            or text(row.get("group_number"))
            or "未分类"
            for row in master_rows
        )
        group_counts = Counter(text(row.get("sku_group")) or "<空>" for row in master_rows)
        result["scope_validation"] = {
            "candidate_sku_count": len(sku_codes),
            "master_sku_matched_count": len(matched_codes),
            "master_sku_missing_count": len(missing_codes),
            "master_sku_missing_codes": missing_codes,
            "sku_group_counts": dict(group_counts),
            "category_counts": dict(category_counts),
            "use_status_counts": dict(Counter(text(row.get("use_status")) or "<空>" for row in master_rows)),
            "status_counts": dict(Counter(text(row.get("status")) or "<空>" for row in master_rows)),
            "controlsale_counts": dict(Counter(text(row.get("controlsale")) or "<空>" for row in master_rows)),
            "sku_cycle_status_counts": dict(Counter(text(row.get("sku_cycle_status")) or "<空>" for row in master_rows)),
            "sku_sale_status_counts": dict(Counter(text(row.get("sku_sale_status")) or "<空>" for row in master_rows)),
        }

        category_sample_map: dict[str, dict[str, Any]] = {}
        for row in master_rows:
            category_code = text(row.get("finished_product_category")) or "未分类"
            item = category_sample_map.setdefault(
                category_code,
                {"finished_product_category": category_code, "sku_count": 0, "sku_examples": []},
            )
            item["sku_count"] += 1
            if len(item["sku_examples"]) < 6:
                item["sku_examples"].append(f"{text(row.get('sku_code'))}:{text(row.get('sku_name'))}")
        category_samples = []
        for category_code in sorted(category_sample_map):
            item = category_sample_map[category_code]
            item["sku_examples"] = " | ".join(item["sku_examples"])
            category_samples.append(item)
        result["scope_validation"]["category_samples"] = category_samples

        warehouse_rows = fetch_all(
            cursor,
            f"""
            SELECT wms_code, wms_desc, wms_type, province, city, wms_status, status
            FROM master_wms
            WHERE wms_desc IN ({warehouse_placeholders})
            """,
            warehouse_names,
        )
        matched_warehouse_names = {text(row.get("wms_desc")) for row in warehouse_rows}
        result["scope_validation"]["warehouse_scope"] = {
            "candidate_warehouse_count": len(warehouse_names),
            "matched_warehouse_count": len(matched_warehouse_names),
            "missing_warehouse_names": [name for name in warehouse_names if name not in matched_warehouse_names],
            "province_counts": dict(Counter(text(row.get("province")) or "<空>" for row in warehouse_rows)),
            "rows": warehouse_rows,
        }

        result["scope_validation"]["sales_amount_quality"] = fetch_one(
            cursor,
            f"""
            SELECT COUNT(*) AS detail_rows,
                   SUM(CASE WHEN IFNULL(d.total_pay_amount, 0) <> 0 THEN 1 ELSE 0 END) AS total_pay_nonzero_rows,
                   SUM(CASE WHEN IFNULL(d.item_total_amount, 0) <> 0 THEN 1 ELSE 0 END) AS item_amount_nonzero_rows,
                   SUM(CASE WHEN IFNULL(d.list_price, 0) <> 0 THEN 1 ELSE 0 END) AS list_price_nonzero_rows,
                   SUM(IFNULL(d.total_pay_amount, 0)) AS total_pay_amount,
                   SUM(IFNULL(d.item_total_amount, 0)) AS item_total_amount,
                   SUM(IFNULL(d.list_price, 0) * IFNULL(d.ship_qty, 0)) AS list_price_amount
            FROM scm_sales_shipment_header h
            INNER JOIN scm_sales_shipment_detail d ON d.outbound_code = h.outbound_code
            WHERE h.outbound_code LIKE 'HJXH%%'
              AND h.ship_at >= DATE_SUB(CURDATE(), INTERVAL 30 DAY)
              AND h.ship_at < CURDATE() + INTERVAL 1 DAY
              AND h.status IN ('700','900','701','901')
              AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
              AND d.sku_code IN ({placeholders})
            """,
            sku_codes,
        )
        result["scope_validation"]["purchase_price_quality"] = fetch_one(
            cursor,
            f"""
            SELECT COUNT(*) AS detail_rows,
                   SUM(CASE WHEN IFNULL(d.actual_purchase_price, 0) > 0 THEN 1 ELSE 0 END) AS actual_price_nonzero_rows,
                   SUM(CASE WHEN IFNULL(d.standard_purchase_price, 0) > 0 THEN 1 ELSE 0 END) AS standard_price_nonzero_rows,
                   MIN(NULLIF(d.actual_purchase_price, 0)) AS min_actual_price,
                   MAX(d.actual_purchase_price) AS max_actual_price,
                   SUM(IFNULL(d.received_quantity, 0) * IFNULL(d.actual_purchase_price, 0)) AS received_actual_amount,
                   SUM(IFNULL(d.received_quantity, 0) * IFNULL(d.standard_purchase_price, 0)) AS received_standard_amount
            FROM scm_po_header h
            INNER JOIN scm_po_detail d ON d.po_code = h.po_code
            WHERE h.cargo_owner_id = 'HJXH'
              AND h.po_type = 'CCP'
              AND d.created_date >= DATE_SUB(CURDATE(), INTERVAL 90 DAY)
              AND d.sku_code IN ({placeholders})
            """,
            sku_codes,
        )

        result["date_coverage"]["inventory_live"] = fetch_one(
            cursor,
            """
            SELECT MAX(created_date) AS max_created_date,
                   COUNT(*) AS row_count,
                   COUNT(DISTINCT sku_code) AS sku_count,
                   COUNT(DISTINCT wms_code) AS warehouse_count
            FROM scm_warehous_manage
            WHERE cargo_owner_id = 'HJXH'
            """,
        )
        result["date_coverage"]["inventory_history"] = fetch_one(
            cursor,
            """
            SELECT MIN(DATE(backups_date)) AS min_date,
                   MAX(DATE(backups_date)) AS max_date,
                   COUNT(DISTINCT DATE(backups_date)) AS day_count
            FROM scm_warehous_manage_backups
            WHERE cargo_owner_id = 'HJXH'
            """,
        )
        result["date_coverage"]["sales"] = fetch_one(
            cursor,
            """
            SELECT MIN(DATE(ship_at)) AS min_date,
                   MAX(DATE(ship_at)) AS max_date,
                   COUNT(*) AS row_count
            FROM scm_sales_shipment_header
            WHERE outbound_code LIKE 'HJXH%%'
              AND ship_at >= DATE_SUB(CURDATE(), INTERVAL 7 MONTH)
              AND status IN ('700','900','701','901')
            """,
        )
        result["date_coverage"]["purchase"] = fetch_one(
            cursor,
            """
            SELECT MIN(DATE(d.created_date)) AS min_date,
                   MAX(DATE(d.created_date)) AS max_date,
                   COUNT(*) AS row_count
            FROM scm_po_header h
            INNER JOIN scm_po_detail d ON d.po_code = h.po_code
            WHERE h.cargo_owner_id = 'HJXH'
              AND h.po_type = 'CCP'
              AND d.created_date >= DATE_SUB(CURDATE(), INTERVAL 7 MONTH)
            """,
        )

        result["scope_validation"]["recent_sales"] = fetch_one(
            cursor,
            f"""
            SELECT COUNT(DISTINCT d.sku_code) AS sold_sku_count,
                   COUNT(DISTINCT h.outbound_code) AS outbound_order_count,
                   SUM(IFNULL(d.ship_qty, 0)) AS ship_qty,
                   SUM(IFNULL(d.ship_qty, 0) * IFNULL(ms.net_weight, 0) / 1000) AS ship_ton
            FROM scm_sales_shipment_header h
            INNER JOIN scm_sales_shipment_detail d ON d.outbound_code = h.outbound_code
            LEFT JOIN master_sku ms ON ms.sku_code = d.sku_code
            WHERE h.outbound_code LIKE 'HJXH%%'
              AND h.ship_at >= DATE_SUB(CURDATE(), INTERVAL 6 MONTH)
              AND h.status IN ('700','900','701','901')
              AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
              AND d.sku_code IN ({placeholders})
            """,
            sku_codes,
        )
        result["scope_validation"]["current_inventory"] = fetch_one(
            cursor,
            f"""
            SELECT COUNT(*) AS inventory_row_count,
                   COUNT(DISTINCT w.sku_code) AS stocked_sku_count,
                   COUNT(DISTINCT w.wms_code) AS warehouse_count,
                   SUM(IFNULL(w.in_stock_quantity, 0)) AS in_stock_qty,
                   SUM(IFNULL(w.in_stock_quantity, 0) * IFNULL(ms.net_weight, 0) / 1000) AS stock_ton
            FROM scm_warehous_manage w
            LEFT JOIN master_sku ms ON ms.sku_code = w.sku_code
            LEFT JOIN master_wms mw ON mw.wms_code = w.wms_code
            WHERE w.cargo_owner_id = 'HJXH'
              AND mw.wms_type IN (10, 11)
              AND w.sku_code IN ({placeholders})
            """,
            sku_codes,
        )
        result["scope_validation"]["recent_purchase"] = fetch_one(
            cursor,
            f"""
            SELECT COUNT(DISTINCT d.sku_code) AS purchased_sku_count,
                   COUNT(DISTINCT h.po_code) AS purchase_order_count,
                   SUM(IFNULL(d.received_quantity, 0)) AS received_qty,
                   SUM(IFNULL(d.received_quantity, 0) * IFNULL(ms.net_weight, 0) / 1000) AS received_ton
            FROM scm_po_header h
            INNER JOIN scm_po_detail d ON d.po_code = h.po_code
            LEFT JOIN master_sku ms ON ms.sku_code = d.sku_code
            WHERE h.cargo_owner_id = 'HJXH'
              AND h.po_type = 'CCP'
              AND h.status = 0
              AND d.status = 0
              AND d.created_date >= DATE_SUB(CURDATE(), INTERVAL 6 MONTH)
              AND d.sku_code IN ({placeholders})
            """,
            sku_codes,
        )

        result["samples"]["master_sku"] = master_rows[:15]
        result["samples"]["warehouse"] = fetch_all(
            cursor,
            """
            SELECT w.wms_code,
                   MAX(mw.wms_desc) AS wms_desc,
                   MAX(mw.wms_type) AS wms_type,
                   COUNT(DISTINCT w.sku_code) AS sku_count,
                   SUM(IFNULL(w.in_stock_quantity, 0)) AS in_stock_qty
            FROM scm_warehous_manage w
            LEFT JOIN master_wms mw ON mw.wms_code = w.wms_code
            WHERE w.cargo_owner_id = 'HJXH'
              AND mw.wms_type IN (10, 11)
            GROUP BY w.wms_code
            ORDER BY in_stock_qty DESC
            LIMIT 20
            """,
        )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )
    print(json.dumps({
        "ok": True,
        "output": str(OUTPUT_FILE),
        "candidate_sku_count": result["scope_validation"]["candidate_sku_count"],
        "matched_count": result["scope_validation"]["master_sku_matched_count"],
        "missing_count": result["scope_validation"]["master_sku_missing_count"],
        "recent_sales": result["scope_validation"]["recent_sales"],
        "current_inventory": result["scope_validation"]["current_inventory"],
        "recent_purchase": result["scope_validation"]["recent_purchase"],
        "date_coverage": result["date_coverage"],
    }, ensure_ascii=False, default=json_default))


if __name__ == "__main__":
    main()
