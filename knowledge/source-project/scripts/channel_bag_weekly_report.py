from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import mimetypes
import os
import sys
import time
import traceback
import urllib.parse
import urllib.request
import uuid
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import pymysql
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from xlsxwriter import Workbook as XlsxWorkbook


APP_NAME = "渠道袋数周报"
TASK_NAME = "Channel_Bag_Weekly_Report_Friday_1200"
TARGET_CATEGORIES = ("肉制品", "面点", "烘焙")
DETAIL_COLUMNS = [
    "发货日期", "仓库编码", "仓库描述", "商品编码", "商品描述", "总发货量",
    "线上汇总", "抖音", "抖音小时达", "拼多多系", "快手系", "淘宝系", "京东系",
    "唯品会", "线上其它", "线下汇总", "线下社区团购_平台", "线下社区团购_分销",
    "线下其它", "其它", "周报分类", "分类来源", "是否纳入汇总",
]


SQL = """
SELECT
    DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d') AS 发货日期,
    h.facility_code AS 仓库编码,
    w.wms_desc AS 仓库描述,
    i.sku_code AS 商品编码,
    m.sku_name AS 商品描述,
    SUM(i.ship_qty) AS 总发货量,
    SUM(CASE WHEN s.type = 'Z004' THEN i.ship_qty ELSE 0 END) AS 线上汇总,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'DY' THEN i.ship_qty ELSE 0 END) AS 抖音,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'DYXD' THEN i.ship_qty ELSE 0 END) AS 抖音小时达,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('PDD','PDDDF') THEN i.ship_qty ELSE 0 END) AS 拼多多系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('KS','KSDF') THEN i.ship_qty ELSE 0 END) AS 快手系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('TGC','MJYX','TB','CT','MAOCHAO') THEN i.ship_qty ELSE 0 END) AS 淘宝系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('JDVMI','JD','JD_MARKET') THEN i.ship_qty ELSE 0 END) AS 京东系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'WPH' THEN i.ship_qty ELSE 0 END) AS 唯品会,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code NOT IN ('DY','DYXD','PDD','PDDDF','KS','KSDF','TGC','MJYX','TB','CT','MAOCHAO','JDVMI','JD','JD_MARKET','WPH') THEN i.ship_qty ELSE 0 END) AS 线上其它,
    SUM(CASE WHEN s.type = 'Z001' THEN i.ship_qty ELSE 0 END) AS 线下汇总,
    SUM(CASE WHEN s.type = 'Z001' AND s.group1 = '03' THEN i.ship_qty ELSE 0 END) AS 线下社区团购_平台,
    SUM(CASE WHEN s.type = 'Z001' AND s.price_group = '122' THEN i.ship_qty ELSE 0 END) AS 线下社区团购_分销,
    SUM(CASE WHEN s.type = 'Z001' AND NOT (s.group1 = '03' OR s.price_group = '122') THEN i.ship_qty ELSE 0 END) AS 线下其它,
    SUM(CASE WHEN s.type NOT IN ('Z001','Z004') OR s.type IS NULL THEN i.ship_qty ELSE 0 END) AS 其它
FROM scm_sales_shipment_header h
INNER JOIN scm_sales_order_header h1 ON h1.sales_order_code = h.ref_order_code
INNER JOIN scm_sales_shipment_detail i ON i.outbound_code = h.outbound_code
LEFT JOIN master_store s ON s.code = h.store_code
LEFT JOIN master_sku m ON m.sku_code = i.sku_code
LEFT JOIN master_wms w ON w.wms_code = h.facility_code
WHERE h.ship_at >= %s AND h.ship_at < %s
  AND h.outbound_code LIKE 'HJXH%%'
  AND h.status IN ('700','900','701','901')
  AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
GROUP BY DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d'), h.facility_code, i.sku_code, m.sku_name, w.wms_desc
ORDER BY DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d'), h.facility_code, i.sku_code
"""


@dataclass(frozen=True)
class WeekWindow:
    name: str
    start: date
    end: date

    @property
    def start_dt(self) -> datetime:
        return datetime.combine(self.start, dt_time.min)

    @property
    def end_exclusive(self) -> datetime:
        return datetime.combine(self.end + timedelta(days=1), dt_time.min)


def app_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    script_path = Path(__file__).resolve()
    return script_path.parents[1] if script_path.parent.name.lower() == "scripts" else script_path.parent


ROOT = app_root()
CONFIG_FILE = ROOT / "config" / "config.env"
DEFAULT_LEGACY = ROOT / "data" / "分仓W30数据.xlsx"
DEFAULT_CURRENT = ROOT / "data" / "基础信息表-SKU信息.xlsx"
OUTPUT_ROOT = ROOT / "outputs"
LOG_ROOT = ROOT / "logs"
LOCK_FILE = ROOT / "runtime" / "weekly_report.lock"


def load_env(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"配置文件不存在：{path}")
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ[key.strip()] = value.strip().strip('"').strip("'")


def env(name: str, default: str = "") -> str:
    return os.environ.get(name, default).strip()


def validate_config(require_webhook: bool) -> dict[str, Any]:
    required = ["PICKUP_DB_HOST", "PICKUP_DB_USER", "PICKUP_DB_PASSWORD"]
    if require_webhook:
        required.append("WECOM_WEBHOOK_URL")
    missing = [name for name in required if not env(name)]
    if missing:
        raise RuntimeError("缺少配置：" + ", ".join(missing))
    placeholders = [
        name for name in required
        if any(mark in os.environ.get(name, "") for mark in ("请填写", "YOUR_", "CHANGE_ME"))
    ]
    if placeholders:
        raise RuntimeError("配置仍是占位值：" + ", ".join(placeholders))
    if require_webhook:
        parsed = urllib.parse.urlparse(env("WECOM_WEBHOOK_URL"))
        key = urllib.parse.parse_qs(parsed.query).get("key", [""])[0]
        if parsed.netloc != "qyapi.weixin.qq.com" or not key:
            raise RuntimeError("WECOM_WEBHOOK_URL 格式不正确")
    for path in [DEFAULT_LEGACY, DEFAULT_CURRENT]:
        if not path.exists() or path.stat().st_size <= 0:
            raise FileNotFoundError(path)
    return {
        "database": f"{env('PICKUP_DB_HOST')}:{env('PICKUP_DB_PORT', '9030')}/{env('PICKUP_DB_NAME', 'ywzt_ods')}",
        "legacy_mapping": str(DEFAULT_LEGACY),
        "current_mapping": str(DEFAULT_CURRENT),
        "webhook_configured": bool(env("WECOM_WEBHOOK_URL")),
    }


def report_windows(report_date: date) -> tuple[WeekWindow, WeekWindow]:
    if report_date.weekday() != 4:
        raise ValueError("report_date 必须是周五")
    current_end = report_date - timedelta(days=1)
    current_start = current_end - timedelta(days=6)
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end - timedelta(days=6)

    def build(start: date, end: date) -> WeekWindow:
        return WeekWindow(f"W{start.isocalendar().week + 1}", start, end)

    return build(previous_start, previous_end), build(current_start, current_end)


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def decimal_value(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    return value if isinstance(value, Decimal) else Decimal(str(value))


def rounded(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def read_legacy_mapping(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["产品信息表"]
        result: dict[str, dict[str, Any]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sku = text(row[1] if len(row) > 1 else None)
            if sku:
                result[sku] = {
                    "产品编码": sku,
                    "旧商品名称": text(row[5] if len(row) > 5 else None),
                    "旧一级分类": text(row[2] if len(row) > 2 else None),
                    "旧二级分类": text(row[3] if len(row) > 3 else None),
                    "旧三级分类": text(row[4] if len(row) > 4 else None),
                    "旧箱规": row[6] if len(row) > 6 else None,
                }
        return result
    finally:
        workbook.close()


def read_current_mapping(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        headers = [text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        index = {name: pos for pos, name in enumerate(headers)}
        required = ["产品编码", "一级分类", "二级分类", "三级分类", "商品名称", "采购/售卖状态", "箱规", "产品大类"]
        missing = [name for name in required if name not in index]
        if missing:
            raise ValueError("新版映射表缺少字段：" + ", ".join(missing))
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sku = text(row[index["产品编码"]])
            if not sku:
                continue
            grouped[sku].append({
                "产品编码": sku,
                "新版商品名称": text(row[index["商品名称"]]),
                "新版一级分类": text(row[index["一级分类"]]),
                "新版二级分类": text(row[index["二级分类"]]),
                "新版三级分类": text(row[index["三级分类"]]),
                "采购/售卖状态": text(row[index["采购/售卖状态"]]),
                "新版箱规": row[index["箱规"]],
                "产品大类": text(row[index["产品大类"]]),
            })
        result: dict[str, dict[str, Any]] = {}
        duplicates: dict[str, int] = {}
        for sku, records in grouped.items():
            result[sku] = records[0]
            result[sku]["新版重复记录数"] = len(records)
            if len(records) > 1:
                duplicates[sku] = len(records)
        return result, duplicates
    finally:
        workbook.close()


def supplemental_category(record: dict[str, Any]) -> tuple[str, str, str]:
    level1 = text(record.get("新版一级分类"))
    major = text(record.get("产品大类"))
    name = text(record.get("新版商品名称"))
    if major != "产成品":
        return "排除", "新版映射", f"产品大类={major or '空'}，非周报目标产成品"
    category_map = {
        "烤肠类": "肉制品", "调理类": "肉制品", "油炸类": "肉制品", "培根类": "肉制品",
        "预制菜肴类": "肉制品", "面制品类": "面点", "烘焙类": "烘焙", "烘焙": "烘焙",
    }
    if level1 in category_map:
        return category_map[level1], "根据商品属性补充", f"产品大类=产成品，一级分类={level1}，商品名称={name}"
    if level1 in {"酱料类", "调味料及复合配料", "周边物料", "冰激凌", "饱饱博士"}:
        return "排除", "新版映射", f"一级分类={level1}，非周报三类"
    return "", "待判断", f"无法根据新版属性可靠归类：一级分类={level1 or '空'}，商品名称={name}"


def build_mapping(legacy: dict[str, dict[str, Any]], current: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    for sku in sorted(set(legacy) | set(current)):
        old, new = legacy.get(sku, {}), current.get(sku, {})
        old_category = text(old.get("旧一级分类"))
        if old_category:
            report_category = old_category if old_category in TARGET_CATEGORIES else "排除"
            source, basis = "原W30产品信息表", f"原一级分类={old_category}，原分类优先"
        elif new:
            report_category, source, basis = supplemental_category(new)
        else:
            report_category, source, basis = "", "待判断", "两份映射表均无记录"
        merged[sku] = {
            "产品编码": sku,
            "商品名称": text(new.get("新版商品名称")) or text(old.get("旧商品名称")),
            "原一级分类": old_category,
            "新版一级分类": text(new.get("新版一级分类")),
            "新版二级分类": text(new.get("新版二级分类")),
            "新版三级分类": text(new.get("新版三级分类")),
            "产品大类": text(new.get("产品大类")),
            "采购/售卖状态": text(new.get("采购/售卖状态")),
            "箱规": new.get("新版箱规") if new.get("新版箱规") not in (None, "") else old.get("旧箱规"),
            "周报分类": report_category,
            "分类来源": source,
            "分类依据": basis,
            "是否纳入汇总": "是" if report_category in TARGET_CATEGORIES else "否",
            "新版重复记录数": new.get("新版重复记录数", 0),
        }
    return merged


def connect_doris():
    return pymysql.connect(
        host=env("PICKUP_DB_HOST"),
        port=int(env("PICKUP_DB_PORT", "9030")),
        user=env("PICKUP_DB_USER"),
        password=os.environ.get("PICKUP_DB_PASSWORD", ""),
        database=env("PICKUP_DB_NAME", "ywzt_ods"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=int(env("DB_CONNECT_TIMEOUT", "20")),
        read_timeout=int(env("DB_READ_TIMEOUT", "300")),
        write_timeout=60,
    )


def query_window(connection, window: WeekWindow) -> list[dict[str, Any]]:
    with connection.cursor() as cursor:
        cursor.execute(SQL, (window.start_dt, window.end_exclusive))
        return cursor.fetchall()


def classify_rows(rows: list[dict[str, Any]], mapping: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    classified = []
    for raw in rows:
        row = dict(raw)
        mapped = mapping.get(text(row.get("商品编码")))
        category = mapped["周报分类"] if mapped else ""
        row["周报分类"] = category
        row["分类来源"] = mapped["分类来源"] if mapped else "待判断"
        row["是否纳入汇总"] = "是" if category in TARGET_CATEGORIES else "否"
        classified.append(row)
    return classified


def aggregate(rows: Iterable[dict[str, Any]]) -> dict[str, dict[str, Decimal]]:
    totals: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    for row in rows:
        category = text(row.get("周报分类"))
        if category not in TARGET_CATEGORIES:
            continue
        for channel in ["抖音", "拼多多系", "快手系", "淘宝系", "京东系", "线下社区团购_平台", "线下社区团购_分销", "线下其它"]:
            totals[category][channel] += decimal_value(row.get(channel))
    return totals


def summary_values(previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]]) -> dict[str, Any]:
    previous, current = aggregate(previous_rows), aggregate(current_rows)
    channel_map = [
        ("抖音", ["抖音"]), ("拼多多", ["拼多多系"]), ("快手", ["快手系"]),
        ("淘宝", ["淘宝系"]), ("京东", ["京东系"]),
        ("社区团购", ["线下社区团购_平台", "线下社区团购_分销"]),
        ("线下大区", ["线下其它"]),
    ]
    output: dict[str, Any] = {}
    for label, fields in channel_map:
        output[label] = {}
        for category in TARGET_CATEGORIES:
            prev = sum((previous[category][field] for field in fields), Decimal(0))
            curr = sum((current[category][field] for field in fields), Decimal(0))
            output[label][category] = {
                "previous_raw": format(prev, "f"), "current_raw": format(curr, "f"),
                "previous": rounded(prev), "current": rounded(curr),
            }
    return output


def build_unmapped(previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]], mapping: dict[str, dict[str, Any]], previous: WeekWindow, current: WeekWindow) -> list[dict[str, Any]]:
    bucket: dict[str, dict[str, Any]] = {}
    for window, rows in [(previous, previous_rows), (current, current_rows)]:
        for row in rows:
            sku = text(row.get("商品编码"))
            mapped = mapping.get(sku)
            if mapped and mapped.get("周报分类"):
                continue
            item = bucket.setdefault(sku, {
                "商品编码": sku, "商品描述": text(row.get("商品描述")),
                f"{previous.name}总发货量": Decimal(0), f"{current.name}总发货量": Decimal(0),
                "首次发货日期": text(row.get("发货日期")), "末次发货日期": text(row.get("发货日期")),
                "原因": mapped.get("分类依据") if mapped else "两份映射表均无此SKU",
            })
            item[f"{window.name}总发货量"] += decimal_value(row.get("总发货量"))
            ship_date = text(row.get("发货日期"))
            item["首次发货日期"] = min(item["首次发货日期"], ship_date)
            item["末次发货日期"] = max(item["末次发货日期"], ship_date)
    return sorted(bucket.values(), key=lambda item: -(decimal_value(item[f"{previous.name}总发货量"]) + decimal_value(item[f"{current.name}总发货量"])))


def build_metadata(previous: WeekWindow, current: WeekWindow, previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]], mapping: dict[str, dict[str, Any]], duplicates: dict[str, int], unmapped: list[dict[str, Any]]) -> dict[str, Any]:
    mapping_columns = [
        "产品编码", "商品名称", "原一级分类", "新版一级分类", "新版二级分类", "新版三级分类",
        "产品大类", "采购/售卖状态", "箱规", "周报分类", "分类来源", "分类依据", "是否纳入汇总", "新版重复记录数",
    ]
    mapping_rows = list(mapping.values())
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "previous": {"name": previous.name, "start": previous.start.isoformat(), "end": previous.end.isoformat()},
        "current": {"name": current.name, "start": current.start.isoformat(), "end": current.end.isoformat()},
        "detail_columns": DETAIL_COLUMNS,
        "mapping_columns": mapping_columns,
        "mapping_rows": [{column: row.get(column) for column in mapping_columns} for row in mapping_rows],
        "unmapped_columns": ["商品编码", "商品描述", f"{previous.name}总发货量", f"{current.name}总发货量", "首次发货日期", "末次发货日期", "原因"],
        "unmapped_rows": unmapped,
        "summary": summary_values(previous_rows, current_rows),
        "audit": {
            "previous_query_rows": len(previous_rows), "current_query_rows": len(current_rows),
            "previous_total_ship_qty": format(sum((decimal_value(row.get("总发货量")) for row in previous_rows), Decimal(0)), "f"),
            "current_total_ship_qty": format(sum((decimal_value(row.get("总发货量")) for row in current_rows), Decimal(0)), "f"),
            "mapping_rows": len(mapping_rows),
            "included_mapping_rows": sum(1 for row in mapping_rows if row["是否纳入汇总"] == "是"),
            "unmapped_skus": len(unmapped), "new_mapping_duplicate_codes": duplicates,
            "classification_source_counts": dict(Counter(row["分类来源"] for row in mapping_rows)),
        },
    }


def build_summary_rows(summary: dict[str, Any]) -> list[tuple[str, dict[str, tuple[int, int]]]]:
    base = ["抖音", "拼多多", "快手", "淘宝", "京东"]
    rows: list[tuple[str, dict[str, tuple[int, int]]]] = []
    for label in base:
        rows.append((label, {cat: (int(summary[label][cat]["previous"]), int(summary[label][cat]["current"])) for cat in TARGET_CATEGORIES}))

    def aggregate_row(label: str, sources: list[str]) -> None:
        rows.append((label, {cat: (
            sum(int(summary[source][cat]["previous"]) for source in sources),
            sum(int(summary[source][cat]["current"]) for source in sources),
        ) for cat in TARGET_CATEGORIES}))

    aggregate_row("线上合计", base)
    for label in ["社区团购", "线下大区"]:
        rows.append((label, {cat: (int(summary[label][cat]["previous"]), int(summary[label][cat]["current"])) for cat in TARGET_CATEGORIES}))
    aggregate_row("线下合计", ["社区团购", "线下大区"])
    online, offline = rows[5][1], rows[8][1]
    rows.append(("合计", {cat: (online[cat][0] + offline[cat][0], online[cat][1] + offline[cat][1]) for cat in TARGET_CATEGORIES}))
    return rows


def ratio(previous: int, current: int) -> tuple[str, str]:
    if previous == 0:
        return ("↑ 0%", "green") if current == 0 else ("新增", "green")
    pct = int(((Decimal(current) - Decimal(previous)) / Decimal(previous) * Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return (f"↑ {pct}%", "green") if pct >= 0 else (f"↓ {pct}%", "red")


def find_font(bold: bool) -> Path:
    configured = env("REPORT_FONT_BOLD_PATH" if bold else "REPORT_FONT_PATH")
    candidates = [Path(configured)] if configured else []
    candidates += [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simfang.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("未找到中文字体，请配置 REPORT_FONT_PATH / REPORT_FONT_BOLD_PATH")


def render_png(metadata: dict[str, Any], output_path: Path) -> None:
    colors = {
        "canvas": "#F5F7FA", "paper": "#FFFFFF", "ink": "#162033", "muted": "#7A8599",
        "faint": "#A5ADBA", "line": "#E8ECF2", "line_strong": "#CCD5E1", "blue": "#2563EB",
        "blue_dark": "#1746A2", "blue_tint": "#F1F6FF", "blue_tint_strong": "#E8F0FF",
        "green": "#14804A", "red": "#D64B3B", "total": "#17243A", "white": "#FFFFFF",
    }
    regular_path, bold_path = find_font(False), find_font(True)

    def font(size: int, bold: bool = False):
        return ImageFont.truetype(str(bold_path if bold else regular_path), size=size)

    image = Image.new("RGB", (2048, 1120), colors["canvas"])
    draw = ImageDraw.Draw(image)
    left, right = 92, 1956
    draw.rounded_rectangle((48, 38, 2000, 1078), radius=20, fill=colors["paper"])
    draw.text((left, 75), "渠道袋数周报", font=font(38, True), fill=colors["ink"])
    draw.rounded_rectangle((364, 80, 452, 118), radius=19, fill=colors["blue_tint_strong"])
    badge = metadata["current"]["name"]
    badge_font = font(19, True)
    badge_width = draw.textbbox((0, 0), badge, font=badge_font)[2]
    draw.text((408 - badge_width / 2, 86), badge, font=badge_font, fill=colors["blue_dark"])
    prev_dates = f'{metadata["previous"]["start"][5:].replace("-", ".")}–{metadata["previous"]["end"][5:].replace("-", ".")}'
    curr_dates = f'{metadata["current"]["start"][5:].replace("-", ".")}–{metadata["current"]["end"][5:].replace("-", ".")}'
    draw.text((left, 132), f'{metadata["previous"]["name"]}（{prev_dates}）  对比  {metadata["current"]["name"]}（{curr_dates}）', font=font(18), fill=colors["muted"])
    unit = "单位：袋"
    unit_width = draw.textbbox((0, 0), unit, font=font(18))[2]
    draw.text((right - unit_width, 89), unit, font=font(18), fill=colors["muted"])

    rows = build_summary_rows(metadata["summary"])
    table_top, channel_width = 192, 190
    group_width = (right - left - channel_width) // 3
    previous_width, current_width, ratio_width = group_width - 188 - 128, 188, 128
    header_bottom, row_h = table_top + 104, 73
    table_bottom = header_bottom + row_h * len(rows)
    x1 = left + channel_width
    for group_index in range(3):
        gx = x1 + group_index * group_width
        draw.rectangle((gx + previous_width, table_top + 56, gx + previous_width + current_width, table_bottom), fill=colors["blue_tint"])
    draw.text((left, table_top + 61), "各渠道", font=font(22, True), fill=colors["ink"])
    for group_index, category in enumerate(TARGET_CATEGORIES):
        gx = x1 + group_index * group_width
        category_font = font(25, True)
        tw = draw.textbbox((0, 0), category, font=category_font)[2]
        draw.text((gx + (group_width - tw) / 2, table_top + 9), category, font=category_font, fill=colors["ink"])
        draw.line((gx + 24, table_top + 53, gx + group_width - 24, table_top + 53), fill=colors["line"], width=1)
        for label, hx, hw, color, bold in [
            (metadata["previous"]["name"], gx, previous_width, colors["muted"], False),
            (metadata["current"]["name"], gx + previous_width, current_width, colors["blue_dark"], True),
            ("环比", gx + previous_width + current_width, ratio_width, colors["muted"], False),
        ]:
            heading_font = font(17, bold)
            tw = draw.textbbox((0, 0), label, font=heading_font)[2]
            draw.text((hx + (hw - tw) / 2, table_top + 70), label, font=heading_font, fill=color)
    draw.line((left, header_bottom, right, header_bottom), fill=colors["line_strong"], width=2)
    for group_index in range(3):
        gx = x1 + group_index * group_width
        draw.line((gx, table_top, gx, table_bottom - row_h), fill=colors["line"], width=1)

    for row_index, (label, values) in enumerate(rows):
        y = header_bottom + row_index * row_h
        is_total, is_subtotal = label == "合计", label in {"线上合计", "线下合计"}
        if is_total:
            draw.rectangle((left, y, right, y + row_h), fill=colors["total"])
        elif is_subtotal:
            draw.rectangle((left, y, right, y + row_h), fill="#FAFBFD")
        if label in {"线上合计", "社区团购", "线下合计", "合计"}:
            draw.line((left, y, right, y), fill=colors["line_strong"], width=2)
        elif row_index:
            draw.line((left, y, right, y), fill=colors["line"], width=1)
        draw.text((left, y + 21), label, font=font(21, is_subtotal or is_total), fill=colors["white"] if is_total else colors["ink"])
        for group_index, category in enumerate(TARGET_CATEGORIES):
            gx = x1 + group_index * group_width
            previous, current = values[category]
            ratio_text, ratio_kind = ratio(previous, current)
            prev_color = "#CBD5E1" if is_total else colors["muted"]
            curr_color = colors["white"] if is_total else (colors["ink"] if is_subtotal else colors["blue_dark"])
            ratio_color = ("#9FD8B7" if ratio_kind == "green" else "#FFB4A9") if is_total else colors[ratio_kind]
            for value, right_x, value_font, value_color in [
                (f"{previous:,}", gx + previous_width - 26, font(21, is_subtotal or is_total), prev_color),
                (f"{current:,}", gx + previous_width + current_width - 26, font(21, True), curr_color),
            ]:
                tw = draw.textbbox((0, 0), value, font=value_font)[2]
                draw.text((right_x - tw, y + 22), value, font=value_font, fill=value_color)
            ratio_font = font(18, True)
            tw = draw.textbbox((0, 0), ratio_text, font=ratio_font)[2]
            draw.text((gx + previous_width + current_width + (ratio_width - tw) / 2, y + 22), ratio_text, font=ratio_font, fill=ratio_color)
    draw.text((left, 1050), "统计口径：实际发货日期｜周五至下周四｜汇总数四舍五入", font=font(14), fill=colors["faint"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG", optimize=True)
    if output_path.stat().st_size > 2 * 1024 * 1024:
        raise RuntimeError(f"PNG 超过企业微信图片 2MB 限制：{output_path.stat().st_size} bytes")


def excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value


def write_workbook(metadata: dict[str, Any], previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = XlsxWorkbook(str(output_path), {"constant_memory": True})
    workbook.set_properties({"title": APP_NAME, "author": "Channel Bag Weekly Report"})
    ink, muted, line, blue, blue_tint, green, red, total = "#162033", "#7A8599", "#E8ECF2", "#1746A2", "#F1F6FF", "#14804A", "#D64B3B", "#17243A"
    fmt = {
        "title": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 20, "bold": True, "font_color": ink, "align": "left", "valign": "vcenter"}),
        "subtitle": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "font_color": muted, "align": "left", "valign": "vcenter"}),
        "group": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 12, "bold": True, "font_color": ink, "align": "center", "valign": "vcenter", "bottom": 1, "bottom_color": line}),
        "head": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": muted, "align": "center", "valign": "vcenter", "bottom": 1, "bottom_color": "#CCD5E1"}),
        "head_current": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": blue, "align": "center", "valign": "vcenter", "bg_color": blue_tint, "bottom": 1, "bottom_color": "#CCD5E1"}),
        "label": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "font_color": ink, "valign": "vcenter", "bottom": 1, "bottom_color": line}),
        "label_bold": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": ink, "valign": "vcenter", "bottom": 1, "bottom_color": "#CCD5E1"}),
        "number_prev": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "font_color": muted, "align": "right", "valign": "vcenter", "num_format": "#,##0", "bottom": 1, "bottom_color": line}),
        "number_curr": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": blue, "align": "right", "valign": "vcenter", "num_format": "#,##0", "bg_color": blue_tint, "bottom": 1, "bottom_color": line}),
        "number_bold": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": ink, "align": "right", "valign": "vcenter", "num_format": "#,##0", "bottom": 1, "bottom_color": "#CCD5E1"}),
        "ratio_green": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": green, "align": "center", "valign": "vcenter", "bottom": 1, "bottom_color": line}),
        "ratio_red": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": red, "align": "center", "valign": "vcenter", "bottom": 1, "bottom_color": line}),
        "total_label": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 11, "bold": True, "font_color": "#FFFFFF", "bg_color": total, "valign": "vcenter"}),
        "total_number": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 11, "bold": True, "font_color": "#FFFFFF", "bg_color": total, "align": "right", "valign": "vcenter", "num_format": "#,##0"}),
        "total_ratio_green": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": "#9FD8B7", "bg_color": total, "align": "center", "valign": "vcenter"}),
        "total_ratio_red": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "bold": True, "font_color": "#FFB4A9", "bg_color": total, "align": "center", "valign": "vcenter"}),
        "raw_head": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": total, "align": "center", "valign": "vcenter", "text_wrap": True}),
        "raw_text": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9}),
        "raw_date": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "num_format": "yyyy-mm-dd"}),
        "raw_number": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "num_format": "#,##0.######"}),
        "map_head": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": total, "align": "center", "valign": "vcenter", "text_wrap": True}),
        "map_text": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "valign": "top"}),
        "map_number": workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 9, "num_format": "#,##0.######"}),
    }

    summary_sheet = workbook.add_worksheet("汇总表")
    summary_sheet.hide_gridlines(2)
    summary_sheet.set_zoom(90)
    summary_sheet.set_column("A:A", 13)
    for col in ["B", "C", "E", "F", "H", "I"]:
        summary_sheet.set_column(f"{col}:{col}", 14)
    for col in ["D", "G", "J"]:
        summary_sheet.set_column(f"{col}:{col}", 10)
    summary_sheet.set_row(0, 34)
    summary_sheet.merge_range("A1:J1", f"渠道袋数周报  {metadata['current']['name']}", fmt["title"])
    summary_sheet.merge_range("A2:J2", f"{metadata['previous']['name']}（{metadata['previous']['start']}—{metadata['previous']['end']}） 对比 {metadata['current']['name']}（{metadata['current']['start']}—{metadata['current']['end']}）｜单位：袋", fmt["subtitle"])
    summary_sheet.merge_range("A3:A4", "各渠道", fmt["group"])
    for start_col, category in zip([1, 4, 7], TARGET_CATEGORIES):
        summary_sheet.merge_range(2, start_col, 2, start_col + 2, category, fmt["group"])
        summary_sheet.write(3, start_col, metadata["previous"]["name"], fmt["head"])
        summary_sheet.write(3, start_col + 1, metadata["current"]["name"], fmt["head_current"])
        summary_sheet.write(3, start_col + 2, "环比", fmt["head"])
    summary_rows = build_summary_rows(metadata["summary"])
    for idx, (label, values) in enumerate(summary_rows, start=4):
        is_total, is_subtotal = label == "合计", label in {"线上合计", "线下合计"}
        summary_sheet.set_row(idx, 25)
        summary_sheet.write(idx, 0, label, fmt["total_label"] if is_total else fmt["label_bold"] if is_subtotal else fmt["label"])
        for group_idx, category in enumerate(TARGET_CATEGORIES):
            previous, current = values[category]
            ratio_text, kind = ratio(previous, current)
            col = 1 + group_idx * 3
            if is_total:
                summary_sheet.write_number(idx, col, previous, fmt["total_number"])
                summary_sheet.write_number(idx, col + 1, current, fmt["total_number"])
                summary_sheet.write(idx, col + 2, ratio_text, fmt[f"total_ratio_{kind}"])
            elif is_subtotal:
                summary_sheet.write_number(idx, col, previous, fmt["number_bold"])
                summary_sheet.write_number(idx, col + 1, current, fmt["number_bold"])
                summary_sheet.write(idx, col + 2, ratio_text, fmt[f"ratio_{kind}"])
            else:
                summary_sheet.write_number(idx, col, previous, fmt["number_prev"])
                summary_sheet.write_number(idx, col + 1, current, fmt["number_curr"])
                summary_sheet.write(idx, col + 2, ratio_text, fmt[f"ratio_{kind}"])
    summary_sheet.freeze_panes(4, 1)
    summary_sheet.set_landscape()
    summary_sheet.fit_to_pages(1, 1)

    def add_detail_sheet(name: str, rows: list[dict[str, Any]]) -> None:
        sheet = workbook.add_worksheet(name)
        sheet.hide_gridlines(2)
        sheet.freeze_panes(1, 5)
        sheet.autofilter(0, 0, len(rows), len(DETAIL_COLUMNS) - 1)
        sheet.set_row(0, 32)
        widths = [12, 16, 26, 20, 36] + [14] * 15 + [13, 15, 13]
        for col, width in enumerate(widths):
            sheet.set_column(col, col, width)
        sheet.write_row(0, 0, DETAIL_COLUMNS, fmt["raw_head"])
        numeric_cols = set(range(5, 20))
        for row_index, row in enumerate(rows, start=1):
            for col_index, column in enumerate(DETAIL_COLUMNS):
                value = row.get(column)
                if col_index == 0 and value:
                    parsed = datetime.strptime(str(value), "%Y-%m-%d")
                    sheet.write_datetime(row_index, col_index, parsed, fmt["raw_date"])
                elif col_index in numeric_cols:
                    sheet.write_number(row_index, col_index, float(decimal_value(value)), fmt["raw_number"])
                else:
                    sheet.write(row_index, col_index, text(value), fmt["raw_text"])

    add_detail_sheet("上周明细", previous_rows)
    add_detail_sheet("本周明细", current_rows)

    def add_record_sheet(name: str, columns: list[str], records: list[dict[str, Any]]) -> None:
        sheet = workbook.add_worksheet(name)
        sheet.hide_gridlines(2)
        sheet.freeze_panes(1, 0)
        sheet.set_row(0, 32)
        sheet.write_row(0, 0, columns, fmt["map_head"])
        for col, column in enumerate(columns):
            width = 16
            if "名称" in column or "描述" in column:
                width = 34
            elif "依据" in column or "原因" in column:
                width = 46
            elif "编码" in column:
                width = 22
            sheet.set_column(col, col, width)
        if records:
            sheet.autofilter(0, 0, len(records), len(columns) - 1)
        for row_index, record in enumerate(records, start=1):
            for col_index, column in enumerate(columns):
                value = record.get(column)
                if isinstance(value, (int, float, Decimal)):
                    sheet.write_number(row_index, col_index, float(value), fmt["map_number"])
                else:
                    sheet.write(row_index, col_index, text(value), fmt["map_text"])

    add_record_sheet("产品分类映射", metadata["mapping_columns"], metadata["mapping_rows"])
    add_record_sheet("未映射SKU", metadata["unmapped_columns"], metadata["unmapped_rows"])
    workbook.close()


def json_post(url: str, payload: dict[str, Any]) -> dict[str, Any]:
    request = urllib.request.Request(url, data=json.dumps(payload, ensure_ascii=False).encode("utf-8"), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(request, timeout=90) as response:
        result = json.loads(response.read().decode("utf-8"))
    if result.get("errcode") != 0:
        raise RuntimeError(f"企业微信接口失败：{result}")
    return result


def send_image(webhook: str, path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    return json_post(webhook, {"msgtype": "image", "image": {"base64": base64.b64encode(data).decode("ascii"), "md5": hashlib.md5(data).hexdigest()}})


def encode_multipart(path: Path) -> tuple[bytes, str]:
    boundary = "----ChannelBag" + uuid.uuid4().hex
    content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    body = b"".join([
        f"--{boundary}\r\n".encode(),
        f'Content-Disposition: form-data; name="media"; filename="{path.name}"\r\n'.encode("utf-8"),
        f"Content-Type: {content_type}\r\n\r\n".encode(), path.read_bytes(), b"\r\n",
        f"--{boundary}--\r\n".encode(),
    ])
    return body, boundary


def send_file(webhook: str, path: Path) -> list[dict[str, Any]]:
    parsed = urllib.parse.urlparse(webhook)
    key = urllib.parse.parse_qs(parsed.query).get("key", [""])[0]
    upload_url = f"{parsed.scheme}://{parsed.netloc}/cgi-bin/webhook/upload_media?key={urllib.parse.quote(key)}&type=file"
    body, boundary = encode_multipart(path)
    request = urllib.request.Request(upload_url, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}", "Content-Length": str(len(body))})
    with urllib.request.urlopen(request, timeout=180) as response:
        upload = json.loads(response.read().decode("utf-8"))
    if upload.get("errcode") != 0 or not upload.get("media_id"):
        raise RuntimeError(f"企业微信文件上传失败：{upload}")
    sent = json_post(webhook, {"msgtype": "file", "file": {"media_id": upload["media_id"]}})
    return [upload, sent]


@contextmanager
def run_lock(force: bool = False):
    LOCK_FILE.parent.mkdir(parents=True, exist_ok=True)
    if LOCK_FILE.exists() and not force:
        age_hours = (time.time() - LOCK_FILE.stat().st_mtime) / 3600
        if age_hours < 6:
            raise RuntimeError(f"已有任务运行中：{LOCK_FILE}")
    LOCK_FILE.write_text(json.dumps({"pid": os.getpid(), "started_at": datetime.now().isoformat()}, ensure_ascii=False), encoding="utf-8")
    try:
        yield
    finally:
        LOCK_FILE.unlink(missing_ok=True)


def run_once(report_date: date, send: bool, output_dir: Path) -> dict[str, Any]:
    previous, current = report_windows(report_date)
    legacy = read_legacy_mapping(DEFAULT_LEGACY)
    latest, duplicates = read_current_mapping(DEFAULT_CURRENT)
    mapping = build_mapping(legacy, latest)
    connection = connect_doris()
    try:
        previous_rows = classify_rows(query_window(connection, previous), mapping)
        current_rows = classify_rows(query_window(connection, current), mapping)
    finally:
        connection.close()
    unmapped = build_unmapped(previous_rows, current_rows, mapping, previous, current)
    metadata = build_metadata(previous, current, previous_rows, current_rows, mapping, duplicates, unmapped)
    stem = f"渠道袋数周报_{previous.name}-{current.name}_{previous.start:%Y%m%d}-{current.end:%Y%m%d}"
    png_path, xlsx_path = output_dir / f"{stem}.png", output_dir / f"{stem}.xlsx"
    for path in [png_path, xlsx_path, output_dir / f"{stem}.audit.json"]:
        path.unlink(missing_ok=True)
    render_png(metadata, png_path)
    write_workbook(metadata, previous_rows, current_rows, xlsx_path)
    sent_results: list[dict[str, Any]] = []
    if send:
        webhook = env("WECOM_WEBHOOK_URL")
        sent_results.append({"type": "image", "result": send_image(webhook, png_path)})
        upload, sent = send_file(webhook, xlsx_path)
        sent_results.extend([{"type": "file_upload", "result": upload}, {"type": "file_send", "result": sent}])
    result = {
        "ok": True, "report_date": report_date.isoformat(), "previous": metadata["previous"], "current": metadata["current"],
        "sent_to_wecom": send, "wecom_results": sent_results,
        "outputs": {"png": {"path": str(png_path), "bytes": png_path.stat().st_size}, "xlsx": {"path": str(xlsx_path), "bytes": xlsx_path.stat().st_size}},
        "audit": metadata["audit"],
    }
    (output_dir / f"{stem}.audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Doris 渠道袋数周报：生成极简 PNG、可追溯 Excel，并可发送企业微信")
    parser.add_argument("--config", type=Path, default=CONFIG_FILE)
    parser.add_argument("--report-date", type=lambda value: datetime.strptime(value, "%Y-%m-%d").date(), default=date.today())
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--send", action="store_true")
    parser.add_argument("--check-config", action="store_true")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--retry-delay-seconds", type=int, default=None)
    args = parser.parse_args()
    load_env(args.config)
    config = validate_config(require_webhook=args.send)
    if args.check_config:
        print(json.dumps({"ok": True, "config": config}, ensure_ascii=False, indent=2))
        return
    output_dir = args.output_dir or (OUTPUT_ROOT / args.report_date.strftime("%Y%m%d"))
    output_dir.mkdir(parents=True, exist_ok=True)
    LOG_ROOT.mkdir(parents=True, exist_ok=True)
    retry_delay = args.retry_delay_seconds if args.retry_delay_seconds is not None else int(env("FULL_RETRY_DELAY_SECONDS", "600"))
    attempts: list[dict[str, Any]] = []
    with run_lock(args.force):
        for attempt in (1, 2):
            try:
                result = run_once(args.report_date, args.send, output_dir)
                result["attempt"] = attempt
                result["retried"] = attempt > 1
                print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
                return
            except Exception as exc:
                attempts.append({"attempt": attempt, "time": datetime.now().isoformat(timespec="seconds"), "error": str(exc), "traceback": traceback.format_exc()})
                if attempt == 1:
                    time.sleep(max(0, retry_delay))
        failure = {"ok": False, "report_date": args.report_date.isoformat(), "sent_to_wecom": False, "attempts": attempts}
        failure_path = output_dir / f"渠道袋数周报_失败_{args.report_date:%Y%m%d}.json"
        failure_path.write_text(json.dumps(failure, ensure_ascii=False, indent=2), encoding="utf-8")
        raise RuntimeError(json.dumps(failure, ensure_ascii=False))


if __name__ == "__main__":
    main()
