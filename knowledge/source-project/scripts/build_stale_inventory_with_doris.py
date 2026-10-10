from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Iterable

import openpyxl
import pymysql


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ENV_FILE = ROOT / "config" / "pickup.env"
DEFAULT_BASE_WORKBOOK = Path(
    r"C:\Users\24493\Documents\Codex\2026-07-03\wo\outputs\滚动更新-仓库滞销库存_组合结果_2026-07-03.xlsx"
)
DEFAULT_OUTPUT_DIR = Path(r"C:\Users\24493\Documents\Codex\2026-07-03\wo\outputs")
TODAY = date(2026, 7, 3)
SHELF_LIFE_DAYS = 365
SHELF_LIFE_MONTH_LABEL = 12
CARGO_OWNER_ID = "HJXH"

REQUESTED_WAREHOUSES = [
    "皇家小虎广西普冷仓库xs",
    "皇家小虎南昌万纬仓xs",
    "皇家小虎合肥万纬仓xs",
    "皇家小虎济南-万维仓xs",
    "皇家小虎泉州万纬仓xs",
    "皇家小虎昆明万纬仓xs",
    "皇家小虎杭州-万纬仓xs",
    "皇家小虎沈阳-鲜生活仓xs",
    "皇家小虎广州-弘一仓xs",
    "皇家小虎南京-普洛斯仓xs",
    "皇家小虎沈阳常温代发仓xs",
    "皇家小虎河北石家庄仓xs",
    "皇家小虎保定代发仓xs",
    "皇家小虎揭阳代发仓xs",
    "皇家小虎徐州代发仓xs",
    "皇家小虎阜阳代发仓xs",
    "皇家小虎山东潍坊代发仓xs-新（2.0）",
    "皇家小虎山东临沂代发仓xs",
    "皇家小虎山东济宁代发仓xs",
    "皇家小虎郑州-常温品代发仓xs",
    "皇家小虎长沙雨花常温-万纬仓xs",
]


@dataclass(frozen=True)
class WarehouseMatch:
    requested_name: str
    wms_code: str | None
    wms_desc: str | None
    match_type: str
    candidates: list[dict]


def load_env_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"env file not found: {path}")
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def connect():
    host = os.environ.get("PICKUP_DB_HOST", "").strip()
    user = os.environ.get("PICKUP_DB_USER", "").strip()
    password = os.environ.get("PICKUP_DB_PASSWORD", "")
    database = os.environ.get("PICKUP_DB_NAME", "ywzt_ods").strip() or "ywzt_ods"
    port = int(os.environ.get("PICKUP_DB_PORT", "9030"))
    missing = [
        name
        for name, value in {
            "PICKUP_DB_HOST": host,
            "PICKUP_DB_USER": user,
            "PICKUP_DB_PASSWORD": password,
        }.items()
        if not value
    ]
    if missing:
        raise SystemExit(f"missing database env vars: {', '.join(missing)}")
    return pymysql.connect(
        host=host,
        port=port,
        user=user,
        password=password,
        database=database,
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=180,
        write_timeout=60,
    )


def normalize_name(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.lower()
    text = text.replace("（", "(").replace("）", ")")
    text = text.replace("万维", "万纬")
    text = text.replace("普洛斯", "普洛斯")
    text = text.replace("雨花常温-万纬", "雨花万纬")
    text = text.replace("长沙雨花常温", "长沙雨花")
    text = re.sub(r"\s+", "", text)
    text = text.replace("-", "")
    text = text.replace("_", "")
    text = text.replace("仓库", "仓")
    text = text.replace("仓xs", "仓")
    text = text.removesuffix("xs")
    return text


def simple_name_tokens(value: str) -> set[str]:
    normalized = normalize_name(value)
    for removable in ("皇家小虎", "常温", "品", "代发", "冷", "仓", "新", "2.0"):
        normalized = normalized.replace(removable, "")
    return {token for token in re.split(r"[^0-9a-zA-Z\u4e00-\u9fff]+", normalized) if token}


def fetch_table_columns(connection, table_name: str) -> list[dict]:
    with connection.cursor() as cursor:
        cursor.execute(f"SHOW COLUMNS FROM {table_name}")
        return list(cursor.fetchall())


def fetch_wms(connection) -> list[dict]:
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT wms_code, wms_desc, wms_type
            FROM master_wms
            WHERE wms_desc LIKE '%%皇家小虎%%'
               OR wms_desc LIKE '%%万纬%%'
               OR wms_desc LIKE '%%万维%%'
               OR wms_desc LIKE '%%普冷%%'
               OR wms_desc LIKE '%%普洛斯%%'
               OR wms_desc LIKE '%%鲜生活%%'
               OR wms_desc LIKE '%%弘一%%'
               OR wms_desc LIKE '%%代发%%'
            """
        )
        return list(cursor.fetchall())


def match_warehouses(requested_names: list[str], wms_rows: list[dict]) -> list[WarehouseMatch]:
    by_normalized: dict[str, list[dict]] = {}
    for row in wms_rows:
        by_normalized.setdefault(normalize_name(row.get("wms_desc")), []).append(row)

    matches: list[WarehouseMatch] = []
    for requested in requested_names:
        requested_norm = normalize_name(requested)
        exact = by_normalized.get(requested_norm, [])
        if len(exact) == 1:
            row = exact[0]
            matches.append(WarehouseMatch(requested, row["wms_code"], row["wms_desc"], "normalized_exact", []))
            continue

        contains = [
            row
            for row in wms_rows
            if requested_norm in normalize_name(row.get("wms_desc"))
            or normalize_name(row.get("wms_desc")) in requested_norm
        ]
        if len(contains) == 1:
            row = contains[0]
            matches.append(WarehouseMatch(requested, row["wms_code"], row["wms_desc"], "contains", []))
            continue

        requested_tokens = simple_name_tokens(requested)
        scored = []
        for row in wms_rows:
            tokens = simple_name_tokens(str(row.get("wms_desc") or ""))
            overlap = len(requested_tokens & tokens)
            if overlap:
                scored.append((overlap, row))
        scored.sort(key=lambda item: (-item[0], str(item[1].get("wms_desc") or "")))
        candidates = [row for _, row in scored[:8]]
        if candidates and scored[0][0] >= 1 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
            row = candidates[0]
            matches.append(WarehouseMatch(requested, row["wms_code"], row["wms_desc"], "token_best", candidates))
        else:
            matches.append(WarehouseMatch(requested, None, None, "unmatched", candidates))
    return matches


def to_float(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def to_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt).date()
        except ValueError:
            pass
    return None


def final_status(days_left: int) -> str:
    if days_left <= 90:
        return "严重临期(3个月内)"
    if days_left <= 180:
        return "一般临期(6个月内)"
    if days_left <= 270:
        return "预警期(9个月内)"
    return "新鲜期"


def integerize(value: float) -> int | float:
    if abs(value - round(value)) < 0.000001:
        return int(round(value))
    return round(value, 6)


def fetch_inventory_rows(connection, matches: list[WarehouseMatch]) -> tuple[list[dict], list[dict]]:
    matched_codes = [item.wms_code for item in matches if item.wms_code]
    if not matched_codes:
        return [], []
    placeholders = ",".join(["%s"] * len(matched_codes))
    sql = f"""
        SELECT
            w.wms_code,
            mw.wms_desc,
            w.sku_code,
            COALESCE(ms.sku_name, MAX(w.sku_name)) AS sku_name,
            DATE(w.production_date) AS production_date,
            SUM(IFNULL(w.in_stock_quantity, 0)) AS in_stock_qty,
            SUM(IFNULL(CAST(w.freeze_quantity AS DECIMAL(20, 4)), 0)) AS freeze_qty,
            SUM(
                CASE
                    WHEN w.stock_lock IS NULL OR w.stock_lock = '' THEN 0
                    ELSE CAST(w.stock_lock AS DECIMAL(20, 4))
                END
            ) AS stock_lock_qty,
            MAX(DATE(w.created_date)) AS max_created_date,
            COUNT(*) AS source_rows
        FROM scm_warehous_manage w
        LEFT JOIN master_wms mw
            ON mw.wms_code = w.wms_code
        LEFT JOIN master_sku ms
            ON ms.sku_code = w.sku_code
        WHERE w.cargo_owner_id = %s
          AND w.wms_code IN ({placeholders})
          AND w.production_date IS NOT NULL
          AND w.sku_code IS NOT NULL
          AND w.sku_code <> ''
          AND (w.inventory_status = 'ZP' OR w.inventory_status IS NULL OR w.inventory_status = '')
          AND IFNULL(w.in_stock_quantity, 0) <> 0
          AND (ms.sku_group = '50' OR ms.sku_group IS NULL)
        GROUP BY
            w.wms_code,
            mw.wms_desc,
            w.sku_code,
            ms.sku_name,
            DATE(w.production_date)
    """
    with connection.cursor() as cursor:
        cursor.execute(sql, [CARGO_OWNER_ID, *matched_codes])
        rows = list(cursor.fetchall())

    summary_sql = f"""
        SELECT
            w.wms_code,
            mw.wms_desc,
            COUNT(*) AS raw_rows,
            COUNT(DISTINCT w.sku_code) AS sku_count,
            COUNT(DISTINCT DATE(w.production_date)) AS production_date_count,
            SUM(IFNULL(w.in_stock_quantity, 0)) AS in_stock_qty,
            MIN(DATE(w.production_date)) AS min_production_date,
            MAX(DATE(w.production_date)) AS max_production_date
        FROM scm_warehous_manage w
        LEFT JOIN master_wms mw
            ON mw.wms_code = w.wms_code
        LEFT JOIN master_sku ms
            ON ms.sku_code = w.sku_code
        WHERE w.cargo_owner_id = %s
          AND w.wms_code IN ({placeholders})
          AND w.production_date IS NOT NULL
          AND w.sku_code IS NOT NULL
          AND w.sku_code <> ''
          AND (w.inventory_status = 'ZP' OR w.inventory_status IS NULL OR w.inventory_status = '')
          AND IFNULL(w.in_stock_quantity, 0) <> 0
          AND (ms.sku_group = '50' OR ms.sku_group IS NULL)
        GROUP BY w.wms_code, mw.wms_desc
        ORDER BY mw.wms_desc
    """
    with connection.cursor() as cursor:
        cursor.execute(summary_sql, [CARGO_OWNER_ID, *matched_codes])
        summary = list(cursor.fetchall())
    return rows, summary


def row_from_inventory(item: dict) -> list[object]:
    production_date = to_date(item.get("production_date"))
    if production_date is None:
        raise ValueError(f"missing production date: {item}")
    expiry_date = production_date + timedelta(days=SHELF_LIFE_DAYS)
    days_left = (expiry_date - TODAY).days
    return [
        item.get("wms_desc") or item.get("wms_code"),
        str(item.get("sku_code") or ""),
        item.get("sku_name") or "",
        production_date,
        expiry_date,
        TODAY,
        days_left,
        integerize(to_float(item.get("in_stock_qty"))),
        SHELF_LIFE_MONTH_LABEL,
        "无",
        "",
        final_status(days_left),
    ]


def date_anomaly_reason(item: dict) -> str:
    production_date = to_date(item.get("production_date"))
    if production_date is None:
        return "缺生产日期"
    if production_date.year >= 2100:
        return "生产日期疑似占位"
    if production_date > TODAY:
        return "生产日期晚于今天"
    if production_date < date(2020, 1, 1):
        return "生产日期早于2020"
    return ""


def find_date_anomalies(inventory_rows: list[dict]) -> list[dict]:
    anomalies = []
    for item in inventory_rows:
        reason = date_anomaly_reason(item)
        if reason:
            row = dict(item)
            row["problem"] = reason
            anomalies.append(row)
    return anomalies


def copy_row_styles(ws, template_row: int, max_col: int):
    styles = []
    for col in range(1, max_col + 1):
        cell = ws.cell(row=template_row, column=col)
        styles.append(
            {
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
            }
        )
    return styles


def apply_row_styles(ws, row: int, styles) -> None:
    for col, style in enumerate(styles, start=1):
        cell = ws.cell(row=row, column=col)
        cell._style = copy(style["style"])
        cell.number_format = style["number_format"]
        cell.font = copy(style["font"])
        cell.fill = copy(style["fill"])
        cell.border = copy(style["border"])
        cell.alignment = copy(style["alignment"])
        cell.protection = copy(style["protection"])


def cell_key(values: Iterable[object]) -> tuple[str, str, str]:
    items = list(values)
    production = to_date(items[3])
    return (
        str(items[0] or "").strip(),
        str(items[1] or "").strip(),
        production.isoformat() if production else "",
    )


def write_workbook(
    base_workbook: Path,
    output_dir: Path,
    inventory_rows: list[dict],
    matches: list[WarehouseMatch],
    warehouse_summary: list[dict],
    table_columns: list[dict],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"滚动更新-仓库滞销库存_补数仓_{TODAY.isoformat()}.xlsx"
    shutil.copy2(base_workbook, output)

    wb = openpyxl.load_workbook(output)
    ws = wb["Sheet1"]
    headers = ["仓库", "货品编码", "货品名称", "生产日期", "失效日期", "今天日期", "剩余天数", "本周数量", "保质期", "效期状态", "", "效期状态"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    styles = copy_row_styles(ws, 2 if ws.max_row >= 2 else 1, 12)
    existing_keys = {
        cell_key(ws.cell(row=row, column=col).value for col in range(1, 13))
        for row in range(2, ws.max_row + 1)
    }
    added = 0
    skipped_existing = 0
    for item in sorted(inventory_rows, key=lambda r: (str(r.get("wms_desc") or ""), str(r.get("sku_code") or ""), str(r.get("production_date") or ""))):
        values = row_from_inventory(item)
        key = cell_key(values)
        if key in existing_keys:
            skipped_existing += 1
            continue
        row_index = ws.max_row + 1
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col, value=value)
        apply_row_styles(ws, row_index, styles)
        for col in (4, 5, 6):
            ws.cell(row=row_index, column=col).number_format = "yyyy-mm-dd"
        ws.cell(row=row_index, column=8).number_format = "0.######"
        existing_keys.add(key)
        added += 1

    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    date_anomalies = find_date_anomalies(inventory_rows)

    if date_anomalies:
        warning_ws = wb["问题清单"] if "问题清单" in wb.sheetnames else wb.create_sheet("问题清单")
        headers = [warning_ws.cell(row=1, column=col).value for col in range(1, warning_ws.max_column + 1)]
        if not headers or headers == [None]:
            headers = ["来源", "仓库", "货品编码", "货品名称", "生产日期", "原数量", "问题"]
            for col, header in enumerate(headers, start=1):
                warning_ws.cell(row=1, column=col, value=header)
        header_index = {str(header): idx + 1 for idx, header in enumerate(headers) if header is not None}
        for item in date_anomalies:
            row_index = warning_ws.max_row + 1
            values = {
                "来源": "数仓",
                "仓库": item.get("wms_desc") or item.get("wms_code"),
                "货品编码": item.get("sku_code"),
                "货品名称": item.get("sku_name"),
                "生产日期": to_date(item.get("production_date")),
                "原数量": integerize(to_float(item.get("in_stock_qty"))),
                "问题": item.get("problem"),
            }
            for header, value in values.items():
                col = header_index.get(header)
                if col:
                    warning_ws.cell(row=row_index, column=col, value=value)
            prod_col = header_index.get("生产日期")
            if prod_col:
                warning_ws.cell(row=row_index, column=prod_col).number_format = "yyyy-mm-dd"

    audit_ws = wb["数仓补充审计"] if "数仓补充审计" in wb.sheetnames else wb.create_sheet("数仓补充审计")
    audit_ws.delete_rows(1, audit_ws.max_row)
    audit_rows = [
        ["项目", "值"],
        ["来源表", "ywzt_ods.scm_warehous_manage"],
        ["货主", CARGO_OWNER_ID],
        ["今天日期", TODAY.isoformat()],
        ["请求仓库数", len(REQUESTED_WAREHOUSES)],
        ["匹配仓库数", sum(1 for item in matches if item.wms_code)],
        ["数仓汇总行数", len(inventory_rows)],
        ["追加到Sheet1行数", added],
        ["已存在跳过行数", skipped_existing],
        ["数仓生产日期异常行数", len(date_anomalies)],
    ]
    for row_idx, row_values in enumerate(audit_rows, start=1):
        for col_idx, value in enumerate(row_values, start=1):
            audit_ws.cell(row=row_idx, column=col_idx, value=value)

    start_row = len(audit_rows) + 3
    match_headers = ["请求仓库", "匹配编码", "匹配仓库", "匹配方式", "候选"]
    for col_idx, value in enumerate(match_headers, start=1):
        audit_ws.cell(row=start_row, column=col_idx, value=value)
    for offset, item in enumerate(matches, start=1):
        candidates = "；".join(f"{row.get('wms_code')}:{row.get('wms_desc')}" for row in item.candidates[:5])
        values = [item.requested_name, item.wms_code, item.wms_desc, item.match_type, candidates]
        for col_idx, value in enumerate(values, start=1):
            audit_ws.cell(row=start_row + offset, column=col_idx, value=value)

    start_row = start_row + len(matches) + 3
    summary_headers = ["仓库编码", "仓库名称", "原始行数", "SKU数", "生产日期数", "库存数量", "最早生产日期", "最晚生产日期"]
    for col_idx, value in enumerate(summary_headers, start=1):
        audit_ws.cell(row=start_row, column=col_idx, value=value)
    for offset, item in enumerate(warehouse_summary, start=1):
        values = [
            item.get("wms_code"),
            item.get("wms_desc"),
            item.get("raw_rows"),
            item.get("sku_count"),
            item.get("production_date_count"),
            integerize(to_float(item.get("in_stock_qty"))),
            to_date(item.get("min_production_date")),
            to_date(item.get("max_production_date")),
        ]
        for col_idx, value in enumerate(values, start=1):
            audit_ws.cell(row=start_row + offset, column=col_idx, value=value)
        for col_idx in (7, 8):
            audit_ws.cell(row=start_row + offset, column=col_idx).number_format = "yyyy-mm-dd"

    start_row = start_row + len(warehouse_summary) + 3
    anomaly_headers = ["仓库编码", "仓库名称", "货品编码", "货品名称", "生产日期", "库存数量", "问题"]
    for col_idx, value in enumerate(anomaly_headers, start=1):
        audit_ws.cell(row=start_row, column=col_idx, value=value)
    if date_anomalies:
        for offset, item in enumerate(date_anomalies, start=1):
            values = [
                item.get("wms_code"),
                item.get("wms_desc"),
                item.get("sku_code"),
                item.get("sku_name"),
                to_date(item.get("production_date")),
                integerize(to_float(item.get("in_stock_qty"))),
                item.get("problem"),
            ]
            for col_idx, value in enumerate(values, start=1):
                audit_ws.cell(row=start_row + offset, column=col_idx, value=value)
            audit_ws.cell(row=start_row + offset, column=5).number_format = "yyyy-mm-dd"
    else:
        audit_ws.cell(row=start_row + 1, column=1, value="无")

    anomaly_ws = wb["数仓日期异常"] if "数仓日期异常" in wb.sheetnames else wb.create_sheet("数仓日期异常")
    anomaly_ws.delete_rows(1, anomaly_ws.max_row)
    for col_idx, value in enumerate(anomaly_headers, start=1):
        anomaly_ws.cell(row=1, column=col_idx, value=value)
    for offset, item in enumerate(date_anomalies, start=1):
        values = [
            item.get("wms_code"),
            item.get("wms_desc"),
            item.get("sku_code"),
            item.get("sku_name"),
            to_date(item.get("production_date")),
            integerize(to_float(item.get("in_stock_qty"))),
            item.get("problem"),
        ]
        for col_idx, value in enumerate(values, start=1):
            anomaly_ws.cell(row=1 + offset, column=col_idx, value=value)
        anomaly_ws.cell(row=1 + offset, column=5).number_format = "yyyy-mm-dd"

    start_row = start_row + len(date_anomalies) + 3
    audit_ws.cell(row=start_row, column=1, value="scm_warehous_manage字段")
    for offset, column in enumerate(table_columns, start=1):
        audit_ws.cell(row=start_row + offset, column=1, value=column.get("Field"))
        audit_ws.cell(row=start_row + offset, column=2, value=column.get("Type"))

    wb.save(output)

    audit_json = {
        "output": str(output),
        "source_table": "ywzt_ods.scm_warehous_manage",
        "cargo_owner_id": CARGO_OWNER_ID,
        "today": TODAY.isoformat(),
        "requested_warehouses": REQUESTED_WAREHOUSES,
        "matches": [item.__dict__ for item in matches],
        "inventory_grouped_rows": len(inventory_rows),
        "added_sheet1_rows": added,
        "skipped_existing_rows": skipped_existing,
        "final_sheet1_rows": ws.max_row - 1,
        "warehouse_summary": warehouse_summary,
        "date_anomalies": date_anomalies,
    }
    (output_dir / f"滚动更新-仓库滞销库存_补数仓_{TODAY.isoformat()}_audit.json").write_text(
        json.dumps(audit_json, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    return output


def inspect_only(connection) -> dict:
    table_columns = fetch_table_columns(connection, "scm_warehous_manage")
    wms_rows = fetch_wms(connection)
    matches = match_warehouses(REQUESTED_WAREHOUSES, wms_rows)
    inventory_rows, summary = fetch_inventory_rows(connection, matches)
    return {
        "source_table": "ywzt_ods.scm_warehous_manage",
        "table_columns": [{"Field": row.get("Field"), "Type": row.get("Type")} for row in table_columns],
        "matches": [item.__dict__ for item in matches],
        "inventory_grouped_rows": len(inventory_rows),
        "warehouse_summary": summary,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="把指定皇家小虎仓库的数仓库存补充到滞销库存组合表")
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV_FILE)
    parser.add_argument("--base-workbook", type=Path, default=DEFAULT_BASE_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--inspect", action="store_true")
    args = parser.parse_args()

    load_env_file(args.env_file)
    connection = connect()
    try:
        table_columns = fetch_table_columns(connection, "scm_warehous_manage")
        wms_rows = fetch_wms(connection)
        matches = match_warehouses(REQUESTED_WAREHOUSES, wms_rows)
        inventory_rows, summary = fetch_inventory_rows(connection, matches)

        if args.inspect:
            print(
                json.dumps(
                    {
                        "source_table": "ywzt_ods.scm_warehous_manage",
                        "columns": [{"Field": row.get("Field"), "Type": row.get("Type")} for row in table_columns],
                        "matches": [item.__dict__ for item in matches],
                        "inventory_grouped_rows": len(inventory_rows),
                        "warehouse_summary": summary,
                    },
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                )
            )
            return 0

        output = write_workbook(args.base_workbook, args.output_dir, inventory_rows, matches, summary, table_columns)
        print(
            json.dumps(
                {
                    "ok": True,
                    "output": str(output),
                    "inventory_grouped_rows": len(inventory_rows),
                    "matched_warehouses": sum(1 for item in matches if item.wms_code),
                    "unmatched_warehouses": [item.requested_name for item in matches if not item.wms_code],
                    "warehouse_count_with_rows": len(summary),
                    "date_anomaly_rows": len(find_date_anomalies(inventory_rows)),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0
    finally:
        connection.close()


if __name__ == "__main__":
    raise SystemExit(main())
