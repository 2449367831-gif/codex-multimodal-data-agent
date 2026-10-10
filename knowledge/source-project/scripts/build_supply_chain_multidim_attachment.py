from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import pymysql
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "outputs" / "supply_chain_monthly_202606"
ENV_FILE = ROOT / "config" / "pickup.env"
OUT_FILE = DATA_DIR / "HJXH供应链6月份月度经营汇报_多维数据附件_2026-06.xlsx"
START = date(2026, 6, 1)
END = date(2026, 7, 1)


def load_env() -> None:
    for raw in ENV_FILE.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ[k.strip()] = v.strip().strip('"').strip("'")


def connect():
    return pymysql.connect(
        host=os.environ["PICKUP_DB_HOST"],
        port=int(os.environ.get("PICKUP_DB_PORT", "9030")),
        user=os.environ["PICKUP_DB_USER"],
        password=os.environ["PICKUP_DB_PASSWORD"],
        database=os.environ.get("PICKUP_DB_NAME", "ywzt_ods"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=600,
        write_timeout=60,
    )


def fetch_all(conn, sql: str, params: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(sql, params or {})
        return list(cur.fetchall())


def latest_base_workbook() -> Path:
    candidates = sorted(
        [p for p in DATA_DIR.glob("*.xlsx") if "多维" not in p.name and not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("No source workbook found")
    return candidates[0]


def read_sheet(wb, name: str) -> list[dict[str, Any]]:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def to_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except Exception:
        return 0.0


def fmt_ratio(value: float) -> float:
    return round(value, 4)


def defect_filter(alias: str) -> str:
    return "\n".join(
        f"AND IFNULL({alias}.wms_desc, '') NOT LIKE '%%{word}%%'"
        for word in ["残次", "次品", "废品", "报废", "残品"]
    )


def query_extra(conn) -> dict[str, list[dict[str, Any]]]:
    wh_filter = defect_filter("mw")
    extra = {
        "warehouse_sku_detail": fetch_all(
            conn,
            f"""
WITH sales AS (
  SELECT h.facility_code warehouse_code, MAX(mw.wms_desc) warehouse_name, d.sku_code,
         MAX(ms.sku_name) sku_name, COUNT(DISTINCT h.outbound_code) outbound_order_count,
         SUM(IFNULL(d.ship_qty,0)) outbound_qty,
         SUM(IFNULL(d.ship_qty,0)*IFNULL(ms.net_weight,0)/1000) outbound_ton
  FROM scm_sales_shipment_header h
  JOIN scm_sales_shipment_detail d ON d.outbound_code=h.outbound_code
  LEFT JOIN master_wms mw ON mw.wms_code=h.facility_code
  LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
  WHERE h.ship_at >= %(start)s AND h.ship_at < %(end)s
    AND h.outbound_code LIKE 'HJXH%%'
    AND h.status IN ('700','900','701','901')
    AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
    AND ms.sku_group='50' AND mw.wms_type IN (10,11)
    {wh_filter}
  GROUP BY h.facility_code, d.sku_code
),
stock AS (
  SELECT w.wms_code warehouse_code, MAX(mw.wms_desc) warehouse_name, w.sku_code,
         MAX(ms.sku_name) sku_name,
         SUM(CASE WHEN DATE(w.backups_date)=DATE_SUB(%(end)s, INTERVAL 1 DAY) THEN IFNULL(w.in_stock_quantity,0) ELSE 0 END) closing_stock_qty,
         SUM(CASE WHEN DATE(w.backups_date)=DATE_SUB(%(end)s, INTERVAL 1 DAY) THEN IFNULL(w.in_stock_quantity,0)*IFNULL(ms.net_weight,0)/1000 ELSE 0 END) closing_stock_ton,
         SUM(IFNULL(w.in_stock_quantity,0)*IFNULL(ms.net_weight,0)/1000)/COUNT(DISTINCT DATE(w.backups_date)) avg_stock_ton
  FROM scm_warehous_manage_backups w
  LEFT JOIN master_wms mw ON mw.wms_code=w.wms_code
  LEFT JOIN master_sku ms ON ms.sku_code=w.sku_code
  WHERE w.backups_date >= %(start)s AND w.backups_date < %(end)s
    AND w.cargo_owner_id='HJXH' AND ms.sku_group='50' AND mw.wms_type IN (10,11)
    {wh_filter}
  GROUP BY w.wms_code, w.sku_code
),
all_keys AS (
  SELECT warehouse_code, sku_code FROM sales
  UNION SELECT warehouse_code, sku_code FROM stock
)
SELECT k.warehouse_code, COALESCE(st.warehouse_name,s.warehouse_name,k.warehouse_code) warehouse_name,
       k.sku_code, COALESCE(st.sku_name,s.sku_name,k.sku_code) sku_name,
       IFNULL(s.outbound_order_count,0) outbound_order_count,
       IFNULL(s.outbound_qty,0) outbound_qty,
       IFNULL(s.outbound_ton,0) outbound_ton,
       IFNULL(st.closing_stock_qty,0) closing_stock_qty,
       IFNULL(st.closing_stock_ton,0) closing_stock_ton,
       IFNULL(st.avg_stock_ton,0) avg_stock_ton,
       CASE WHEN IFNULL(s.outbound_ton,0)>0 THEN IFNULL(st.closing_stock_ton,0)/(IFNULL(s.outbound_ton,0)/30) ELSE NULL END closing_cover_days
FROM all_keys k
LEFT JOIN sales s ON s.warehouse_code=k.warehouse_code AND s.sku_code=k.sku_code
LEFT JOIN stock st ON st.warehouse_code=k.warehouse_code AND st.sku_code=k.sku_code
ORDER BY outbound_ton DESC, closing_stock_ton DESC
LIMIT 5000
""",
            {"start": START, "end": END},
        ),
        "warehouse_age": fetch_all(
            conn,
            f"""
SELECT w.wms_code warehouse_code, MAX(mw.wms_desc) warehouse_name,
       CASE
         WHEN w.production_date IS NULL THEN '无生产日期'
         WHEN DATEDIFF(DATE_SUB(%(end)s, INTERVAL 1 DAY), DATE(w.production_date)) <= 30 THEN '0-30天'
         WHEN DATEDIFF(DATE_SUB(%(end)s, INTERVAL 1 DAY), DATE(w.production_date)) <= 60 THEN '31-60天'
         WHEN DATEDIFF(DATE_SUB(%(end)s, INTERVAL 1 DAY), DATE(w.production_date)) <= 90 THEN '61-90天'
         WHEN DATEDIFF(DATE_SUB(%(end)s, INTERVAL 1 DAY), DATE(w.production_date)) <= 180 THEN '91-180天'
         WHEN DATEDIFF(DATE_SUB(%(end)s, INTERVAL 1 DAY), DATE(w.production_date)) <= 270 THEN '181-270天'
         ELSE '270天以上'
       END age_bucket,
       COUNT(DISTINCT w.sku_code) sku_count,
       SUM(IFNULL(w.in_stock_quantity,0)) stock_qty,
       SUM(IFNULL(w.in_stock_quantity,0)*IFNULL(ms.net_weight,0)/1000) stock_ton
FROM scm_warehous_manage_backups w
LEFT JOIN master_wms mw ON mw.wms_code=w.wms_code
LEFT JOIN master_sku ms ON ms.sku_code=w.sku_code
WHERE DATE(w.backups_date)=DATE_SUB(%(end)s, INTERVAL 1 DAY)
  AND w.cargo_owner_id='HJXH' AND ms.sku_group='50' AND mw.wms_type IN (10,11)
  {wh_filter}
GROUP BY w.wms_code, age_bucket
ORDER BY warehouse_name, stock_ton DESC
""",
            {"end": END},
        ),
        "purchase_wms_type": fetch_all(
            conn,
            """
SELECT IFNULL(CAST(mw.wms_type AS CHAR),'未维护') wms_type,
       CASE
         WHEN mw.wms_type=10 THEN '仓储/中心仓'
         WHEN mw.wms_type=11 THEN '分仓/仓储'
         WHEN mw.wms_type=12 THEN '供应商/外部'
         WHEN mw.wms_type=16 THEN '工厂/成品仓'
         ELSE '其他/未维护'
       END wms_type_name,
       COUNT(DISTINCT h.po_code) purchase_order_count,
       COUNT(*) detail_rows,
       SUM(IFNULL(d.received_quantity,0)) purchase_qty,
       SUM(IFNULL(d.received_quantity,0)*IFNULL(ms.net_weight,0)/1000) purchase_ton
FROM scm_po_header h
JOIN scm_po_detail d ON d.po_code=h.po_code
LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
LEFT JOIN master_wms mw ON mw.wms_code=h.take_delivery_wms
WHERE d.created_date >= %(start)s AND d.created_date < %(end)s
  AND h.cargo_owner_id='HJXH' AND h.po_type='CCP' AND h.status=0 AND d.status=0
  AND ms.sku_group='50'
GROUP BY wms_type, wms_type_name
ORDER BY purchase_ton DESC
""",
            {"start": START, "end": END},
        ),
        "purchase_vendor": fetch_all(
            conn,
            """
SELECT IFNULL(h.vendor_code,'') vendor_code, IFNULL(h.vendor_name,'未维护供应商') vendor_name,
       COUNT(DISTINCT h.po_code) purchase_order_count,
       COUNT(DISTINCT d.sku_code) sku_count,
       SUM(IFNULL(d.received_quantity,0)) purchase_qty,
       SUM(IFNULL(d.received_quantity,0)*IFNULL(ms.net_weight,0)/1000) purchase_ton
FROM scm_po_header h
JOIN scm_po_detail d ON d.po_code=h.po_code
LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
WHERE d.created_date >= %(start)s AND d.created_date < %(end)s
  AND h.cargo_owner_id='HJXH' AND h.po_type='CCP' AND h.status=0 AND d.status=0
  AND ms.sku_group='50'
GROUP BY h.vendor_code, h.vendor_name
ORDER BY purchase_ton DESC
LIMIT 200
""",
            {"start": START, "end": END},
        ),
        "purchase_destination": fetch_all(
            conn,
            """
SELECT h.take_delivery_wms warehouse_code, IFNULL(mw.wms_desc,'未维护到货仓') warehouse_name,
       IFNULL(CAST(mw.wms_type AS CHAR),'未维护') wms_type,
       COUNT(DISTINCT h.po_code) purchase_order_count,
       COUNT(DISTINCT d.sku_code) sku_count,
       SUM(IFNULL(d.received_quantity,0)) purchase_qty,
       SUM(IFNULL(d.received_quantity,0)*IFNULL(ms.net_weight,0)/1000) purchase_ton
FROM scm_po_header h
JOIN scm_po_detail d ON d.po_code=h.po_code
LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
LEFT JOIN master_wms mw ON mw.wms_code=h.take_delivery_wms
WHERE d.created_date >= %(start)s AND d.created_date < %(end)s
  AND h.cargo_owner_id='HJXH' AND h.po_type='CCP' AND h.status=0 AND d.status=0
  AND ms.sku_group='50'
GROUP BY h.take_delivery_wms, mw.wms_desc, mw.wms_type
ORDER BY purchase_ton DESC
LIMIT 200
""",
            {"start": START, "end": END},
        ),
        "outbound_type": fetch_all(
            conn,
            f"""
SELECT h.ref_order_type,
       CASE
         WHEN h.ref_order_type='JYCK' THEN '交易出库'
         WHEN h.ref_order_type LIKE 'SO%%' THEN '销售订单出库'
         ELSE h.ref_order_type
       END outbound_type_name,
       COUNT(DISTINCT h.outbound_code) outbound_order_count,
       COUNT(DISTINCT d.sku_code) sku_count,
       SUM(IFNULL(d.ship_qty,0)) outbound_qty,
       SUM(IFNULL(d.ship_qty,0)*IFNULL(ms.net_weight,0)/1000) outbound_ton
FROM scm_sales_shipment_header h
JOIN scm_sales_shipment_detail d ON d.outbound_code=h.outbound_code
LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
LEFT JOIN master_wms mw ON mw.wms_code=h.facility_code
WHERE h.ship_at >= %(start)s AND h.ship_at < %(end)s
  AND h.outbound_code LIKE 'HJXH%%'
  AND h.status IN ('700','900','701','901')
  AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
  AND ms.sku_group='50' AND mw.wms_type IN (10,11)
  {wh_filter}
GROUP BY h.ref_order_type
ORDER BY outbound_ton DESC
""",
            {"start": START, "end": END},
        ),
    }
    fulfillment_base = f"""
WITH order_base AS (
  SELECT h.outbound_code,
         DATE(h.ship_at) ship_date,
         DATE(h.created) created_date,
         h.created,
         h.ship_at,
         h.facility_code warehouse_code,
         MAX(mw.wms_desc) warehouse_name,
         h.ref_order_type,
         CASE
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z004' THEN '线上-JYCK/Z004'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' AND ms_store.group1='03' THEN '线下KA-JYCK/group1=03'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' AND ms_store.price_group='122' THEN '线下特价组-JYCK/price_group=122'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' THEN '线下经销-JYCK/Z001'
           WHEN h.ref_order_type LIKE 'SO%%' THEN CONCAT('SO销售-', h.ref_order_type)
           ELSE h.ref_order_type
         END channel_group,
         h.status,
         IFNULL(h.carrier_code, '') carrier_code,
         IFNULL(h.error_code, '') error_code,
         IFNULL(h.error_desc, '') error_desc,
         IFNULL(h.hold_code, '') hold_code,
         IFNULL(h.status_flow_code, '') status_flow_code,
         COUNT(DISTINCT d.sku_code) sku_count,
         SUM(IFNULL(d.ship_qty,0)) outbound_qty,
         SUM(IFNULL(d.ship_qty,0)*IFNULL(ms.net_weight,0)/1000) outbound_ton,
         TIMESTAMPDIFF(MINUTE, h.created, h.ship_at) / 60.0 fulfill_hours
  FROM scm_sales_shipment_header h
  JOIN scm_sales_shipment_detail d ON d.outbound_code=h.outbound_code
  LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
  LEFT JOIN master_wms mw ON mw.wms_code=h.facility_code
  LEFT JOIN master_store ms_store ON ms_store.code=h.store_code
  WHERE h.ship_at >= %(start)s AND h.ship_at < %(end)s
    AND h.created IS NOT NULL AND h.ship_at IS NOT NULL
    AND h.outbound_code LIKE 'HJXH%%'
    AND h.status IN ('700','900','701','901')
    AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
    AND ms.sku_group='50' AND mw.wms_type IN (10,11)
    {wh_filter}
  GROUP BY h.outbound_code, DATE(h.ship_at), DATE(h.created), h.created, h.ship_at, h.facility_code, h.ref_order_type,
           channel_group, h.status, h.carrier_code, h.error_code, h.error_desc, h.hold_code, h.status_flow_code
)
"""
    extra["fulfillment_overview"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT '6月履约' metric,
       COUNT(*) order_count,
       SUM(outbound_qty) outbound_qty,
       SUM(outbound_ton) outbound_ton,
       AVG(fulfill_hours) avg_fulfill_hours,
       PERCENTILE_APPROX(fulfill_hours, 0.5) median_fulfill_hours,
       SUM(CASE WHEN fulfill_hours <= 24 THEN 1 ELSE 0 END) within_24h_orders,
       SUM(CASE WHEN fulfill_hours <= 48 THEN 1 ELSE 0 END) within_48h_orders,
       SUM(CASE WHEN fulfill_hours > 48 THEN 1 ELSE 0 END) over_48h_orders,
       SUM(CASE WHEN fulfill_hours <= 24 THEN 1 ELSE 0 END) / COUNT(*) within_24h_rate,
       SUM(CASE WHEN fulfill_hours <= 48 THEN 1 ELSE 0 END) / COUNT(*) within_48h_rate
FROM order_base
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_daily"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT ship_date,
       COUNT(*) order_count,
       SUM(outbound_qty) outbound_qty,
       SUM(outbound_ton) outbound_ton,
       AVG(fulfill_hours) avg_fulfill_hours,
       PERCENTILE_APPROX(fulfill_hours, 0.5) median_fulfill_hours,
       SUM(CASE WHEN fulfill_hours <= 24 THEN 1 ELSE 0 END) / COUNT(*) within_24h_rate,
       SUM(CASE WHEN fulfill_hours <= 48 THEN 1 ELSE 0 END) / COUNT(*) within_48h_rate,
       SUM(CASE WHEN fulfill_hours > 48 THEN 1 ELSE 0 END) over_48h_orders
FROM order_base
GROUP BY ship_date
ORDER BY ship_date
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_warehouse"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT warehouse_code, warehouse_name,
       COUNT(*) order_count,
       SUM(outbound_qty) outbound_qty,
       SUM(outbound_ton) outbound_ton,
       AVG(fulfill_hours) avg_fulfill_hours,
       PERCENTILE_APPROX(fulfill_hours, 0.5) median_fulfill_hours,
       SUM(CASE WHEN fulfill_hours <= 24 THEN 1 ELSE 0 END) / COUNT(*) within_24h_rate,
       SUM(CASE WHEN fulfill_hours <= 48 THEN 1 ELSE 0 END) / COUNT(*) within_48h_rate,
       SUM(CASE WHEN fulfill_hours > 48 THEN 1 ELSE 0 END) over_48h_orders
FROM order_base
GROUP BY warehouse_code, warehouse_name
ORDER BY order_count DESC
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_channel"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT channel_group, ref_order_type,
       COUNT(*) order_count,
       SUM(outbound_qty) outbound_qty,
       SUM(outbound_ton) outbound_ton,
       AVG(fulfill_hours) avg_fulfill_hours,
       PERCENTILE_APPROX(fulfill_hours, 0.5) median_fulfill_hours,
       SUM(CASE WHEN fulfill_hours <= 24 THEN 1 ELSE 0 END) / COUNT(*) within_24h_rate,
       SUM(CASE WHEN fulfill_hours <= 48 THEN 1 ELSE 0 END) / COUNT(*) within_48h_rate,
       SUM(CASE WHEN fulfill_hours > 48 THEN 1 ELSE 0 END) over_48h_orders
FROM order_base
GROUP BY channel_group, ref_order_type
ORDER BY order_count DESC
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_overview"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT COUNT(*) completed_orders,
       SUM(CASE WHEN fulfill_hours > 48 THEN 1 ELSE 0 END) over48_orders,
       SUM(CASE WHEN fulfill_hours > 72 THEN 1 ELSE 0 END) over72_orders,
       SUM(CASE WHEN fulfill_hours > 120 THEN 1 ELSE 0 END) over120_orders,
       SUM(CASE WHEN fulfill_hours > 48 THEN outbound_qty ELSE 0 END) over48_qty,
       SUM(CASE WHEN fulfill_hours > 48 THEN outbound_ton ELSE 0 END) over48_ton,
       AVG(CASE WHEN fulfill_hours > 48 THEN fulfill_hours ELSE NULL END) over48_avg_hours,
       MAX(CASE WHEN fulfill_hours > 48 THEN fulfill_hours ELSE NULL END) over48_max_hours
FROM order_base
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_warehouse"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT warehouse_code, warehouse_name,
       COUNT(*) over48_orders,
       SUM(outbound_qty) over48_qty,
       SUM(outbound_ton) over48_ton,
       AVG(fulfill_hours) avg_hours,
       MAX(fulfill_hours) max_hours,
       SUM(CASE WHEN fulfill_hours > 72 THEN 1 ELSE 0 END) over72_orders,
       SUM(CASE WHEN fulfill_hours > 120 THEN 1 ELSE 0 END) over120_orders
FROM order_base
WHERE fulfill_hours > 48
GROUP BY warehouse_code, warehouse_name
ORDER BY over48_ton DESC, over48_orders DESC
LIMIT 200
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_channel"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT channel_group, ref_order_type,
       COUNT(*) over48_orders,
       SUM(outbound_qty) over48_qty,
       SUM(outbound_ton) over48_ton,
       AVG(fulfill_hours) avg_hours,
       MAX(fulfill_hours) max_hours,
       SUM(CASE WHEN fulfill_hours > 72 THEN 1 ELSE 0 END) over72_orders,
       SUM(CASE WHEN fulfill_hours > 120 THEN 1 ELSE 0 END) over120_orders
FROM order_base
WHERE fulfill_hours > 48
GROUP BY channel_group, ref_order_type
ORDER BY over48_ton DESC, over48_orders DESC
LIMIT 100
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_carrier"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT IFNULL(NULLIF(carrier_code,''),'未维护') carrier_code,
       COUNT(*) over48_orders,
       SUM(outbound_qty) over48_qty,
       SUM(outbound_ton) over48_ton,
       AVG(fulfill_hours) avg_hours,
       MAX(fulfill_hours) max_hours,
       SUM(CASE WHEN fulfill_hours > 72 THEN 1 ELSE 0 END) over72_orders,
       SUM(CASE WHEN fulfill_hours > 120 THEN 1 ELSE 0 END) over120_orders
FROM order_base
WHERE fulfill_hours > 48
GROUP BY IFNULL(NULLIF(carrier_code,''),'未维护')
ORDER BY over48_ton DESC, over48_orders DESC
LIMIT 100
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_created_date"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT created_date,
       COUNT(*) over48_orders,
       SUM(outbound_qty) over48_qty,
       SUM(outbound_ton) over48_ton,
       AVG(fulfill_hours) avg_hours,
       MAX(fulfill_hours) max_hours
FROM order_base
WHERE fulfill_hours > 48
GROUP BY created_date
ORDER BY created_date
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over48_error_hold"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT IFNULL(NULLIF(error_code,''),'无错误码') error_code,
       IFNULL(NULLIF(LEFT(error_desc, 180),''),'无错误描述') error_desc,
       IFNULL(NULLIF(hold_code,''),'无hold') hold_code,
       IFNULL(NULLIF(status_flow_code,''),'无状态流') status_flow_code,
       COUNT(*) over48_orders,
       SUM(outbound_qty) over48_qty,
       SUM(outbound_ton) over48_ton,
       AVG(fulfill_hours) avg_hours,
       MAX(fulfill_hours) max_hours
FROM order_base
WHERE fulfill_hours > 48
GROUP BY IFNULL(NULLIF(error_code,''),'无错误码'),
         IFNULL(NULLIF(LEFT(error_desc, 180),''),'无错误描述'),
         IFNULL(NULLIF(hold_code,''),'无hold'),
         IFNULL(NULLIF(status_flow_code,''),'无状态流')
ORDER BY over48_orders DESC, over48_ton DESC
LIMIT 200
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_bucket"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT CASE
         WHEN fulfill_hours <= 12 THEN '0-12小时'
         WHEN fulfill_hours <= 24 THEN '12-24小时'
         WHEN fulfill_hours <= 48 THEN '24-48小时'
         WHEN fulfill_hours <= 72 THEN '48-72小时'
         ELSE '72小时以上'
       END fulfill_bucket,
       COUNT(*) order_count,
       SUM(outbound_qty) outbound_qty,
       SUM(outbound_ton) outbound_ton
FROM order_base
GROUP BY fulfill_bucket
ORDER BY CASE fulfill_bucket
  WHEN '0-12小时' THEN 1
  WHEN '12-24小时' THEN 2
  WHEN '24-48小时' THEN 3
  WHEN '48-72小时' THEN 4
  ELSE 5 END
""",
        {"start": START, "end": END},
    )
    extra["order_status_distribution"] = fetch_all(
        conn,
        f"""
SELECT DATE(h.created) created_date,
       h.status,
       COUNT(DISTINCT h.outbound_code) order_count,
       SUM(IFNULL(d.ship_qty,0)) qty
FROM scm_sales_shipment_header h
JOIN scm_sales_shipment_detail d ON d.outbound_code=h.outbound_code
LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
LEFT JOIN master_wms mw ON mw.wms_code=h.facility_code
WHERE h.created >= %(start)s AND h.created < %(end)s
  AND h.outbound_code LIKE 'HJXH%%'
  AND ms.sku_group='50' AND mw.wms_type IN (10,11)
  {wh_filter}
GROUP BY DATE(h.created), h.status
ORDER BY created_date, h.status
""",
        {"start": START, "end": END},
    )
    created_abnormal_base = f"""
WITH created_order AS (
  SELECT h.outbound_code,
         DATE(h.created) created_date,
         h.created,
         h.ship_at,
         h.facility_code warehouse_code,
         MAX(mw.wms_desc) warehouse_name,
         h.ref_order_type,
         CASE
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z004' THEN '线上-JYCK/Z004'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' AND ms_store.group1='03' THEN '线下KA-JYCK/group1=03'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' AND ms_store.price_group='122' THEN '线下特价组-JYCK/price_group=122'
           WHEN h.ref_order_type='JYCK' AND ms_store.type='Z001' THEN '线下经销-JYCK/Z001'
           WHEN h.ref_order_type LIKE 'SO%%' THEN CONCAT('SO销售-', h.ref_order_type)
           ELSE h.ref_order_type
         END channel_group,
         h.status,
         IFNULL(h.carrier_code, '') carrier_code,
         IFNULL(h.error_code, '') error_code,
         IFNULL(LEFT(h.error_desc, 180), '') error_desc,
         IFNULL(h.hold_code, '') hold_code,
         IFNULL(h.status_flow_code, '') status_flow_code,
         COUNT(DISTINCT d.sku_code) sku_count,
         SUM(IFNULL(d.request_qty,0)) request_qty,
         SUM(IFNULL(d.ship_qty,0)) ship_qty,
         SUM(IFNULL(d.request_qty,0)*IFNULL(ms.net_weight,0)/1000) request_ton,
         SUM(IFNULL(d.ship_qty,0)*IFNULL(ms.net_weight,0)/1000) ship_ton,
         TIMESTAMPDIFF(MINUTE, h.created, %(end)s) / 60.0 age_hours
  FROM scm_sales_shipment_header h
  JOIN scm_sales_shipment_detail d ON d.outbound_code=h.outbound_code
  LEFT JOIN master_sku ms ON ms.sku_code=d.sku_code
  LEFT JOIN master_wms mw ON mw.wms_code=h.facility_code
  LEFT JOIN master_store ms_store ON ms_store.code=h.store_code
  WHERE h.created >= %(start)s AND h.created < %(end)s
    AND h.outbound_code LIKE 'HJXH%%'
    AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
    AND ms.sku_group='50' AND mw.wms_type IN (10,11)
    {wh_filter}
  GROUP BY h.outbound_code, DATE(h.created), h.created, h.ship_at, h.facility_code, h.ref_order_type,
           channel_group, h.status, h.carrier_code, h.error_code, LEFT(h.error_desc, 180), h.hold_code, h.status_flow_code
)
"""
    created_abnormal_where = "WHERE CAST(status AS CHAR) NOT IN ('700','900','701','901') OR ship_at IS NULL"
    extra["fulfillment_created_status_abnormal"] = fetch_all(
        conn,
        created_abnormal_base
        + f"""
SELECT status,
       IFNULL(NULLIF(error_code,''),'无错误码') error_code,
       IFNULL(NULLIF(error_desc,''),'无错误描述') error_desc,
       IFNULL(NULLIF(hold_code,''),'无hold') hold_code,
       IFNULL(NULLIF(status_flow_code,''),'无状态流') status_flow_code,
       IFNULL(NULLIF(carrier_code,''),'未维护') carrier_code,
       COUNT(*) order_count,
       SUM(request_qty) request_qty,
       SUM(ship_qty) ship_qty,
       SUM(request_ton) request_ton,
       SUM(ship_ton) ship_ton,
       AVG(age_hours) avg_age_hours
FROM created_order
{created_abnormal_where}
GROUP BY status,
         IFNULL(NULLIF(error_code,''),'无错误码'),
         IFNULL(NULLIF(error_desc,''),'无错误描述'),
         IFNULL(NULLIF(hold_code,''),'无hold'),
         IFNULL(NULLIF(status_flow_code,''),'无状态流'),
         IFNULL(NULLIF(carrier_code,''),'未维护')
ORDER BY order_count DESC, request_ton DESC
LIMIT 500
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_created_abnormal_warehouse"] = fetch_all(
        conn,
        created_abnormal_base
        + f"""
SELECT warehouse_code, warehouse_name, status,
       COUNT(*) order_count,
       SUM(request_qty) request_qty,
       SUM(ship_qty) ship_qty,
       SUM(request_ton) request_ton,
       SUM(ship_ton) ship_ton,
       AVG(age_hours) avg_age_hours
FROM created_order
{created_abnormal_where}
GROUP BY warehouse_code, warehouse_name, status
ORDER BY order_count DESC, request_ton DESC
LIMIT 300
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_created_abnormal_detail"] = fetch_all(
        conn,
        created_abnormal_base
        + f"""
SELECT outbound_code, created, ship_at, warehouse_code, warehouse_name, channel_group, ref_order_type,
       status, carrier_code, error_code, error_desc, hold_code, status_flow_code,
       sku_count, request_qty, ship_qty, request_ton, ship_ton, age_hours
FROM created_order
{created_abnormal_where}
ORDER BY age_hours DESC, request_ton DESC
LIMIT 1000
""",
        {"start": START, "end": END},
    )
    extra["fulfillment_over_48_detail"] = fetch_all(
        conn,
        fulfillment_base
        + """
SELECT outbound_code, created, ship_at, warehouse_code, warehouse_name, channel_group, ref_order_type,
       status, carrier_code, error_code, error_desc, hold_code, status_flow_code, sku_count, outbound_qty, outbound_ton, fulfill_hours
FROM order_base
WHERE fulfill_hours > 48
ORDER BY fulfill_hours DESC
LIMIT 1000
""",
        {"start": START, "end": END},
    )
    return extra


def add_week_and_risk(base: dict[str, list[dict[str, Any]]]) -> None:
    daily = base["daily_flow"]
    weekly: dict[str, dict[str, Any]] = {}
    for row in daily:
        dt = row["dt"]
        if isinstance(dt, str):
            day = date.fromisoformat(dt[:10])
        else:
            day = dt
        week_idx = ((day.day - 1) // 7) + 1
        label = f"第{week_idx}周"
        item = weekly.setdefault(
            label,
            {
                "week": label,
                "date_range": "",
                "outbound_order_count": 0,
                "outbound_qty": 0.0,
                "outbound_ton": 0.0,
                "purchase_order_count": 0,
                "purchase_qty": 0.0,
                "purchase_ton": 0.0,
                "stock_ton_end": 0.0,
            },
        )
        for key in ["outbound_order_count", "outbound_qty", "outbound_ton", "purchase_order_count", "purchase_qty", "purchase_ton"]:
            item[key] += to_float(row.get(key))
        item["stock_ton_end"] = to_float(row.get("stock_ton"))
        start_day = (week_idx - 1) * 7 + 1
        end_day = min(week_idx * 7, 30)
        item["date_range"] = f"6/{start_day}-6/{end_day}"
    base["weekly_trend"] = list(weekly.values())

    risks = []
    for row in base["sku_summary"]:
        outbound = to_float(row.get("outbound_ton"))
        closing = to_float(row.get("closing_stock_ton"))
        cover = row.get("closing_cover_days")
        cover_num = None if cover is None else to_float(cover)
        if outbound <= 0 and closing <= 0:
            continue
        risk = "正常"
        action = "按常规节奏跟进"
        if outbound > 100 and cover_num is not None and cover_num < 15:
            risk = "断货风险"
            action = "优先确认在途、采购排产和仓间调拨"
        elif closing > 100 and (cover_num is None or cover_num > 60):
            risk = "高覆盖"
            action = "评估调拨、促销或采购节奏下调"
        elif outbound <= 0 and closing > 20:
            risk = "无动销库存"
            action = "核对是否停采、清仓或渠道转移"
        risks.append(
            {
                "sku_code": row.get("sku_code"),
                "sku_name": row.get("sku_name"),
                "risk_tag": risk,
                "outbound_ton": outbound,
                "closing_stock_ton": closing,
                "closing_cover_days": cover,
                "purchase_ton": to_float(row.get("purchase_ton")),
                "suggested_action": action,
            }
        )
    order = {"断货风险": 0, "高覆盖": 1, "无动销库存": 2, "正常": 3}
    risks.sort(key=lambda x: (order.get(x["risk_tag"], 9), -to_float(x["outbound_ton"]), -to_float(x["closing_stock_ton"])))
    base["sku_risk"] = risks


def pct(numerator: float, denominator: float) -> float:
    return numerator / denominator if denominator else 0.0


def weighted_avg(rows: list[dict[str, Any]], value_key: str, weight_key: str) -> float:
    total_weight = sum(to_float(row.get(weight_key)) for row in rows)
    if not total_weight:
        return 0.0
    return sum(to_float(row.get(value_key)) * to_float(row.get(weight_key)) for row in rows) / total_weight


def top_names(rows: list[dict[str, Any]], name_key: str, value_key: str, limit: int = 5) -> str:
    filtered = [row for row in rows if to_float(row.get(value_key)) > 0]
    filtered.sort(key=lambda row: to_float(row.get(value_key)), reverse=True)
    return "；".join(f"{row.get(name_key)} {to_float(row.get(value_key)):.2f}吨" for row in filtered[:limit])


def find_row(rows: list[dict[str, Any]], predicate) -> dict[str, Any] | None:
    for row in rows:
        if predicate(row):
            return row
    return None


def classify_created_abnormal(row: dict[str, Any]) -> tuple[str, str]:
    text = " ".join(
        str(row.get(key) or "")
        for key in ["error_code", "error_desc", "hold_code", "status_flow_code", "carrier_code"]
    )
    status = str(row.get("status") or "")
    rules = [
        (
            "面单额度/余额",
            ["下单额度", "面单余额", "额度已经使用完毕", "面单额度", "账户余额", "余额不足", "余额不", "单余额"],
            "建立面单余额和下单额度日监控；余额低于阈值提前充值或切换备用承运商。",
        ),
        (
            "地址不详/超区/停发",
            ["地址不详细", "收方地址", "收件地址", "服务范围", "超出", "停发", "停业整改", "交通管制", "违规整改", "无法到达", "目的地"],
            "在订单推送前增加地址完整性和服务范围校验；超区订单自动改派可覆盖承运商。",
        ),
        (
            "订单已关闭/取消",
            ["未支付已关闭", "已关闭订单", "取消失败", "无需取单", "已取消"],
            "将已关闭或取消单从履约异常池剔除，并保留独立的关闭原因台账。",
        ),
        (
            "主数据/配置缺失",
            ["未找到对应的配置", "货主快递信息", "仓库编码不能为空", "主数据", "未维护"],
            "补齐货主-仓库-承运商配置；上线配置缺失拦截，避免订单进入仓端后再失败。",
        ),
        (
            "平台/接口异常",
            ["服务不可达", "系统异常", "开放平台", "接口", "调用", "超时", "callback", "OrderLine"],
            "建立失败重试队列和接口异常看板；超过30分钟未恢复的接口错误转人工处理。",
        ),
        (
            "可发库存不足",
            ["可发库存不足", "库存不足", "缺货"],
            "按SKU同步OMS与WMS可发库存；对缺货SKU先冻结销售或改仓发货。",
        ),
        (
            "平台规则/资料不符",
            ["资料不一致", "禁止", "禁寄", "泄露", "不符合", "拦截", "违规", "黑名单"],
            "把平台规则类失败前置到下单校验；禁寄、资料不符订单不进入仓端履约池。",
        ),
    ]
    for category, keywords, suggestion in rules:
        if any(keyword in text for keyword in keywords):
            return category, suggestion
    if not text.strip() or all(token in text for token in ["无错误码", "无错误描述"]):
        return (
            f"状态{status}无错误描述",
            "先确认状态码业务含义；将取消/关闭类订单与真实待履约订单分开统计，避免异常池失真。",
        )
    return "其他异常描述", "保留原始错误描述，按订单明细分派到平台、仓库或系统接口责任侧处理。"


def aggregate_created_abnormal(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    suggestions: dict[str, str] = {}
    samples: dict[str, list[str]] = {}
    for row in rows:
        category, suggestion = classify_created_abnormal(row)
        item = grouped.setdefault(
            category,
            {
                "原因类别": category,
                "订单数": 0.0,
                "申请件数": 0.0,
                "已发件数": 0.0,
                "申请吨数": 0.0,
                "已发吨数": 0.0,
                "加权平均滞留小时": 0.0,
                "_age_weight": 0.0,
            },
        )
        orders = to_float(row.get("order_count"))
        item["订单数"] += orders
        item["申请件数"] += to_float(row.get("request_qty"))
        item["已发件数"] += to_float(row.get("ship_qty"))
        item["申请吨数"] += to_float(row.get("request_ton"))
        item["已发吨数"] += to_float(row.get("ship_ton"))
        item["加权平均滞留小时"] += to_float(row.get("avg_age_hours")) * orders
        item["_age_weight"] += orders
        suggestions[category] = suggestion
        desc = str(row.get("error_desc") or "").strip()
        if desc and desc not in {"无错误描述", "None"}:
            samples.setdefault(category, [])
            if desc not in samples[category] and len(samples[category]) < 3:
                samples[category].append(desc[:120])
    output = []
    for category, item in grouped.items():
        weight = item.pop("_age_weight") or 0.0
        item["加权平均滞留小时"] = item["加权平均滞留小时"] / weight if weight else 0.0
        item["代表性错误"] = "；".join(samples.get(category, [])) or "无错误描述"
        item["处理建议"] = suggestions.get(category, "")
        output.append(item)
    output.sort(key=lambda row: (to_float(row.get("订单数")), to_float(row.get("申请吨数"))), reverse=True)
    return output


def build_fulfillment_issue_summary(base: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    overview = (base.get("fulfillment_over48_overview") or [{}])[0]
    completed_orders = to_float(overview.get("completed_orders"))
    over48_orders = to_float(overview.get("over48_orders"))
    over48_ton = to_float(overview.get("over48_ton"))
    over48_rate = pct(over48_orders, completed_orders)
    channel_rows = base.get("fulfillment_over48_channel", [])
    carrier_rows = base.get("fulfillment_over48_carrier", [])
    warehouse_rows = base.get("fulfillment_over48_warehouse", [])
    created_rows = base.get("fulfillment_over48_created_date", [])
    status_rows = base.get("fulfillment_created_status_abnormal", [])
    issue_rows = aggregate_created_abnormal(status_rows)

    so01 = find_row(channel_rows, lambda row: str(row.get("ref_order_type")) == "SO01")
    zt = find_row(carrier_rows, lambda row: str(row.get("carrier_code")) == "ZT")
    jd = find_row(carrier_rows, lambda row: str(row.get("carrier_code")) == "JD")
    old_created = [
        row
        for row in created_rows
        if row.get("created_date") and str(row.get("created_date"))[:10] < START.isoformat()
    ]
    old_orders = sum(to_float(row.get("over48_orders")) for row in old_created)
    old_ton = sum(to_float(row.get("over48_ton")) for row in old_created)
    top_issue = issue_rows[0] if issue_rows else {}

    rows = [
        {
            "问题": "已发货超48小时",
            "数据表现": f"{over48_orders:.0f}单，占已完成发货{over48_rate:.2%}；影响{over48_ton:.2f}吨，平均{to_float(overview.get('over48_avg_hours')):.1f}小时，最长{to_float(overview.get('over48_max_hours')):.1f}小时。",
            "原因判断": "异常单量占比不高，但吨数集中，不能只看订单数；需按吨数和仓库优先级处理。",
            "处理建议": "建立>24小时预警和>48小时闭环清单，按吨数排序跟进；超72/120小时订单单独复盘。",
        },
        {
            "问题": "大吨位超时集中在SO01/中通",
            "数据表现": (
                f"SO01超48为{to_float((so01 or {}).get('over48_orders')):.0f}单、{to_float((so01 or {}).get('over48_ton')):.2f}吨；"
                f"ZT承运超48为{to_float((zt or {}).get('over48_orders')):.0f}单、{to_float((zt or {}).get('over48_ton')):.2f}吨。"
            ),
            "原因判断": "SO销售大单/中通线路是吨数影响主因，和线上小单不是同一类问题。",
            "处理建议": "对SO01+ZT建立预约、波次、装车、回传四段时效；重点仓每日输出未发清单并注明卡点。",
        },
        {
            "问题": "仓库集中度高",
            "数据表现": top_names(warehouse_rows, "warehouse_name", "over48_ton", 6),
            "原因判断": "超48吨数集中在少数仓，适合按仓库点名复盘，而不是泛化为全网履约问题。",
            "处理建议": "按仓库拉取超48明细，复盘是否为预约排队、波次滞后、承运商揽收或系统回传延迟。",
        },
        {
            "问题": "历史创建小单拉高平均时长",
            "数据表现": f"6月前创建但6月发货的超48订单{old_orders:.0f}单、{old_ton:.2f}吨；JD承运商超48为{to_float((jd or {}).get('over48_orders')):.0f}单、{to_float((jd or {}).get('over48_ton')):.2f}吨。",
            "原因判断": "这类订单吨数小但时长极长，会显著拉高平均履约小时，主要应视为历史单/状态回传清理问题。",
            "处理建议": "将历史创建单从当月经营SLA中单列；清理京东代发历史状态和回传链路，避免干扰真实当月履约判断。",
        },
        {
            "问题": "未完成/未发货异常池",
            "数据表现": f"最大原因类别：{top_issue.get('原因类别','无')}，{to_float(top_issue.get('订单数')):.0f}单、申请{to_float(top_issue.get('申请吨数')):.2f}吨。",
            "原因判断": "创建在6月但未完成的订单需要先区分取消/关闭、真实待履约、取号失败和系统接口失败。",
            "处理建议": "先确认状态码含义，再把取消关闭类剔除；剩余订单按原因类别分派到平台、仓库、承运商和系统接口处理。",
        },
    ]
    for item in issue_rows[:6]:
        rows.append(
            {
                "问题": item["原因类别"],
                "数据表现": f"{to_float(item.get('订单数')):.0f}单，申请{to_float(item.get('申请件数')):.0f}件/{to_float(item.get('申请吨数')):.2f}吨，平均滞留{to_float(item.get('加权平均滞留小时')):.1f}小时。",
                "原因判断": str(item.get("代表性错误") or "无错误描述")[:240],
                "处理建议": item.get("处理建议", ""),
            }
        )
    return rows


def make_readme(summary_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {"工作表": "00_目录说明", "内容": "附件工作表目录、口径说明、可用于汇报的维度"},
        {"工作表": "01_核心指标", "内容": "出库、采购、库存、周转、覆盖天数等核心指标"},
        {"工作表": "02_日趋势", "内容": "6月每日出库、采购、库存"},
        {"工作表": "03_周趋势", "内容": "按自然周分组的出库、采购、周末库存"},
        {"工作表": "04_渠道结构", "内容": "按销售渠道/订单类型拆分出库"},
        {"工作表": "05_出库类型", "内容": "按 ref_order_type 拆分出库"},
        {"工作表": "06_仓库汇总", "内容": "仓库级出库、采购到货、库存和覆盖天数"},
        {"工作表": "07_仓库SKU明细", "内容": "仓库+SKU粒度的出库、期末库存、覆盖天数"},
        {"工作表": "08_SKU汇总", "内容": "SKU级出库、采购、库存、周转和覆盖天数"},
        {"工作表": "09_SKU风险", "内容": "按断货风险、高覆盖、无动销库存等标签排序"},
        {"工作表": "10_品类汇总", "内容": "按SKU主数据分类聚合"},
        {"工作表": "11_库龄结构", "内容": "期末库存按生产日期库龄段拆分"},
        {"工作表": "12_仓库库龄", "内容": "仓库+库龄段拆分"},
        {"工作表": "13_采购仓类型", "内容": "采购实收到货仓类型拆分，解释工厂/成品仓口径"},
        {"工作表": "14_采购供应商", "内容": "供应商维度采购实收"},
        {"工作表": "15_采购到货仓", "内容": "到货仓维度采购实收"},
        {"工作表": "16_履约异常结论", "内容": "只看异常履约：超48、未完成/未发货原因判断和处理建议"},
        {"工作表": "17_未完成原因汇总", "内容": "6月创建但未完成或未发货订单，按错误描述归类后的原因和建议"},
        {"工作表": "18_超48总览", "内容": "已发货订单中超过48小时的订单数、吨数、平均/最长时长"},
        {"工作表": "19_超48仓库", "内容": "超48小时已发货订单按仓库拆分，按吨数排序"},
        {"工作表": "20_超48渠道", "内容": "超48小时已发货订单按渠道/订单类型拆分"},
        {"工作表": "21_超48承运商", "内容": "超48小时已发货订单按 carrier_code 拆分"},
        {"工作表": "22_超48创建日", "内容": "超48小时订单按创建日期拆分，用于识别历史创建长尾"},
        {"工作表": "23_超48错误字段", "内容": "超48小时订单按 error/hold/status_flow 字段拆分"},
        {"工作表": "24_未完成状态原因", "内容": "6月创建但未完成或未发货订单的状态、错误描述、承运商拆分"},
        {"工作表": "25_未完成仓库", "内容": "6月创建但未完成或未发货订单按仓库和状态拆分"},
        {"工作表": "26_超48明细", "内容": "超48小时已发货订单样例明细，最多1000行"},
        {"工作表": "27_未完成明细", "内容": "6月创建但未完成或未发货订单样例明细，最多1000行"},
        {"工作表": "28_履约总览参考", "内容": "正常履约总览，仅作参考，不作为异常分析重点"},
        {"工作表": "29_履约日趋势参考", "内容": "按发货日期查看履约单量、吨数、履约小时和达成率，仅作参考"},
        {"工作表": "30_仓库履约参考", "内容": "仓库维度履约效率，仅作参考"},
        {"工作表": "31_渠道履约参考", "内容": "渠道维度履约效率，仅作参考"},
        {"工作表": "32_履约时效分布", "内容": "0-12小时、12-24小时、24-48小时、48小时以上等分布"},
        {"工作表": "33_订单状态分布", "内容": "按销售出库单创建日期查看订单状态"},
        {"工作表": "34_数据质量", "内容": "日期覆盖、行数、仓库数、SKU数"},
        {"工作表": "口径", "内容": "货主HJXH；成品sku_group=50；库存/出库限定10/11仓并剔除残次等；采购为全量CCP成品实收"},
    ]


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], freeze: str = "A2") -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["无数据"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E2EC")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="17365D")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, (int, float)):
                if any(key in str(ws.cell(1, cell.column).value) for key in ["ton", "days", "ratio"]):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
    for idx, header in enumerate(headers, 1):
        width = min(max(len(str(header)) + 2, 12), 38)
        for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=min(ws.max_row, 300)):
            for c in cell:
                width = min(max(width, len(str(c.value or "")) + 2), 42)
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_charts(wb: Workbook) -> None:
    if "02_日趋势" in wb.sheetnames:
        ws = wb["02_日趋势"]
        chart = LineChart()
        chart.title = "每日出库/采购吨数"
        chart.y_axis.title = "吨"
        data = Reference(ws, min_col=4, max_col=6, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 18
        chart.height = 8
        ws.add_chart(chart, "K2")
    if "04_渠道结构" in wb.sheetnames:
        ws = wb["04_渠道结构"]
        top = min(ws.max_row, 10)
        chart = BarChart()
        chart.title = "渠道出库吨数Top"
        data = Reference(ws, min_col=8, min_row=1, max_row=top)
        cats = Reference(ws, min_col=1, min_row=2, max_row=top)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 16
        chart.height = 8
        ws.add_chart(chart, "K2")
    if "13_采购仓类型" in wb.sheetnames:
        ws = wb["13_采购仓类型"]
        chart = BarChart()
        chart.title = "采购实收按到货仓类型"
        data = Reference(ws, min_col=6, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=2, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 15
        chart.height = 7
        ws.add_chart(chart, "H2")


def main() -> int:
    base_path = latest_base_workbook()
    src = load_workbook(base_path, data_only=True, read_only=True)
    base = {
        "summary_kpi": read_sheet(src, "summary_kpi"),
        "data_quality": read_sheet(src, "data_quality"),
        "daily_flow": read_sheet(src, "daily_flow"),
        "warehouse_summary": read_sheet(src, "warehouse_summary"),
        "sku_summary": read_sheet(src, "sku_summary"),
        "category_summary": read_sheet(src, "category_summary"),
        "channel_summary": read_sheet(src, "channel_summary"),
        "age_bucket": read_sheet(src, "age_bucket"),
        "top_aging_sku": read_sheet(src, "top_aging_sku"),
    }
    add_week_and_risk(base)
    load_env()
    with connect() as conn:
        base.update(query_extra(conn))
    base["fulfillment_created_issue_summary"] = aggregate_created_abnormal(base["fulfillment_created_status_abnormal"])
    base["fulfillment_issue_summary"] = build_fulfillment_issue_summary(base)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "00_目录说明", make_readme(base["summary_kpi"]))
    write_sheet(wb, "01_核心指标", base["summary_kpi"])
    write_sheet(wb, "02_日趋势", base["daily_flow"])
    write_sheet(wb, "03_周趋势", base["weekly_trend"])
    write_sheet(wb, "04_渠道结构", base["channel_summary"])
    write_sheet(wb, "05_出库类型", base["outbound_type"])
    write_sheet(wb, "06_仓库汇总", base["warehouse_summary"])
    write_sheet(wb, "07_仓库SKU明细", base["warehouse_sku_detail"])
    write_sheet(wb, "08_SKU汇总", base["sku_summary"])
    write_sheet(wb, "09_SKU风险", base["sku_risk"])
    write_sheet(wb, "10_品类汇总", base["category_summary"])
    write_sheet(wb, "11_库龄结构", base["age_bucket"])
    write_sheet(wb, "12_仓库库龄", base["warehouse_age"])
    write_sheet(wb, "13_采购仓类型", base["purchase_wms_type"])
    write_sheet(wb, "14_采购供应商", base["purchase_vendor"])
    write_sheet(wb, "15_采购到货仓", base["purchase_destination"])
    write_sheet(wb, "16_履约异常结论", base["fulfillment_issue_summary"])
    write_sheet(wb, "17_未完成原因汇总", base["fulfillment_created_issue_summary"])
    write_sheet(wb, "18_超48总览", base["fulfillment_over48_overview"])
    write_sheet(wb, "19_超48仓库", base["fulfillment_over48_warehouse"])
    write_sheet(wb, "20_超48渠道", base["fulfillment_over48_channel"])
    write_sheet(wb, "21_超48承运商", base["fulfillment_over48_carrier"])
    write_sheet(wb, "22_超48创建日", base["fulfillment_over48_created_date"])
    write_sheet(wb, "23_超48错误字段", base["fulfillment_over48_error_hold"])
    write_sheet(wb, "24_未完成状态原因", base["fulfillment_created_status_abnormal"])
    write_sheet(wb, "25_未完成仓库", base["fulfillment_created_abnormal_warehouse"])
    write_sheet(wb, "26_超48明细", base["fulfillment_over_48_detail"])
    write_sheet(wb, "27_未完成明细", base["fulfillment_created_abnormal_detail"])
    write_sheet(wb, "28_履约总览参考", base["fulfillment_overview"])
    write_sheet(wb, "29_履约日趋势参考", base["fulfillment_daily"])
    write_sheet(wb, "30_仓库履约参考", base["fulfillment_warehouse"])
    write_sheet(wb, "31_渠道履约参考", base["fulfillment_channel"])
    write_sheet(wb, "32_履约时效分布", base["fulfillment_bucket"])
    write_sheet(wb, "33_订单状态分布", base["order_status_distribution"])
    write_sheet(wb, "34_数据质量", base["data_quality"])
    add_charts(wb)
    wb.save(OUT_FILE)
    print(OUT_FILE)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
