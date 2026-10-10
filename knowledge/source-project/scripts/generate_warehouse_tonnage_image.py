from __future__ import annotations

import argparse
import base64
import hashlib
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import pymysql
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


def app_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parents[1]


ROOT = app_root()
OUTPUT_DIR = ROOT / "outputs"
DEFAULT_TEMPLATE_DIRS = [
    ROOT / "templates",
    Path(
    r"C:\Users\24493\Documents\WXWork\1688854282686093\Cache\File\2026-07"
    ),
]
TEMPLATE_NAME = "\u4ed3\u5e93\u51fa\u5165\u5e93\u5428\u6570\u53ef\u89c6\u5316\u56fe\u8868.xlsx"

ORANGE = "#F4B183"
HEADER_GRAY = "#D9D9D9"
WHITE = "#FFFFFF"
BLACK = "#000000"
RED = "#FF0000"

PREFERRED_WMS_CODES = {
    "\u7687\u5bb6\u5c0f\u864e\u4fdd\u5b9a\u4ee3\u53d1\u4ed3\uff08\u65b0\uff09xs": "WMS000281",
    "\u7687\u5bb6\u5c0f\u864e\u9752\u5c9b\u987a\u4e30\u4ed3\uff08\u65b0\uff09xs": "WMS000285",
}


@dataclass
class Metrics:
    toc_ton: float = 0.0
    tob_sales_ton: float = 0.0
    transfer_out_ton: float = 0.0
    purchase_in_ton: float = 0.0
    transfer_in_ton: float = 0.0

    @property
    def offline_total_ton(self) -> float:
        return self.tob_sales_ton + self.transfer_out_ton

    @property
    def inbound_ton(self) -> float:
        return self.purchase_in_ton + self.transfer_in_ton

    def add(self, other: "Metrics") -> None:
        self.toc_ton += other.toc_ton
        self.tob_sales_ton += other.tob_sales_ton
        self.transfer_out_ton += other.transfer_out_ton
        self.purchase_in_ton += other.purchase_in_ton
        self.transfer_in_ton += other.transfer_in_ton


def default_env_file() -> Path:
    for candidate in [ROOT / "config" / "config.env", ROOT / "config" / "pickup.env"]:
        if candidate.exists():
            return candidate
    return ROOT / "config" / "config.env"


def load_env_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"env file not found: {path}")
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def connect(env_file: Path) -> pymysql.connections.Connection:
    load_env_file(env_file)
    return pymysql.connect(
        host=os.environ["PICKUP_DB_HOST"],
        port=int(os.environ.get("PICKUP_DB_PORT", "9030")),
        user=os.environ["PICKUP_DB_USER"],
        password=os.environ["PICKUP_DB_PASSWORD"],
        database=os.environ.get("PICKUP_DB_NAME", "ywzt_ods"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=600,
        write_timeout=60,
    )


def previous_day() -> date:
    return datetime.now().date() - timedelta(days=1)


def find_template(path: Path | None) -> Path:
    if path:
        return path
    for directory in DEFAULT_TEMPLATE_DIRS:
        candidate = directory / TEMPLATE_NAME
        if candidate.exists():
            return candidate
    for directory in DEFAULT_TEMPLATE_DIRS:
        if directory.exists():
            files = sorted(directory.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
            if files:
                return files[0]
    raise FileNotFoundError("No template xlsx found in templates directory or WXWork cache")


def read_names(ws, col: int) -> list[str]:
    names: list[str] = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, col).value
        if value is None:
            continue
        text = str(value).strip()
        if text in {"\u5408\u8ba1", "\u603b\u8ba1\u5168\u4ed3"}:
            break
        if text:
            names.append(canonical_template_name(text))
    return names


def read_template(path: Path) -> dict[str, Any]:
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    return {
        "path": str(path),
        "headers": [ws.cell(1, col).value or "" for col in range(1, 15)],
        "toc_third": read_names(ws, 1),
        "toc_self": read_names(ws, 3),
        "tob_third": read_names(ws, 5),
        "tob_self": read_names(ws, 10),
    }


def norm_name(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text).lower()
    replacements = [
        ("\uff08", "("),
        ("\uff09", ")"),
        ("\u4e07\u7ef4", "\u4e07\u7eac"),
        ("\u5e38\u6e29-\u4e07\u7eac", "\u4e07\u7eac"),
        ("\u4ed3\u5e93", "\u4ed3"),
        ("\u4fdd\u5b9a-\u4ee3\u53d1\u4ed3(\u65b0)", "\u4fdd\u5b9a\u4ee3\u53d1\u4ed3"),
        ("\u4fdd\u5b9a\u4ee3\u53d1\u4ed3(\u65b0)", "\u4fdd\u5b9a\u4ee3\u53d1\u4ed3"),
        ("\u4fdd\u5b9a-\u4ee3\u53d1\u4ed3", "\u4fdd\u5b9a\u4ee3\u53d1\u4ed3"),
        ("\u9752\u5c9b\u987a\u4e30\u4ed3(\u65b0)", "\u9752\u5c9b\u987a\u4e30\u4ed3\u65b0"),
        ("\u65b0(2.0)", ""),
        ("(2.0)", ""),
        ("xs-\u65b0", ""),
        ("xs", ""),
        ("-", ""),
        ("_", ""),
        (" ", ""),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    text = re.sub(r"[()\[\]{}]", "", text)
    return text


def display_name(value: str) -> str:
    text = value.replace("\u7687\u5bb6\u5c0f\u864e", "")
    text = text.replace("\uff08" + "2.0" + "\uff09", "")
    text = text.replace("(2.0)", "")
    if text.endswith("xs"):
        text = text[:-2]
    return text


def canonical_template_name(value: str) -> str:
    mapping = {
        "\u7687\u5bb6\u5c0f\u864e\u4fdd\u5b9a\u4ee3\u53d1\u4ed3xs": "\u7687\u5bb6\u5c0f\u864e\u4fdd\u5b9a\u4ee3\u53d1\u4ed3\uff08\u65b0\uff09xs",
        "\u7687\u5bb6\u5c0f\u864e\u9752\u5c9b\u987a\u4e30\u4ed3xs": "\u7687\u5bb6\u5c0f\u864e\u9752\u5c9b\u987a\u4e30\u4ed3\uff08\u65b0\uff09xs",
    }
    return mapping.get(value, value)


def to_float(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except Exception:
        return 0.0


def fetch_metrics(conn: pymysql.connections.Connection, t1: str, t2: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    sql = """
WITH wms_sum_num AS (
    SELECT date(h.actual_arrival_date) dt,
           CASE WHEN h.warehous_type = '10' THEN 'cgrk'
                WHEN h.warehous_type = '11' THEN 'dbrk'
                WHEN h.warehous_type = '12' THEN 'thrk'
                WHEN h.warehous_type = 'TZRK' THEN 'tzrk'
                ELSE 'qtrk' END AS types,
           d.wms_code,
           d.sku_code,
           SUM(IFNULL(d.received_quantity, 0)) AS num
    FROM scm_warehous_detail d
    LEFT JOIN scm_warehous_header h
      ON h.warehous_code = d.warehous_code
     AND h.`status` = 0
    WHERE d.`status` = 0
      AND d.received_quantity IS NOT NULL
      AND d.received_quantity != 0
      AND h.actual_arrival_date BETWEEN %(t1)s AND %(t2)s
      AND h._fdl_marked_deleted = 'false'
      AND d._fdl_marked_deleted = 'false'
    GROUP BY h.warehous_type, d.sku_code, d.wms_code, date(h.actual_arrival_date), types

    UNION ALL

    SELECT date(h.ship_at) dt,
           CASE WHEN h.ref_order_type = 'DBCK' THEN 'dbck'
                WHEN h.ref_order_type = '\u5e93\u5b58\u6570\u91cf\u8c03\u6574\u5355' THEN 'tzck'
                WHEN h.terms_of_shipment = '20' THEN '2cck'
                WHEN h.terms_of_shipment != '20' THEN '2bck' END AS types,
           d.facility_code AS wms_code,
           d.sku_code,
           SUM(IFNULL(d.ship_qty, 0)) AS num
    FROM scm_sales_shipment_detail d
    LEFT JOIN scm_sales_shipment_header h
      ON h.outbound_code = d.outbound_code
     AND h.delete_flag = 0
    WHERE d.delete_flag = 0
      AND d.ship_qty IS NOT NULL
      AND d.ship_qty != 0
      AND h.ship_at BETWEEN %(t1)s AND %(t2)s
      AND h._fdl_marked_deleted = 'false'
      AND d._fdl_marked_deleted = 'false'
    GROUP BY h.ref_order_type, h.terms_of_shipment, d.facility_code, d.sku_code, date(h.ship_at), types
)
SELECT r.wms_code,
       w.wms_desc,
       SUM(IF(r.types = '2cck', r.num * IF(IFNULL(s.sku_class, '') LIKE '50%%', IFNULL(s.net_weight, 0), 0), 0)) / 1000 AS toc_ton,
       SUM(IF(r.types = '2bck', r.num * IF(IFNULL(s.sku_class, '') LIKE '50%%', IFNULL(s.net_weight, 0), 0), 0)) / 1000 AS tob_sales_ton,
       SUM(IF(r.types = 'dbck', r.num * IF(IFNULL(s.sku_class, '') LIKE '50%%', IFNULL(s.net_weight, 0), 0), 0)) / 1000 AS transfer_out_ton,
       SUM(IF(r.types = 'cgrk', r.num * IF(IFNULL(s.sku_class, '') LIKE '50%%', IFNULL(s.net_weight, 0), 0), 0)) / 1000 AS purchase_in_ton,
       SUM(IF(r.types = 'dbrk', r.num * IF(IFNULL(s.sku_class, '') LIKE '50%%', IFNULL(s.net_weight, 0), 0), 0)) / 1000 AS transfer_in_ton
FROM wms_sum_num r
LEFT JOIN master_sku s
  ON r.sku_code = s.sku_code
 AND s.`status` = 0
LEFT JOIN master_wms w
  ON r.wms_code = w.wms_code
 AND w.`status` = 0
GROUP BY r.wms_code, w.wms_desc
"""
    with conn.cursor() as cur:
        cur.execute(sql, {"t1": t1, "t2": t2})
        rows = list(cur.fetchall())
        cur.execute("SELECT wms_code, wms_desc, `status` FROM master_wms WHERE `status` = 0")
        master_rows = list(cur.fetchall())
    return rows, master_rows


def row_to_metrics(row: dict[str, Any]) -> Metrics:
    return Metrics(
        toc_ton=to_float(row.get("toc_ton")),
        tob_sales_ton=to_float(row.get("tob_sales_ton")),
        transfer_out_ton=to_float(row.get("transfer_out_ton")),
        purchase_in_ton=to_float(row.get("purchase_in_ton")),
        transfer_in_ton=to_float(row.get("transfer_in_ton")),
    )


def build_metric_lookup(
    data_rows: list[dict[str, Any]], master_rows: list[dict[str, Any]], template_names: Iterable[str]
) -> tuple[dict[str, Metrics], list[dict[str, Any]]]:
    data_by_code = {str(row["wms_code"]): row_to_metrics(row) for row in data_rows if row.get("wms_code")}
    data_norm: dict[str, list[dict[str, Any]]] = {}
    master_norm: dict[str, list[dict[str, Any]]] = {}
    for row in data_rows:
        data_norm.setdefault(norm_name(row.get("wms_desc")), []).append(row)
    for row in master_rows:
        master_norm.setdefault(norm_name(row.get("wms_desc")), []).append(row)

    lookup: dict[str, Metrics] = {}
    audit: list[dict[str, Any]] = []
    for name in template_names:
        normalized = norm_name(name)
        matched_codes: list[str] = []
        match_type = "none"
        matched_desc: list[str] = []

        preferred_code = PREFERRED_WMS_CODES.get(name)
        if preferred_code:
            matched_codes = [preferred_code]
            matched_desc = [
                str(row.get("wms_desc") or "")
                for row in master_rows
                if str(row.get("wms_code")) == preferred_code
            ]
            match_type = "preferred_code"
        elif normalized in data_norm:
            matched_codes = [str(row["wms_code"]) for row in data_norm[normalized] if row.get("wms_code")]
            matched_desc = [str(row.get("wms_desc") or "") for row in data_norm[normalized]]
            match_type = "data_exact"
        elif normalized in master_norm:
            matched_codes = [str(row["wms_code"]) for row in master_norm[normalized] if row.get("wms_code")]
            matched_desc = [str(row.get("wms_desc") or "") for row in master_norm[normalized]]
            match_type = "master_exact"
        else:
            candidates: list[dict[str, Any]] = []
            for key, rows in master_norm.items():
                if normalized and (normalized in key or key in normalized):
                    candidates.extend(rows)
            if candidates:
                candidates.sort(
                    key=lambda row: (
                        0 if str(row.get("wms_code")) in data_by_code else 1,
                        len(norm_name(row.get("wms_desc"))),
                    )
                )
                first = candidates[0]
                matched_codes = [str(first["wms_code"])]
                matched_desc = [str(first.get("wms_desc") or "")]
                match_type = "master_contains"

        metrics = Metrics()
        for code in matched_codes:
            metrics.add(data_by_code.get(code, Metrics()))
        lookup[name] = metrics
        audit.append(
            {
                "template_name": name,
                "normalized": normalized,
                "match_type": match_type,
                "matched_codes": matched_codes,
                "matched_desc": matched_desc,
                "metrics": metrics.__dict__,
            }
        )
    return lookup, audit


def total_metrics(items: Iterable[tuple[str, Metrics]]) -> Metrics:
    total = Metrics()
    for _, metrics in items:
        total.add(metrics)
    return total


def filter_template_names(template: dict[str, Any], keep_names: set[str]) -> dict[str, Any]:
    filtered = dict(template)
    for key in ["toc_third", "toc_self", "tob_third", "tob_self"]:
        filtered[key] = [name for name in template[key] if name in keep_names]
    return filtered


def summary_metrics(template: dict[str, Any], lookup: dict[str, Metrics]) -> dict[str, float]:
    toc_total = Metrics()
    tob_total = Metrics()
    for name in template["toc_third"] + template["toc_self"]:
        toc_total.add(lookup.get(name, Metrics()))
    for name in template["tob_third"] + template["tob_self"]:
        tob_total.add(lookup.get(name, Metrics()))
    return {
        "toc_out_ton": toc_total.toc_ton,
        "tob_total_out_ton": tob_total.offline_total_ton,
        "tob_sales_out_ton": tob_total.tob_sales_ton,
        "transfer_out_ton": tob_total.transfer_out_ton,
        "inbound_ton": tob_total.inbound_ton,
    }


def wecom_post(webhook: str, payload: dict[str, Any]) -> dict[str, Any]:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        webhook,
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            text = response.read().decode("utf-8", errors="replace")
    except urllib.error.URLError as exc:
        raise RuntimeError(f"WeCom webhook request failed: {exc}") from exc
    try:
        result = json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"WeCom webhook returned non-JSON response: {text[:200]}") from exc
    if result.get("errcode") not in (0, None):
        raise RuntimeError(f"WeCom webhook error: {result}")
    return result


def send_wecom_report(webhook: str, report_date: date, image_path: Path, summary: dict[str, float]) -> list[dict[str, Any]]:
    title = (
        f"{report_date:%Y-%m-%d} "
        "\u4ed3\u5e93\u51fa\u5165\u5e93\u5428\u6570\u65e5\u62a5\u5df2\u751f\u6210"
    )
    content = "\n".join(
        [
            title,
            f"TOC\u51fa\u5e93: {fmt_num(summary['toc_out_ton'])}\u5428",
            f"TOB\u603b\u51fa\u5e93: {fmt_num(summary['tob_total_out_ton'])}\u5428",
            f"TOB\u9500\u552e\u51fa\u5e93: {fmt_num(summary['tob_sales_out_ton'])}\u5428",
            f"\u8c03\u62e8\u51fa\u5e93: {fmt_num(summary['transfer_out_ton'])}\u5428",
            f"\u5165\u5e93: {fmt_num(summary['inbound_ton'])}\u5428",
        ]
    )
    results = [wecom_post(webhook, {"msgtype": "text", "text": {"content": content}})]
    image_bytes = image_path.read_bytes()
    image_payload = {
        "msgtype": "image",
        "image": {
            "base64": base64.b64encode(image_bytes).decode("ascii"),
            "md5": hashlib.md5(image_bytes).hexdigest(),
        },
    }
    results.append(wecom_post(webhook, image_payload))
    return results


def fmt_num(value: float) -> str:
    return f"{value:.2f}"


def font(path: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(path, size=size)


def get_fonts() -> dict[str, ImageFont.FreeTypeFont]:
    font_dir = Path(r"C:\Windows\Fonts")
    regular = font_dir / "msyh.ttc"
    bold = font_dir / "msyhbd.ttc"
    if not regular.exists():
        regular = font_dir / "simhei.ttf"
    if not bold.exists():
        bold = font_dir / "simhei.ttf"
    return {
        "title": font(str(bold), 30),
        "header": font(str(bold), 25),
        "body": font(str(bold), 25),
        "small": font(str(bold), 23),
    }


def text_bbox(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def draw_center(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fnt: ImageFont.FreeTypeFont,
    fill: str = BLACK,
) -> None:
    x1, y1, x2, y2 = box
    w, h = text_bbox(draw, text, fnt)
    draw.text((x1 + (x2 - x1 - w) / 2, y1 + (y2 - y1 - h) / 2 - 1), text, font=fnt, fill=fill)


def draw_left(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fnt: ImageFont.FreeTypeFont,
    fill: str = BLACK,
) -> None:
    x1, y1, _, y2 = box
    _, h = text_bbox(draw, text, fnt)
    draw.text((x1 + 8, y1 + (y2 - y1 - h) / 2 - 1), text, font=fnt, fill=fill)


def draw_multiline_center(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    lines: list[tuple[str, str]],
    fnt: ImageFont.FreeTypeFont,
    line_gap: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    heights = [text_bbox(draw, line, fnt)[1] for line, _ in lines]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for (line, color), h in zip(lines, heights):
        w, _ = text_bbox(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=color)
        y += h + line_gap


def header_lines(text: str, col_index: int) -> list[tuple[str, str]]:
    if col_index in {1, 3, 5, 10}:
        return [(text, BLACK)]
    custom: dict[int, list[tuple[str, str]]] = {
        2: [("\u5f53\u65e5\u51fa", BLACK), ("\u5e93\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        4: [("\u5f53\u65e5\u51fa\u5e93", BLACK), ("\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        6: [("\u5f53\u65e5\u7ebf\u4e0b", BLACK), ("\u603b\u51fa\u5e93", RED), ("\u91cd\u91cf(\u5428\uff09", BLACK)],
        7: [("\u5f53\u65e5TOB", BLACK), ("\u9500\u552e\u51fa\u5e93", RED), ("\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        8: [("\u5f53\u65e5\u8c03\u62e8", RED), ("\u51fa\u5e93\u91cd\u91cf", RED), ("(\u5428)", BLACK)],
        9: [("\u5f53\u65e5\u5165\u5e93", BLACK), ("\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        11: [("\u5f53\u65e5TOB", BLACK), ("\u603b\u51fa\u5e93", RED), ("\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        12: [("\u5f53\u65e5TOB", BLACK), ("\u9500\u552e\u51fa\u5e93", RED), ("\u91cd\u91cf", BLACK), ("(\u5428\uff09", BLACK)],
        13: [("\u5f53\u65e5\u8c03\u62e8", RED), ("\u51fa\u5e93\u91cd\u91cf", RED), ("(\u5428\uff0c\u7cfb", BLACK), ("\u7edf\u51fa\u5e93\u65f6", BLACK), ("\u95f4)", BLACK)],
        14: [("\u5f53\u65e5\u5165\u5e93", BLACK), ("\u91cd\u91cf(\u5428\uff09", BLACK)],
    }
    return custom.get(col_index, [(text, BLACK)])


def draw_cell(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    fill: str,
    border: str = BLACK,
    width: int = 1,
) -> None:
    draw.rectangle(box, fill=fill, outline=border, width=width)


def render_image(
    report_date: date,
    template: dict[str, Any],
    lookup: dict[str, Metrics],
    output: Path,
) -> None:
    fonts = get_fonts()
    width, height = 2540, 1905
    title_h, header_h = 68, 305
    body_rows = len(template["toc_third"]) + 1
    row_h = (height - title_h - header_h) / body_rows
    col_widths = [320, 110, 320, 110, 320, 126, 126, 126, 126, 320, 126, 126, 126, 128]
    x_edges = [0]
    for col_w in col_widths:
        x_edges.append(x_edges[-1] + col_w)

    image = Image.new("RGB", (width, height), WHITE)
    draw = ImageDraw.Draw(image)

    left_title = f"{report_date.year}/{report_date.month}/{report_date.day} \u4ed3\u5e93TOC\u53d1\u8d27\u6570\u636e\uff08\u51c0\u91cd\uff09"
    right_title = (
        f"{report_date.year}/{report_date.month}/{report_date.day} "
        "\u4ed3\u5e93\u6536\u8d27\u548c\u8c03\u62e8\u53d1\u8d27\u53caTOB\u9500\u552e\u53d1\u8d27\u6570\u636e\uff08\u51c0\u91cd\uff09"
    )
    draw_cell(draw, (0, 0, x_edges[4], title_h), ORANGE)
    draw_cell(draw, (x_edges[4], 0, width, title_h), ORANGE)
    draw_center(draw, (0, 0, x_edges[4], title_h), left_title, fonts["title"])
    draw_center(draw, (x_edges[4], 0, width, title_h), right_title, fonts["title"])

    for idx, header in enumerate(template["headers"], start=1):
        box = (x_edges[idx - 1], title_h, x_edges[idx], title_h + header_h)
        draw_cell(draw, box, HEADER_GRAY)
        draw_multiline_center(draw, box, header_lines(str(header), idx), fonts["header"])

    toc_third = sorted(
        [(name, lookup[name]) for name in template["toc_third"]],
        key=lambda item: (item[1].toc_ton, display_name(item[0])),
    )
    toc_self = sorted(
        [(name, lookup[name]) for name in template["toc_self"]],
        key=lambda item: (item[1].toc_ton, display_name(item[0])),
    )
    tob_third = sorted(
        [(name, lookup[name]) for name in template["tob_third"]],
        key=lambda item: (item[1].offline_total_ton, display_name(item[0])),
    )
    tob_self = sorted(
        [(name, lookup[name]) for name in template["tob_self"]],
        key=lambda item: (item[1].offline_total_ton, display_name(item[0])),
    )

    toc_third_total = total_metrics(toc_third)
    toc_self_total = total_metrics(toc_self)
    toc_all_total = Metrics()
    toc_all_total.add(toc_third_total)
    toc_all_total.add(toc_self_total)

    tob_third_total = total_metrics(tob_third)
    tob_self_total = total_metrics(tob_self)
    tob_all_total = Metrics()
    tob_all_total.add(tob_third_total)
    tob_all_total.add(tob_self_total)

    def row_box(row_idx: int, col_idx: int) -> tuple[int, int, int, int]:
        y1 = int(round(title_h + header_h + row_idx * row_h))
        y2 = int(round(title_h + header_h + (row_idx + 1) * row_h))
        return x_edges[col_idx - 1], y1, x_edges[col_idx], y2

    def draw_name_value(row_idx: int, name_col: int, value_col: int, name: str, value: float, total: bool = False) -> None:
        fill = HEADER_GRAY if total else WHITE
        draw_cell(draw, row_box(row_idx, name_col), fill)
        draw_cell(draw, row_box(row_idx, value_col), fill)
        draw_left(draw, row_box(row_idx, name_col), name, fonts["body"])
        draw_center(draw, row_box(row_idx, value_col), fmt_num(value), fonts["body"])

    def draw_right_row(row_idx: int, start_col: int, name: str, metrics: Metrics, total: bool = False) -> None:
        fill = HEADER_GRAY if total else WHITE
        values = [metrics.offline_total_ton, metrics.tob_sales_ton, metrics.transfer_out_ton, metrics.inbound_ton]
        for col in range(start_col, start_col + 5):
            draw_cell(draw, row_box(row_idx, col), fill)
        draw_left(draw, row_box(row_idx, start_col), name, fonts["body"])
        for offset, value in enumerate(values, start=1):
            draw_center(draw, row_box(row_idx, start_col + offset), fmt_num(value), fonts["body"])

    for idx, (name, metrics) in enumerate(toc_third):
        draw_name_value(idx, 1, 2, display_name(name), metrics.toc_ton)
    draw_name_value(len(toc_third), 1, 2, "\u5408\u8ba1", toc_third_total.toc_ton, total=True)

    for idx, (name, metrics) in enumerate(toc_self):
        draw_name_value(idx, 3, 4, display_name(name), metrics.toc_ton)
    draw_name_value(len(toc_self), 3, 4, "\u5408\u8ba1", toc_self_total.toc_ton, total=True)

    merge_start = len(toc_self) + 1
    merge_end = body_rows
    for col in [3, 4]:
        box = (
            x_edges[col - 1],
            int(round(title_h + header_h + merge_start * row_h)),
            x_edges[col],
            int(round(title_h + header_h + merge_end * row_h)),
        )
        draw_cell(draw, box, ORANGE)
        draw_center(draw, box, "\u603b\u8ba1\u5168\u4ed3" if col == 3 else fmt_num(toc_all_total.toc_ton), fonts["body"])

    for idx, (name, metrics) in enumerate(tob_third):
        draw_right_row(idx, 5, display_name(name), metrics)
    draw_right_row(len(tob_third), 5, "\u5408\u8ba1", tob_third_total, total=True)

    for idx, (name, metrics) in enumerate(tob_self):
        draw_right_row(idx, 10, display_name(name), metrics)
    draw_right_row(len(tob_self), 10, "\u5408\u8ba1", tob_self_total, total=True)

    merge_start = len(tob_self) + 1
    for col in range(10, 15):
        box = (
            x_edges[col - 1],
            int(round(title_h + header_h + merge_start * row_h)),
            x_edges[col],
            int(round(title_h + header_h + body_rows * row_h)),
        )
        draw_cell(draw, box, ORANGE)
        if col == 10:
            text = "\u603b\u8ba1\u5168\u4ed3"
        elif col == 11:
            text = fmt_num(tob_all_total.offline_total_ton)
        elif col == 12:
            text = fmt_num(tob_all_total.tob_sales_ton)
        elif col == 13:
            text = fmt_num(tob_all_total.transfer_out_ton)
        else:
            text = fmt_num(tob_all_total.inbound_ton)
        draw_center(draw, box, text, fonts["body"])

    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=None)
    parser.add_argument("--env-file", type=Path, default=None)
    parser.add_argument("--date", default=None, help="Report date in YYYY-MM-DD. Defaults to yesterday.")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--send", action="store_true", help="Send the generated image to WeCom.")
    parser.add_argument("--webhook", default=None, help="Override WECOM_WEBHOOK_URL from config.env.")
    args = parser.parse_args()

    report_date = datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else previous_day()
    t1 = f"{report_date:%Y-%m-%d} 00:00:00"
    t2 = f"{report_date:%Y-%m-%d} 23:59:59"
    env_file = args.env_file or default_env_file()
    template_path = find_template(args.template)
    template = read_template(template_path)

    conn = connect(env_file)
    try:
        data_rows, master_rows = fetch_metrics(conn, t1, t2)
    finally:
        conn.close()

    all_template_names = list(
        dict.fromkeys(template["toc_third"] + template["toc_self"] + template["tob_third"] + template["tob_self"])
    )
    lookup, audit = build_metric_lookup(data_rows, master_rows, all_template_names)
    keep_names = {item["template_name"] for item in audit if item["match_type"] != "none"}
    render_template = filter_template_names(template, keep_names)

    out_png = args.output_dir / f"warehouse_tonnage_{report_date:%Y%m%d}.png"
    out_audit = args.output_dir / f"warehouse_tonnage_{report_date:%Y%m%d}_audit.json"
    render_image(report_date, render_template, lookup, out_png)
    summary = summary_metrics(render_template, lookup)
    sent_to_wecom = False
    wecom_results: list[dict[str, Any]] = []
    wecom_error: str | None = None
    if args.send:
        webhook = args.webhook or os.environ.get("WECOM_WEBHOOK_URL", "").strip()
        if not webhook:
            raise SystemExit("missing WECOM_WEBHOOK_URL in config.env or --webhook")
        try:
            wecom_results = send_wecom_report(webhook, report_date, out_png, summary)
            sent_to_wecom = True
        except Exception as exc:
            wecom_error = str(exc)
    out_audit.write_text(
        json.dumps(
            {
                "report_date": str(report_date),
                "t1": t1,
                "t2": t2,
                "env_file": str(env_file),
                "template": template["path"],
                "doris_rows": len(data_rows),
                "summary": summary,
                "template_counts": {
                    "toc_third": len(template["toc_third"]),
                    "toc_self": len(template["toc_self"]),
                    "tob_third": len(template["tob_third"]),
                    "tob_self": len(template["tob_self"]),
                },
                "render_counts": {
                    "toc_third": len(render_template["toc_third"]),
                    "toc_self": len(render_template["toc_self"]),
                    "tob_third": len(render_template["tob_third"]),
                    "tob_self": len(render_template["tob_self"]),
                },
                "excluded_inactive_or_unmatched": [item for item in audit if item["match_type"] == "none"],
                "unmatched": [item for item in audit if item["match_type"] == "none"],
                "sent_to_wecom": sent_to_wecom,
                "wecom_error": wecom_error,
                "wecom_results": wecom_results,
                "matches": audit,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "png": str(out_png),
                "audit": str(out_audit),
                "sent_to_wecom": sent_to_wecom,
                "wecom_error": wecom_error,
                "summary": summary,
            },
            ensure_ascii=False,
        )
    )
    if wecom_error:
        raise SystemExit(wecom_error)


if __name__ == "__main__":
    main()
