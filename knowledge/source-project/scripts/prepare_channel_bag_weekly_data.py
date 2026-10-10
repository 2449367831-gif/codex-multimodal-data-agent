from __future__ import annotations

import argparse
import csv
import json
import os
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pymysql
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ENV = ROOT / "config" / "pickup.env"
DEFAULT_LEGACY_WORKBOOK = Path(r"D:\导出数据草稿D盘\分仓W30数据.xlsx")
DEFAULT_CURRENT_MAPPING = Path(
    r"C:\Users\24493\Documents\WXWork\1688854282686093\Cache\File\2026-07\基础信息表-SKU信息.xlsx"
)
TARGET_CATEGORIES = {"肉制品", "面点", "烘焙"}
RAW_COLUMNS = [
    "发货日期",
    "仓库编码",
    "仓库描述",
    "商品编码",
    "商品描述",
    "总发货量",
    "线上汇总",
    "抖音",
    "抖音小时达",
    "拼多多系",
    "快手系",
    "淘宝系",
    "京东系",
    "唯品会",
    "线上其它",
    "线下汇总",
    "线下社区团购_平台",
    "线下社区团购_分销",
    "线下其它",
    "其它",
]
DETAIL_COLUMNS = [*RAW_COLUMNS, "周报分类", "分类来源", "是否纳入汇总"]


SQL = """
SELECT
    DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d') AS 发货日期,
    h.facility_code AS 仓库编码,
    w.wms_desc AS 仓库描述,
    i.sku_code AS 商品编码,
    m.sku_name AS 商品描述,
    SUM(i.ship_qty) AS 总发货量,
    SUM(CASE WHEN s.type = 'Z004' THEN i.ship_qty ELSE 0 END) AS 线上汇总,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'DY' THEN i.ship_qty ELSE 0 END) AS 抖音,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'DYXD' THEN i.ship_qty ELSE 0 END) AS 抖音小时达,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('PDD','PDDDF') THEN i.ship_qty ELSE 0 END) AS 拼多多系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('KS','KSDF') THEN i.ship_qty ELSE 0 END) AS 快手系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('TGC','MJYX','TB','CT','MAOCHAO') THEN i.ship_qty ELSE 0 END) AS 淘宝系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code IN ('JDVMI','JD','JD_MARKET') THEN i.ship_qty ELSE 0 END) AS 京东系,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code = 'WPH' THEN i.ship_qty ELSE 0 END) AS 唯品会,
    SUM(CASE WHEN s.type = 'Z004' AND s.platform_code NOT IN ('DY','DYXD','PDD','PDDDF','KS','KSDF','TGC','MJYX','TB','CT','MAOCHAO','JDVMI','JD','JD_MARKET','WPH') THEN i.ship_qty ELSE 0 END) AS 线上其它,
    SUM(CASE WHEN s.type = 'Z001' THEN i.ship_qty ELSE 0 END) AS 线下汇总,
    SUM(CASE WHEN s.type = 'Z001' AND s.group1 = '03' THEN i.ship_qty ELSE 0 END) AS 线下社区团购_平台,
    SUM(CASE WHEN s.type = 'Z001' AND s.price_group = '122' THEN i.ship_qty ELSE 0 END) AS 线下社区团购_分销,
    SUM(CASE WHEN s.type = 'Z001' AND NOT (s.group1 = '03' OR s.price_group = '122') THEN i.ship_qty ELSE 0 END) AS 线下其它,
    SUM(CASE WHEN s.type NOT IN ('Z001','Z004') OR s.type IS NULL THEN i.ship_qty ELSE 0 END) AS 其它
FROM scm_sales_shipment_header h
INNER JOIN scm_sales_order_header h1 ON h1.sales_order_code = h.ref_order_code
INNER JOIN scm_sales_shipment_detail i ON i.outbound_code = h.outbound_code
LEFT JOIN master_store s ON s.code = h.store_code
LEFT JOIN master_sku m ON m.sku_code = i.sku_code
LEFT JOIN master_wms w ON w.wms_code = h.facility_code
WHERE h.ship_at >= %s AND h.ship_at < %s
  AND h.outbound_code LIKE 'HJXH%%'
  AND h.status IN ('700','900','701','901')
  AND (h.ref_order_type IN ('JYCK','HHCK','BFCK','QTCK') OR h.ref_order_type LIKE 'SO%%')
GROUP BY DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d'), h.facility_code, i.sku_code, m.sku_name, w.wms_desc
ORDER BY DATE_FORMAT(h.ship_at, '%%Y-%%m-%%d'), h.facility_code, i.sku_code
"""


@dataclass(frozen=True)
class WeekWindow:
    name: str
    start: date
    end: date

    @property
    def start_dt(self) -> datetime:
        return datetime.combine(self.start, time.min)

    @property
    def end_exclusive(self) -> datetime:
        return datetime.combine(self.end + timedelta(days=1), time.min)


def load_env_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def connect_doris():
    required = ["PICKUP_DB_HOST", "PICKUP_DB_USER", "PICKUP_DB_PASSWORD"]
    missing = [key for key in required if not os.environ.get(key, "").strip()]
    if missing:
        raise RuntimeError(f"缺少 Doris 配置：{', '.join(missing)}")
    return pymysql.connect(
        host=os.environ["PICKUP_DB_HOST"].strip(),
        port=int(os.environ.get("PICKUP_DB_PORT", "9030")),
        user=os.environ["PICKUP_DB_USER"].strip(),
        password=os.environ["PICKUP_DB_PASSWORD"],
        database=os.environ.get("PICKUP_DB_NAME", "ywzt_ods").strip() or "ywzt_ods",
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=240,
        write_timeout=60,
    )


def report_windows(report_date: date) -> tuple[WeekWindow, WeekWindow]:
    if report_date.weekday() != 4:
        raise ValueError("report_date 必须是周五")
    current_end = report_date - timedelta(days=1)
    current_start = current_end - timedelta(days=6)
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end - timedelta(days=6)

    def make_window(start: date, end: date) -> WeekWindow:
        return WeekWindow(f"W{start.isocalendar().week + 1}", start, end)

    return make_window(previous_start, previous_end), make_window(current_start, current_end)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_legacy_mapping(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["产品信息表"]
        mapping: dict[str, dict[str, Any]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sku = cell_text(row[1] if len(row) > 1 else None)
            if not sku:
                continue
            mapping[sku] = {
                "产品编码": sku,
                "旧商品名称": cell_text(row[5] if len(row) > 5 else None),
                "旧一级分类": cell_text(row[2] if len(row) > 2 else None),
                "旧二级分类": cell_text(row[3] if len(row) > 3 else None),
                "旧三级分类": cell_text(row[4] if len(row) > 4 else None),
                "旧箱规": row[6] if len(row) > 6 else None,
            }
        return mapping
    finally:
        workbook.close()


def read_current_mapping(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        headers = [cell_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        index = {name: position for position, name in enumerate(headers)}
        required = ["产品编码", "一级分类", "二级分类", "三级分类", "商品名称", "采购/售卖状态", "箱规", "产品大类"]
        missing = [name for name in required if name not in index]
        if missing:
            raise ValueError(f"新版映射表缺少字段：{', '.join(missing)}")

        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for values in sheet.iter_rows(min_row=2, values_only=True):
            sku = cell_text(values[index["产品编码"]])
            if not sku:
                continue
            grouped[sku].append(
                {
                    "产品编码": sku,
                    "新版商品名称": cell_text(values[index["商品名称"]]),
                    "新版一级分类": cell_text(values[index["一级分类"]]),
                    "新版二级分类": cell_text(values[index["二级分类"]]),
                    "新版三级分类": cell_text(values[index["三级分类"]]),
                    "德州产线": cell_text(values[index.get("德州产线", -1)]) if "德州产线" in index else "",
                    "采购/售卖状态": cell_text(values[index["采购/售卖状态"]]),
                    "新版箱规": values[index["箱规"]],
                    "产品大类": cell_text(values[index["产品大类"]]),
                }
            )

        mapping: dict[str, dict[str, Any]] = {}
        duplicates: dict[str, int] = {}
        for sku, records in grouped.items():
            mapping[sku] = records[0]
            mapping[sku]["新版重复记录数"] = len(records)
            mapping[sku]["新版全部商品名称"] = " / ".join(dict.fromkeys(record["新版商品名称"] for record in records))
            if len(records) > 1:
                duplicates[sku] = len(records)
        return mapping, duplicates
    finally:
        workbook.close()


def supplemental_category(record: dict[str, Any]) -> tuple[str, str, str]:
    level1 = cell_text(record.get("新版一级分类"))
    major = cell_text(record.get("产品大类"))
    name = cell_text(record.get("新版商品名称"))
    if major != "产成品":
        return "排除", "新版映射", f"产品大类={major or '空'}，非周报目标产成品"

    category_map = {
        "烤肠类": "肉制品",
        "调理类": "肉制品",
        "油炸类": "肉制品",
        "培根类": "肉制品",
        "预制菜肴类": "肉制品",
        "面制品类": "面点",
        "烘焙类": "烘焙",
        "烘焙": "烘焙",
    }
    category = category_map.get(level1)
    if category:
        return category, "根据商品属性补充", f"产品大类=产成品，一级分类={level1}，商品名称={name}"
    if level1 in {"酱料类", "调味料及复合配料", "周边物料", "冰激凌", "饱饱博士"}:
        return "排除", "新版映射", f"一级分类={level1}，非周报三类"
    return "", "待判断", f"无法根据新版属性可靠归类：一级分类={level1 or '空'}，商品名称={name}"


def build_mapping(legacy: dict[str, dict[str, Any]], current: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    for sku in sorted(set(legacy) | set(current)):
        old = legacy.get(sku, {})
        new = current.get(sku, {})
        old_category = cell_text(old.get("旧一级分类"))
        if old_category:
            report_category = old_category if old_category in TARGET_CATEGORIES else "排除"
            source = "原W30产品信息表"
            basis = f"原一级分类={old_category}，原分类优先"
        elif new:
            report_category, source, basis = supplemental_category(new)
        else:
            report_category, source, basis = "", "待判断", "两份映射表均无记录"
        merged[sku] = {
            "产品编码": sku,
            "商品名称": cell_text(new.get("新版商品名称")) or cell_text(old.get("旧商品名称")),
            "原一级分类": old_category,
            "新版一级分类": cell_text(new.get("新版一级分类")),
            "新版二级分类": cell_text(new.get("新版二级分类")),
            "新版三级分类": cell_text(new.get("新版三级分类")),
            "产品大类": cell_text(new.get("产品大类")),
            "采购/售卖状态": cell_text(new.get("采购/售卖状态")),
            "箱规": new.get("新版箱规") if new.get("新版箱规") not in (None, "") else old.get("旧箱规"),
            "周报分类": report_category,
            "分类来源": source,
            "分类依据": basis,
            "是否纳入汇总": "是" if report_category in TARGET_CATEGORIES else "否",
            "新版重复记录数": new.get("新版重复记录数", 0),
        }
    return merged


def decimal_value(value: Any) -> Decimal:
    if value is None:
        return Decimal(0)
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def csv_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None:
        return ""
    return value


def query_window(connection, window: WeekWindow) -> list[dict[str, Any]]:
    with connection.cursor() as cursor:
        cursor.execute(SQL, (window.start_dt, window.end_exclusive))
        return cursor.fetchall()


def classify_rows(rows: list[dict[str, Any]], mapping: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    classified = []
    for source_row in rows:
        row = dict(source_row)
        sku = cell_text(row.get("商品编码"))
        mapped = mapping.get(sku)
        if mapped:
            category = mapped["周报分类"]
            source = mapped["分类来源"]
        else:
            category = ""
            source = "待判断"
        row["周报分类"] = category
        row["分类来源"] = source
        row["是否纳入汇总"] = "是" if category in TARGET_CATEGORIES else "否"
        classified.append(row)
    return classified


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: csv_value(row.get(column)) for column in columns})


def aggregate(rows: list[dict[str, Any]]) -> dict[str, dict[str, Decimal]]:
    totals: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    for row in rows:
        category = cell_text(row.get("周报分类"))
        if category not in TARGET_CATEGORIES:
            continue
        for channel in ["抖音", "拼多多系", "快手系", "淘宝系", "京东系", "线下社区团购_平台", "线下社区团购_分销", "线下其它"]:
            totals[category][channel] += decimal_value(row.get(channel))
    return totals


def rounded(number: Decimal) -> int:
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def summary_values(previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]]) -> dict[str, Any]:
    previous = aggregate(previous_rows)
    current = aggregate(current_rows)
    channels = [
        ("抖音", ["抖音"]),
        ("拼多多", ["拼多多系"]),
        ("快手", ["快手系"]),
        ("淘宝", ["淘宝系"]),
        ("京东", ["京东系"]),
        ("社区团购", ["线下社区团购_平台", "线下社区团购_分销"]),
        ("线下大区", ["线下其它"]),
    ]
    output: dict[str, Any] = {}
    for label, source_columns in channels:
        output[label] = {}
        for category in ["肉制品", "面点", "烘焙"]:
            previous_total = sum((previous[category][name] for name in source_columns), Decimal(0))
            current_total = sum((current[category][name] for name in source_columns), Decimal(0))
            output[label][category] = {
                "previous_raw": format(previous_total, "f"),
                "current_raw": format(current_total, "f"),
                "previous": rounded(previous_total),
                "current": rounded(current_total),
            }
    return output


def unmapped_rows(
    previous_rows: list[dict[str, Any]], current_rows: list[dict[str, Any]], mapping: dict[str, dict[str, Any]], previous: WeekWindow, current: WeekWindow
) -> list[dict[str, Any]]:
    bucket: dict[str, dict[str, Any]] = {}
    for window, rows in [(previous, previous_rows), (current, current_rows)]:
        for row in rows:
            sku = cell_text(row.get("商品编码"))
            mapped = mapping.get(sku)
            if mapped and mapped.get("周报分类"):
                continue
            item = bucket.setdefault(
                sku,
                {
                    "商品编码": sku,
                    "商品描述": cell_text(row.get("商品描述")),
                    f"{previous.name}总发货量": Decimal(0),
                    f"{current.name}总发货量": Decimal(0),
                    "首次发货日期": cell_text(row.get("发货日期")),
                    "末次发货日期": cell_text(row.get("发货日期")),
                    "原因": mapped.get("分类依据") if mapped else "两份映射表均无此SKU",
                },
            )
            item[f"{window.name}总发货量"] += decimal_value(row.get("总发货量"))
            ship_date = cell_text(row.get("发货日期"))
            item["首次发货日期"] = min(item["首次发货日期"], ship_date)
            item["末次发货日期"] = max(item["末次发货日期"], ship_date)
    result = []
    for item in bucket.values():
        result.append({key: csv_value(value) for key, value in item.items()})
    return sorted(result, key=lambda item: -(Decimal(item[f"{previous.name}总发货量"] or 0) + Decimal(item[f"{current.name}总发货量"] or 0)))


def build_metadata(
    previous: WeekWindow,
    current: WeekWindow,
    previous_rows: list[dict[str, Any]],
    current_rows: list[dict[str, Any]],
    mapping: dict[str, dict[str, Any]],
    duplicates: dict[str, int],
    unmapped: list[dict[str, Any]],
) -> dict[str, Any]:
    mapping_rows = list(mapping.values())
    mapping_columns = [
        "产品编码",
        "商品名称",
        "原一级分类",
        "新版一级分类",
        "新版二级分类",
        "新版三级分类",
        "产品大类",
        "采购/售卖状态",
        "箱规",
        "周报分类",
        "分类来源",
        "分类依据",
        "是否纳入汇总",
        "新版重复记录数",
    ]
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "previous": {"name": previous.name, "start": previous.start.isoformat(), "end": previous.end.isoformat()},
        "current": {"name": current.name, "start": current.start.isoformat(), "end": current.end.isoformat()},
        "detail_columns": DETAIL_COLUMNS,
        "mapping_columns": mapping_columns,
        "mapping_rows": [{column: csv_value(row.get(column)) for column in mapping_columns} for row in mapping_rows],
        "unmapped_columns": ["商品编码", "商品描述", f"{previous.name}总发货量", f"{current.name}总发货量", "首次发货日期", "末次发货日期", "原因"],
        "unmapped_rows": unmapped,
        "summary": summary_values(previous_rows, current_rows),
        "audit": {
            "previous_query_rows": len(previous_rows),
            "current_query_rows": len(current_rows),
            "previous_total_ship_qty": format(sum((decimal_value(row.get("总发货量")) for row in previous_rows), Decimal(0)), "f"),
            "current_total_ship_qty": format(sum((decimal_value(row.get("总发货量")) for row in current_rows), Decimal(0)), "f"),
            "mapping_rows": len(mapping_rows),
            "included_mapping_rows": sum(1 for row in mapping_rows if row["是否纳入汇总"] == "是"),
            "unmapped_skus": len(unmapped),
            "new_mapping_duplicate_codes": duplicates,
            "classification_source_counts": dict(Counter(row["分类来源"] for row in mapping_rows)),
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="准备渠道袋数 W 周报数据")
    parser.add_argument("--report-date", type=lambda value: datetime.strptime(value, "%Y-%m-%d").date(), default=date.today())
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV)
    parser.add_argument("--legacy-workbook", type=Path, default=DEFAULT_LEGACY_WORKBOOK)
    parser.add_argument("--mapping-workbook", type=Path, default=DEFAULT_CURRENT_MAPPING)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    previous, current = report_windows(args.report_date)
    load_env_file(args.env_file)
    legacy = read_legacy_mapping(args.legacy_workbook)
    latest, duplicates = read_current_mapping(args.mapping_workbook)
    mapping = build_mapping(legacy, latest)

    connection = connect_doris()
    try:
        previous_rows = classify_rows(query_window(connection, previous), mapping)
        current_rows = classify_rows(query_window(connection, current), mapping)
    finally:
        connection.close()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "previous.csv", previous_rows, DETAIL_COLUMNS)
    write_csv(args.output_dir / "current.csv", current_rows, DETAIL_COLUMNS)
    missing = unmapped_rows(previous_rows, current_rows, mapping, previous, current)
    metadata = build_metadata(previous, current, previous_rows, current_rows, mapping, duplicates, missing)
    (args.output_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "ok": True,
                "output_dir": str(args.output_dir),
                "previous": metadata["previous"],
                "current": metadata["current"],
                "audit": metadata["audit"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
