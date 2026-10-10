from __future__ import annotations

import argparse
import json
import os
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import pymysql
import xlsxwriter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ENV_FILE = ROOT / "config" / "pickup.env"
DEFAULT_OUTPUT_DIR = ROOT / "outputs" / "计划部周报"


CARGO_OWNER_ID = "HJXH"
REPORT_DATE = date.today()
FDS_SKU_CODE = "5006070600000100"


PRODUCTS = [
    {"sku_code": "XHDDKCYW500g（LSHM）", "name": "小虎地道烤肠原味500g（LS）", "category": "地道烤肠", "box_spec": 20},
    {"sku_code": "XHDDKCHJW500g（LSHM）", "name": "小虎地道烤肠黑椒味500g（LS）", "category": "地道烤肠", "box_spec": 20},
    {"sku_code": "XHNRXB520g（LSHM）", "name": "小虎牛肉馅饼520g（LS）", "category": "牛肉馅饼", "box_spec": 18},
    {"sku_code": "XH-SZB800G（LSHM）", "name": "小虎好手抓饼10片装800G（LS）", "category": "手抓饼", "box_spec": 10},
    {"sku_code": "XHSGLMXLB200G（LSHM）", "name": "小虎手工老面小笼包200g（LS）", "category": "小笼包", "box_spec": 30},
    {"sku_code": "XHXGXRSM500g（LSHM）", "name": "小虎香菇猪肉烧麦500g（LS）", "category": "烧麦", "box_spec": 20},
    {"sku_code": "XH-MXGFWJRJARLW140G（LSHM）", "name": "小虎墨西哥风味鸡肉卷奥尔良味140g（LS）", "category": "鸡肉卷", "box_spec": 60},
    {"sku_code": "5001070014000500", "name": "皇家小虎蛋挞组合装0500G00", "category": "蛋挞组合装", "box_spec": 20},
    {"sku_code": "XHMZSJB600G", "name": "小虎拇指生煎包600G", "category": "生煎包", "box_spec": 20},
    {"sku_code": "5001020040000220", "name": "皇家小虎霸气手枪腿（奥尔良味）0220G00", "category": "手枪腿", "box_spec": 40},
    {"sku_code": "5001070012220600", "name": "皇家小虎商用葡式蛋挞皮0600G22", "category": "蛋挞皮", "box_spec": 20},
    {"sku_code": "5001070013220907", "name": "皇家小虎商用葡式蛋挞液0907G22", "category": "蛋挞液", "box_spec": 12},
    {"sku_code": "XHAELCG320g", "name": "小虎奥尔良翅根320g", "category": "翅根", "box_spec": 40},
    {"sku_code": "XHAELCZ300g", "name": "小虎奥尔良翅中300g", "category": "翅中", "box_spec": 40},
    {"sku_code": "XHZSNRJ120g", "name": "小虎芝士牛肉卷120g", "category": "牛肉卷", "box_spec": 60},
    {"sku_code": "5101010010009", "name": "零食很忙联名款热食盒2个装", "category": "蛋挞盒（小）", "box_spec": 600},
    {"sku_code": "5101010010008", "name": "零食很忙联名款热食盒6个装", "category": "蛋挞盒（大）", "box_spec": 300},
    {"sku_code": "5001070081210220", "name": "皇家小虎芝士培根披萨0220G21", "category": "披萨", "box_spec": 30},
    {"sku_code": "5001070073210220", "name": "皇家小虎榴莲果肉披萨0220G21", "category": "披萨", "box_spec": 30},
    {"sku_code": "5001070072210220", "name": "皇家小虎黑椒牛肉披萨0220G21", "category": "披萨", "box_spec": 30},
    {"sku_code": "5001010102213500", "name": "皇家小虎虎皮肉肠（黑椒味）3500G21", "category": "虎皮肉肠", "box_spec": 4},
    {"sku_code": "5001010101213500", "name": "皇家小虎虎皮肉肠（原味）3500G21", "category": "虎皮肉肠", "box_spec": 4},
    {"sku_code": "5006070600000100", "name": "肥大圣金枕榴莲果肉（去核）0100G00", "category": "肥大圣", "box_spec": 30},
    {"sku_code": "5001080020210002", "name": "皇家小虎香辣蘸料0002G21", "category": "香辣蘸料", "box_spec": 1000},
    {"sku_code": "5101010010013", "name": "皇家小虎烤肠船型盒", "category": "船型盒", "box_spec": 1000},
]


WAREHOUSES = [
    {"name": "皇家小虎长沙黄兴万纬仓xs", "short": "长沙"},
    {"name": "皇家小虎广州-弘一仓xs", "short": "广州"},
    {"name": "皇家小虎昆明万纬仓xs", "short": "昆明"},
    {"name": "皇家小虎郑州-万纬仓xs", "short": "郑州"},
    {"name": "皇家小虎成都-万纬仓", "short": "成都"},
    {"name": "皇家小虎南昌万纬仓xs", "short": "南昌"},
    {"name": "皇家小虎济南-万维仓xs", "short": "济南"},
    {"name": "皇家小虎合肥万纬仓xs", "short": "合肥"},
    {"name": "皇家小虎天津-万纬仓", "short": "天津"},
    {"name": "皇家小虎泉州万纬仓xs", "short": "泉州"},
]


FDS_WAREHOUSE_BY_REPORT_WAREHOUSE = {
    "皇家小虎长沙黄兴万纬仓xs": {"wms_code": "WMS000190", "wms_desc": "肥大圣长沙黄兴万纬仓"},
    "皇家小虎合肥万纬仓xs": {"wms_code": "WMS000187", "wms_desc": "肥大圣合肥万纬仓xs"},
    "皇家小虎南昌万纬仓xs": {"wms_code": "WMS000182", "wms_desc": "肥大圣南昌万纬仓xs"},
    "皇家小虎泉州万纬仓xs": {"wms_code": "WMS000183", "wms_desc": "肥大圣泉州万纬仓xs"},
}


BOX_SPEC_BY_SKU = {str(product["sku_code"]): int(product["box_spec"]) for product in PRODUCTS}


STORE_NAMES = [
    "线下-广西盛世齐帆供应链有限公司",
    "线下-广东横琴雨燕供应链管理有限公司",
    "线下-线下促销-零食很忙（赵一鸣）",
    "线下-线下促销样品-零食很忙（赵一鸣）",
    "线下-南京芙莱茜供应链管理有限公司",
    "线下-长沙富楚供应链管理有限公司",
    "线下-（地配）甘肃锦祥合盛商务企业(有限合伙)",
    "线下-新佳宜鲜食",
    "线下-天津润通供应链管理有限公司",
    "线下-陕西乐优益商贸有限公司",
    "线下-地配-牙克石市兔牙啃啃省钱超市(个体工商户)",
    "线下-地配-达拉特旗凯茜零食店(个体工商户)",
    "线下-上海柳城贸易有限公司",
    "地配-线下-阿巴嘎旗赵一鸣生活超市(个体工商户)",
    "线下-海南路每家实业有限公司",
    "地配-线下-锡林浩特市金汐食客零食超市(个体工商户)",
    "线下-阿荣旗那食品(个体工商户)",
    "线下-宿迁虎家食品销售有限公司",
    "线下-地配-开鲁县佳必得零食店(个体工商户)",
    "地配-线下-科尔沁右翼中旗思博零食店",
    "地配-线下-海拉区金灿灿零食店(个体工商户)",
    "线下-牙克石市鑫荣升副食超市",
    "地配-线下-海拉尔区九紫零食店(个体工商户)",
    "线下-恰货铺子",
    "线下-地配-海拉尔区溪源省钱超市(个体工商户)",
    "地配-线下-鄂伦春自治旗阿里河镇博博食品超市(个体工商户)",
    "线下-地配-扎赉特旗范洪刚便利店(个体工商户)",
    "地配-线下-鄂伦春自治旗大杨树镇米哆生活超市",
    "地配-线下-鄂温克族自治旗大雁镇追梦人食品店",
    "线下-地配-东胜区灿丰零食商行(个体工商户)",
    "线下-地配-苏尼特左旗有有水产店(个体工商户)",
    "地配-线下-土默特左旗世有零食店(个体工商户)",
    "地配-线下-乌兰浩特市妙食屋零食铺(个体工商户)",
    "线下-地配-内蒙古二连浩特市贝加尔街赵一鸣零食店",
    "线下-地配-库伦旗库伦镇欢欢零食店(个体工商户)",
    "地配-线下-阿荣旗那食品(个体工商户)",
    "线下-地配-扎兰屯市湘宁零食店(个体工商户)",
    "线下-地配-科尔沁区万柿大集小明商贸行(个体工商户)",
    "线下-地配-阿荣旗那吉镇超亮副食超市(个体工商户)",
    "地配-线下-丹巴县八零很忙超市(个体工商户)",
    "线下-地配-新巴尔虎左旗百亿食品销售便利店(个体工商户)",
    "地配-线下-锡林浩市一鸣零食店(个体工商户)",
    "线下-地配-新巴尔虎右旗阿拉坦额莫勒镇赵一鸣零食店(个体工商户)",
    "地配-线下-科尔沁右翼前旗随便批发超市(个体工商户)",
    "线下-地配-察哈尔右翼中旗建成零食店(个体工商户)",
    "地配-线下-武都区寒梅食品店(个体工商户)",
    "地配-线下-新龙县风云副食店(个体工商户)",
    "地配-线下-莫力达瓦达斡尔族自治旗尼尔基镇慕理零食店",
    "线下-地配-鄂伦春自治旗阿里河镇晟亿零食店(个体工商户)",
    "线下-地配-甘孜县小张零食店(个体工商户)",
    "线下-地配-炉霍县长实副实品经营店",
    "线下-地配-满洲里市婷婷食品店(个体工商户)",
    "线下-地配-福泽零食店",
    "线下-地配-白玉县小芒副食品批发店(个体工商户)",
    "线下-地配-海拉尔区哈晗超市(个体工商户)",
    "地配-线下-霍林郭勒市林零食店(个体工商户)",
    "地配-线下-东胜区思鸣便利便利店(个体工商户)",
    "地配-线下-旗八仙筒镇进进零食店(个体工商户)",
    "地配-线下-海拉尔区九紫零食店(全体工商户)",
    "线下-地配-化德县五角星量贩零食店(个体工商户)",
    "地配-线下-赛罕区天天红便利店(个体工商户)",
    "线下-地配-东胜区福亿超市",
    "线下-地配-乌兰浩特市赵一鸣零食店(个体工商户)",
    "线下-地配-西乌珠穆沁旗福泽零食店",
    "地配-线下-九龙县得万利百货超市(个体工商户)",
    "线下-地配-镶黄旗文梅超市(个体工商户)",
    "地配-线下-乌兰浩特市佳森零食铺(个体工商户)",
    "线下-地配-科尔沁区发辽零食店(个体工商户)",
    "地配-线下-宕昌县炳阳便利超市(个体工商户)",
    "地配-线下-海拉尔区金满仓零食店(个体工商户)",
    "地配-线下-包头市东河区乐满家零食店(个体工商户)",
    "线下-地配-锡林郭勒盟乌拉盖管理区嘉亿零食店(个体工商户)",
    "线下-地配-稻城利百德超市(个体工商户)",
    "地配-线下-陈巴尔虎旗巴彦库仁镇张大嘴零食超市(个体工商户)",
    "线下-地配-四子王旗品淳赵一鸣零售店(个体工商户)",
    "线下-地配-巴林左旗林东东城峰艳零食店",
    "线下-地配-武都区璟璟便利店(个体工商户)",
    "线下-地配-秦州区磨牙部落食品销售店(个体工商户)",
    "线下-地配-玉泉区明艺省芊零食店（个体工商户）",
    "线下-线下渠道-耿马冰花商贸有限公司（很忙云南临沧代配商）",
    "线下-地配-正镶白旗省钱超市(个体工商户)",
    "线下-地配-会理唐昊食品销售店(个体工商户)",
    "线下-地配-甘州区赵一鸣食品店(个体工商户)",
    "线下-长沙正深食品贸易有限公司",
    "地配-线下-阿鲁科尔沁旗赵一鸣零食店(个体工商户)",
    "地配-线下-牙克石市山水生鲜超市(个体工商户)",
    "地配-线下-鄂尔多斯市东胜区星鸣零食超市",
    "线下-长沙鸣望供应链管理有限公司",
    "地配-线下-鄂尔多斯市东胜区宏鸣零食超市(个体工商户",
    "地配-线下-满洲里市兔牙啃啃省钱超市(个体工商户)",
    "线下-地配-西昌市世达超市",
    "地配-线下-鄂尔多斯市东胜区盛鸣超市",
    "线下-芙莱茜（杭州）供应链有限公司",
    "线下-满洲里市兔牙啃啃省钱超市(个体工商户)",
    "线下-线下渠道-广东赵一鸣",
]


def load_env_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"env file not found: {path}")
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def connect():
    host = os.environ.get("PICKUP_DB_HOST", "").strip()
    user = os.environ.get("PICKUP_DB_USER", "").strip()
    password = os.environ.get("PICKUP_DB_PASSWORD", "")
    database = os.environ.get("PICKUP_DB_NAME", "ywzt_ods").strip() or "ywzt_ods"
    port = int(os.environ.get("PICKUP_DB_PORT", "9030"))
    missing = [
        name
        for name, value in {
            "PICKUP_DB_HOST": host,
            "PICKUP_DB_USER": user,
            "PICKUP_DB_PASSWORD": password,
        }.items()
        if not value
    ]
    if missing:
        raise SystemExit(f"missing database env vars: {', '.join(missing)}")
    return pymysql.connect(
        host=host,
        port=port,
        user=user,
        password=password,
        database=database,
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=20,
        read_timeout=240,
        write_timeout=60,
    )


def to_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def excel_number(value: Any) -> int | float:
    number = to_number(value)
    if abs(number - round(number)) < 0.000001:
        return int(round(number))
    return round(number, 6)


def box_number(value: Any) -> int:
    return int(round(to_number(value)))


def week_windows(report_date: date) -> list[dict[str, Any]]:
    end_date = report_date - timedelta(days=1)
    start_date = end_date - timedelta(days=27)
    weeks = []
    for offset in range(4):
        start = start_date + timedelta(days=offset * 7)
        end = start + timedelta(days=6)
        # The user-defined name is the Saturday calendar week + 1.
        week_no = start.isocalendar().week + 1
        weeks.append(
            {
                "name": f"W{week_no}",
                "start_date": start,
                "end_date": end,
                "start_dt": datetime.combine(start, time.min),
                "end_exclusive": datetime.combine(end + timedelta(days=1), time.min),
            }
        )
    return weeks


def placeholders(values: list[Any]) -> str:
    if not values:
        raise ValueError("empty values are not allowed")
    return ",".join(["%s"] * len(values))


def route_cte_sql(route_count: int) -> str:
    selects = []
    for idx in range(route_count):
        if idx == 0:
            selects.append(
                """
                SELECT
                    %s AS sku_code,
                    %s AS report_warehouse_name,
                    %s AS source_wms_code,
                    %s AS source_wms_desc
                """
            )
        else:
            selects.append("SELECT %s, %s, %s, %s")
    return "\nUNION ALL\n".join(selects)


def build_sales_routes(connection) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    fds_codes = [item["wms_code"] for item in FDS_WAREHOUSE_BY_REPORT_WAREHOUSE.values()]
    with connection.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT wms_code, wms_desc, status
            FROM master_wms
            WHERE wms_code IN ({placeholders(fds_codes)})
            """,
            fds_codes,
        )
        fds_master_rows = list(cursor.fetchall())

    active_fds_by_code = {
        str(row.get("wms_code") or ""): row
        for row in fds_master_rows
        if str(row.get("status") or "") == "0"
    }
    routes: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []

    for report_warehouse in WAREHOUSES:
        report_name = report_warehouse["name"]
        fds_warehouse = FDS_WAREHOUSE_BY_REPORT_WAREHOUSE.get(report_name)
        if not fds_warehouse:
            audit.append(
                {
                    "sku_code": FDS_SKU_CODE,
                    "report_warehouse_name": report_name,
                    "source_wms_code": None,
                    "source_wms_desc": None,
                    "route_type": "fds_no_system_mapping",
                    "enabled": False,
                }
            )
            continue

        source_wms_code = fds_warehouse["wms_code"]
        master_row = active_fds_by_code.get(source_wms_code)
        if not master_row:
            audit.append(
                {
                    "sku_code": FDS_SKU_CODE,
                    "report_warehouse_name": report_name,
                    "source_wms_code": source_wms_code,
                    "source_wms_desc": fds_warehouse["wms_desc"],
                    "route_type": "fds_not_found_or_disabled",
                    "enabled": False,
                }
            )
            continue

        source_wms_desc = str(master_row.get("wms_desc") or fds_warehouse["wms_desc"])
        routes.append(
            {
                "sku_code": FDS_SKU_CODE,
                "report_warehouse_name": report_name,
                "source_wms_code": source_wms_code,
                "source_wms_desc": source_wms_desc,
                "route_type": "fds_independent_warehouse",
            }
        )
        audit.append(
            {
                "sku_code": FDS_SKU_CODE,
                "report_warehouse_name": report_name,
                "source_wms_code": source_wms_code,
                "source_wms_desc": source_wms_desc,
                "route_type": "fds_independent_warehouse",
                "enabled": True,
            }
        )

    for product in PRODUCTS:
        sku_code = product["sku_code"]
        if sku_code == FDS_SKU_CODE:
            continue
        for report_warehouse in WAREHOUSES:
            report_name = report_warehouse["name"]
            routes.append(
                {
                    "sku_code": sku_code,
                    "report_warehouse_name": report_name,
                    "source_wms_code": None,
                    "source_wms_desc": report_name,
                    "route_type": "report_warehouse_name",
                }
            )

    return routes, audit


def fetch_sales_rows(connection, weeks: list[dict[str, Any]], sales_routes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not sales_routes:
        return []
    sku_codes = [item["sku_code"] for item in PRODUCTS]
    range_start = weeks[0]["start_dt"]
    range_end = weeks[-1]["end_exclusive"]
    route_sql = route_cte_sql(len(sales_routes))
    sql = f"""
        WITH route AS (
            {route_sql}
        )
        SELECT
            CASE
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
            END AS week_name,
            DATE(h.created) AS created_date,
            DATE(h.trade_created_at) AS trade_created_date,
            r.report_warehouse_name AS warehouse_name,
            h.facility_code AS source_wms_code,
            MAX(wms.wms_desc) AS source_warehouse_name,
            h.store_name AS store_name,
            d.sku_code AS sku_code,
            COALESCE(MAX(sku.sku_name), MAX(d.sku_name)) AS source_sku_name,
            SUM(d.ship_qty) AS ship_qty
        FROM scm_sales_order_header h
        INNER JOIN scm_sales_order_detail d
            ON h.sales_order_code = d.sales_order_code
           AND d.status = 0
           AND h.status NOT IN ('800','999')
        LEFT JOIN master_wms wms
            ON h.facility_code = wms.wms_code
        LEFT JOIN master_sku sku
            ON sku.sku_code = d.sku_code
        INNER JOIN route r
            ON r.sku_code = d.sku_code
           AND (
                (r.source_wms_code IS NOT NULL AND h.facility_code = r.source_wms_code)
                OR (r.source_wms_code IS NULL AND wms.wms_desc = r.source_wms_desc)
           )
        WHERE h.created >= %s
          AND h.created < %s
          AND h.store_name IN ({placeholders(STORE_NAMES)})
          AND d.sku_code IN ({placeholders(sku_codes)})
          AND h.biz_channel = 20
          AND d.ship_qty <> 0
        GROUP BY
            CASE
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
                WHEN h.created >= %s AND h.created < %s THEN %s
            END,
            DATE(h.created),
            DATE(h.trade_created_at),
            r.report_warehouse_name,
            h.facility_code,
            h.store_name,
            d.sku_code
        HAVING week_name IS NOT NULL
    """
    route_params: list[Any] = []
    for route in sales_routes:
        route_params.extend(
            [
                route["sku_code"],
                route["report_warehouse_name"],
                route["source_wms_code"],
                route["source_wms_desc"],
            ]
        )
    week_params: list[Any] = []
    for week in weeks:
        week_params.extend([week["start_dt"], week["end_exclusive"], week["name"]])
    params = (
        route_params
        + week_params
        + [range_start, range_end]
        + STORE_NAMES
        + sku_codes
        + week_params
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return list(cursor.fetchall())


def build_lookups(rows: list[dict[str, Any]]):
    week_sku_warehouse = defaultdict(float)
    raw_detail_rows = []
    for row in rows:
        week = str(row.get("week_name") or "")
        sku = str(row.get("sku_code") or "")
        warehouse = str(row.get("warehouse_name") or "")
        qty = to_number(row.get("ship_qty"))
        box_qty = box_number(qty / BOX_SPEC_BY_SKU[sku])
        week_sku_warehouse[(week, sku, warehouse)] += box_qty
        raw_detail_rows.append(
            {
                "week": week,
                "created_date": row.get("created_date"),
                "trade_created_date": row.get("trade_created_date"),
                "warehouse_name": warehouse,
                "source_wms_code": row.get("source_wms_code"),
                "source_warehouse_name": row.get("source_warehouse_name"),
                "store_name": row.get("store_name"),
                "sku_code": sku,
                "source_sku_name": row.get("source_sku_name"),
                "ship_qty": qty,
                "box_qty": box_qty,
            }
        )
    return week_sku_warehouse, raw_detail_rows


def category_rows(weeks: list[dict[str, Any]], week_sku_warehouse: dict[tuple[str, str, str], float]):
    rows = []
    categories: list[tuple[str, int]] = []
    seen = set()
    for product in PRODUCTS:
        key = (product["category"], product["box_spec"])
        if key not in seen:
            seen.add(key)
            categories.append(key)
    for category, box_spec in categories:
        week_values = []
        for week in weeks:
            total = 0.0
            for product in PRODUCTS:
                if product["category"] == category and product["box_spec"] == box_spec:
                    for warehouse in WAREHOUSES:
                        total += week_sku_warehouse.get((week["name"], product["sku_code"], warehouse["name"]), 0.0)
            week_values.append(excel_number(total))
        rows.append({"category": category, "box_spec": box_spec, "week_values": week_values})
    return rows


def write_workbook(
    output: Path,
    weeks: list[dict[str, Any]],
    week_sku_warehouse: dict[tuple[str, str, str], float],
    raw_detail_rows: list[dict[str, Any]],
    query_row_count: int,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(output)

    default = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 9, "align": "center", "valign": "vcenter", "border": 1}
    )
    left = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 9, "align": "left", "valign": "vcenter", "border": 1}
    )
    small = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 8, "align": "center", "valign": "vcenter", "border": 1}
    )
    small_left = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 8, "align": "left", "valign": "vcenter", "border": 1}
    )
    header = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 9,
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#F2F2F2",
        }
    )
    warehouse_header = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 8,
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#E2F0D9",
            "text_wrap": True,
        }
    )
    city_header = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 9,
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#FCE4D6",
        }
    )
    pink_header = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 10,
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#F4B6C2",
        }
    )
    total_fmt = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 9, "align": "center", "valign": "vcenter", "border": 1, "bg_color": "#FFF2CC"}
    )
    note_header = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 10,
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#E2F0D9",
        }
    )
    note_body = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 10, "align": "left", "valign": "vcenter", "border": 1}
    )
    qty_fmt = workbook.add_format(
        {"font_name": "微软雅黑", "font_size": 8, "align": "center", "valign": "vcenter", "border": 1, "num_format": "0"}
    )
    total_qty_fmt = workbook.add_format(
        {
            "font_name": "微软雅黑",
            "font_size": 8,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "bg_color": "#FFF2CC",
            "num_format": "0",
        }
    )

    def write_detail_sheet(sheet_name: str, week_names: list[str]) -> None:
        ws = workbook.add_worksheet(sheet_name)
        ws.hide_gridlines(2)
        ws.freeze_panes(3, 5)
        ws.set_column(0, 0, 5)
        ws.set_column(1, 1, 34)
        ws.set_column(2, 2, 12)
        ws.set_column(3, 3, 10)
        ws.set_column(4, 4, 7)
        ws.set_column(5, 14, 16)
        ws.set_column(15, 15, 10)
        ws.set_row(0, 18)
        ws.set_row(1, 18)
        ws.merge_range(2, 0, 2, 15, "单位：箱    范围：截图10仓    渠道：零食渠道", header)

        for col, name in enumerate(["序号", "产品名称", "分类", "渠道", "箱规"]):
            ws.merge_range(0, col, 1, col, name, header)
        for idx, warehouse in enumerate(WAREHOUSES, start=5):
            ws.write(0, idx, warehouse["name"], warehouse_header)
            ws.write(1, idx, warehouse["short"], city_header)
        ws.merge_range(0, 15, 1, 15, "合计", total_fmt)

        for row_idx, product in enumerate(PRODUCTS, start=3):
            excel_row = row_idx + 1
            ws.set_row(row_idx, 15)
            ws.write(row_idx, 0, row_idx - 2, small)
            ws.write(row_idx, 1, product["name"], small_left)
            ws.write(row_idx, 2, product["category"], small)
            ws.write(row_idx, 3, "零食渠道", small)
            ws.write(row_idx, 4, product["box_spec"], small)
            row_total = 0.0
            for col_idx, warehouse in enumerate(WAREHOUSES, start=5):
                value = 0.0
                for week_name in week_names:
                    value += week_sku_warehouse.get((week_name, product["sku_code"], warehouse["name"]), 0.0)
                row_total += value
                ws.write_number(row_idx, col_idx, excel_number(value), qty_fmt)
            ws.write_formula(row_idx, 15, f"=SUM(F{excel_row}:O{excel_row})", total_qty_fmt, excel_number(row_total))
        ws.autofilter(1, 0, 2 + len(PRODUCTS), 15)
        ws.set_landscape()
        ws.fit_to_pages(1, 0)

    ws = workbook.add_worksheet("口径说明")
    ws.hide_gridlines(2)
    ws.set_column(0, 0, 18)
    ws.set_column(1, 1, 88)
    notes = [
        ("运行日", "每周六"),
        ("本次运行日", REPORT_DATE.isoformat()),
        ("数据开始", f"{weeks[0]['start_date'].isoformat()} 00:00:00"),
        ("数据截止", f"{weeks[-1]['end_date'].isoformat()} 23:59:59"),
        ("周区间", "周六至周五，向前推4周"),
        ("周命名", ", ".join([f"{w['name']}={w['start_date']}至{w['end_date']}" for w in weeks])),
        ("切周字段", "scm_sales_order_header.created"),
        ("基础条件", "biz_channel=20; d.status=0; h.status NOT IN ('800','999'); d.ship_qty<>0"),
        ("输出口径", "汇总页、明细页和截图均按箱数展示：实发数量 ÷ 箱规"),
        ("仓库范围", f"{len(WAREHOUSES)}个截图仓库"),
        ("商品范围", f"{len(PRODUCTS)}个SKU，前两个地道烤肠按LS编码取数并展示LS名称"),
        ("SQL返回分组行数", query_row_count),
        ("原始明细行数", len(raw_detail_rows)),
    ]
    ws.write(0, 0, "项目", note_header)
    ws.write(0, 1, "内容", note_header)
    for row, (key, value) in enumerate(notes, start=1):
        ws.write(row, 0, key, note_body)
        ws.write(row, 1, value, note_body)

    summary = workbook.add_worksheet("分类趋势汇总")
    summary.hide_gridlines(2)
    summary.set_column(0, 0, 16)
    summary.set_column(1, 1, 7)
    summary.set_column(2, 5, 9)
    summary.set_column(6, 6, 10)
    summary.set_column(7, 7, 13)
    for col, value in enumerate(["产品分类", "箱规", *[w["name"] for w in weeks], "合计", "趋势图"]):
        summary.write(0, col, value, pink_header)

    for row_idx, row in enumerate(category_rows(weeks, week_sku_warehouse), start=1):
        excel_row = row_idx + 1
        summary.write(row_idx, 0, row["category"], default)
        summary.write(row_idx, 1, row["box_spec"], default)
        for col_idx, value in enumerate(row["week_values"], start=2):
            summary.write_number(row_idx, col_idx, value, default)
        summary.write_number(row_idx, 6, excel_number(sum(row["week_values"])), total_fmt)
        summary.write_blank(row_idx, 7, None, default)
        summary.add_sparkline(
            row_idx,
            7,
            {
                "range": f"分类趋势汇总!C{excel_row}:F{excel_row}",
                "type": "line",
                "markers": True,
                "series_color": "#5B9BD5",
                "markers_color": "#FF0000",
            },
        )
    summary.freeze_panes(1, 2)
    summary.set_landscape()
    summary.fit_to_pages(1, 0)

    write_detail_sheet("四周合计出库", [w["name"] for w in weeks])
    for week in weeks:
        write_detail_sheet(f"{week['name']} 出库", [week["name"]])

    raw = workbook.add_worksheet("取数明细")
    raw.hide_gridlines(2)
    raw_headers = [
        "周",
        "创建日期",
        "下单日期",
        "库存地点",
        "来源仓编码",
        "来源库存地点",
        "客户/店铺名称",
        "产品编码",
        "源产品名称",
        "实发数量",
        "箱数",
    ]
    for col, value in enumerate(raw_headers):
        raw.write(0, col, value, note_header)
    raw.set_column(0, 0, 8)
    raw.set_column(1, 2, 12)
    raw.set_column(3, 3, 24)
    raw.set_column(4, 4, 14)
    raw.set_column(5, 5, 24)
    raw.set_column(6, 6, 42)
    raw.set_column(7, 8, 28)
    raw.set_column(9, 10, 12)
    for row_idx, row in enumerate(raw_detail_rows, start=1):
        raw.write(row_idx, 0, row["week"], default)
        raw.write(row_idx, 1, str(row["created_date"]), default)
        raw.write(row_idx, 2, str(row["trade_created_date"]), default)
        raw.write(row_idx, 3, row["warehouse_name"], left)
        raw.write(row_idx, 4, row["source_wms_code"], left)
        raw.write(row_idx, 5, row["source_warehouse_name"], left)
        raw.write(row_idx, 6, row["store_name"], left)
        raw.write(row_idx, 7, row["sku_code"], left)
        raw.write(row_idx, 8, row["source_sku_name"], left)
        raw.write_number(row_idx, 9, excel_number(row["ship_qty"]), default)
        raw.write_number(row_idx, 10, box_number(row["box_qty"]), qty_fmt)
    if raw_detail_rows:
        raw.autofilter(0, 0, len(raw_detail_rows), len(raw_headers) - 1)
    raw.freeze_panes(1, 0)

    workbook.close()


def write_audit(
    output: Path,
    weeks: list[dict[str, Any]],
    rows: list[dict[str, Any]],
    raw_detail_rows: list[dict[str, Any]],
    route_count: int,
    route_audit: list[dict[str, Any]],
) -> Path:
    audit = {
        "output": str(output),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "report_date": REPORT_DATE.isoformat(),
        "range_start": f"{weeks[0]['start_date'].isoformat()} 00:00:00",
        "range_end": f"{weeks[-1]['end_date'].isoformat()} 23:59:59",
        "weeks": [
            {"name": w["name"], "start": w["start_date"].isoformat(), "end": w["end_date"].isoformat()}
            for w in weeks
        ],
        "sku_count": len(PRODUCTS),
        "warehouse_count": len(WAREHOUSES),
        "store_name_count": len(STORE_NAMES),
        "route_count": route_count,
        "fds_enabled_route_count": len([row for row in route_audit if row.get("enabled", False)]),
        "fds_route_audit": route_audit,
        "query_group_rows": len(rows),
        "raw_detail_rows": len(raw_detail_rows),
        "total_ship_qty": excel_number(sum(to_number(row.get("ship_qty")) for row in rows)),
        "total_box_qty": box_number(sum(to_number(row.get("box_qty")) for row in raw_detail_rows)),
    }
    audit_path = output.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    return audit_path


def verify_workbook(output: Path) -> dict[str, Any]:
    wb = load_workbook(output, read_only=True, data_only=False)
    result = {
        "sheets": wb.sheetnames,
        "summary_rows": wb["分类趋势汇总"].max_row,
        "summary_cols": wb["分类趋势汇总"].max_column,
        "total_sheet_rows": wb["四周合计出库"].max_row,
        "raw_rows": wb["取数明细"].max_row - 1,
    }
    wb.close()
    return result


def main() -> None:
    global REPORT_DATE

    parser = argparse.ArgumentParser(description="生成计划部周报：零食渠道四周出库")
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV_FILE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--report-date", type=lambda value: datetime.strptime(value, "%Y-%m-%d").date(), default=REPORT_DATE)
    args = parser.parse_args()

    REPORT_DATE = args.report_date

    load_env_file(args.env_file)
    weeks = week_windows(REPORT_DATE)
    output = args.output_dir / f"计划部周报_零食渠道出库_{weeks[0]['start_date']:%Y%m%d}_{weeks[-1]['end_date']:%Y%m%d}.xlsx"

    with connect() as connection:
        sales_routes, route_audit = build_sales_routes(connection)
        rows = fetch_sales_rows(connection, weeks, sales_routes)

    week_sku_warehouse, raw_detail_rows = build_lookups(rows)
    write_workbook(output, weeks, week_sku_warehouse, raw_detail_rows, len(rows))
    audit_path = write_audit(output, weeks, rows, raw_detail_rows, len(sales_routes), route_audit)
    verification = verify_workbook(output)

    print(json.dumps({"ok": True, "output": str(output), "audit": str(audit_path), **verification}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
